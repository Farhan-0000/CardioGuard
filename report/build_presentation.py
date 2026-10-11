#!/usr/bin/env python
"""
Build the review presentation ``report/MedBridge_Review_Presentation.pptx``.

Slide text and speaker notes live in ``report/presentation/slides.toml``; every number is
read at build time from the real results / measurement files through the report's results
context (``report/lib/results_context.py``) or computed from the code base
(``report/presentation/deck_data.py``). A result that does not exist yet is shown as a
highlighted ``[PENDING: ...]`` marker, and the three hallucination-study result slides
become "pending" slides until ``evaluation/results/summary.json`` exists. Re-running the builder
after new results land regenerates the whole deck.

Usage (from the repository root)::

    python report/build_presentation.py                 # PPTX only
    python report/build_presentation.py --pdf           # PPTX + PDF (LibreOffice)
    python report/build_presentation.py --png DIR       # also render every slide to PNG
    python report/build_presentation.py --strict        # exit 1 on layout or citation problems

Exit codes: 0 success; 1 ``--strict`` and a layout issue, unknown citation key or an
emergency example whose classifier output differs from the content file; 2 content file
missing or invalid; 3 ``--pdf``/``--png`` requested but LibreOffice could not convert.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import tomllib
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

REPORT_DIR = Path(__file__).resolve().parent
REPO_ROOT = REPORT_DIR.parent
if str(REPORT_DIR) not in sys.path:
    sys.path.insert(0, str(REPORT_DIR))

from lib import references as refs_lib  # noqa: E402
from lib import tables  # noqa: E402
from lib.libreoffice import find_soffice, pdf_page_count  # noqa: E402
from lib.results_context import ResultsContext  # noqa: E402
from pptx import Presentation  # noqa: E402
from presentation import deck_data  # noqa: E402
from presentation import pptx_kit as K  # noqa: E402
from presentation.pptx_kit import ACTION_TONES, TONES, P, Tone  # noqa: E402

DEFAULT_OUT = REPORT_DIR / "MedBridge_Review_Presentation.pptx"
DEFAULT_CONTENT = REPORT_DIR / "presentation" / "slides.toml"
DEFAULT_LOG = REPORT_DIR / "presentation" / "build" / "presentation_build_log.json"

# Layout grid (inches, 16:9 slide of 13.333 x 7.5 in)
MX = 0.6  # side margin
CW = K.SLIDE_W - 2 * MX  # content width
BODY_TOP = 1.45
BODY_BOTTOM = 6.55
SOURCE_Y = 6.6
FOOTER_Y = 7.02

VALUE_RE = re.compile(r"\{\{([^{}]+)\}\}")
CITE_RE = re.compile(r"\[(@[^\[\]]+)\]")
NUMBER_UNIT_RE = re.compile(r"(\d) (%|pp\b|ms\b|s\b|GB\b|mmHg\b)")
#: Slide-friendly names for results that do not exist yet (full paths go to the notes and log).
PENDING_LABELS: tuple[tuple[str, str], ...] = (
    ("study", "hallucination study"),
    ("deck.stats", "hallucination study"),
    ("deck.gate3_withdrawal", "V2 run records"),
    ("perf_latest.after", "after-optimisation latency probe"),
    ("deck.perf", "after-optimisation latency probe"),
    ("deck.robust.rabbits_before", "pre-fix RABBITS run"),
    ("deck.robust.adversarial_before", "pre-fix adversarial run"),
    ("deck.robust", "live run"),
    ("tests", "full test-run summary"),
    ("aq", "live MedBridge-AQ run"),
    ("rabbits", "live RABBITS run"),
    ("adversarial", "live adversarial run"),
    ("load", "load test"),
    ("calibration", "dev-split calibration"),
    ("human_agreement", "human annotation"),
)
#: Load-test request kinds that never call the LLM (fast-path and small-talk templates).
FAST_KINDS = ("emergency", "small_talk")
ORG_ABBREVIATIONS = {"World Health Organization": "WHO", "Parliament of India": "Govt. of India"}

REFERENCES_NOTE = "report/references/references.json (full entries in the notes)"
SCREENSHOT_CAPTURE = (
    "python -m medbridge.web.screenshots --base-url http://127.0.0.1:8000 --out report/screenshots"
)
LIVE_COMMANDS = {
    "aq": "python -m tests.benchmark.benchmark_runner --mode live --all --resume",
    "rabbits": "python -m tests.adversarial.rabbits_runner --mode live --resume",
    "adversarial": "python -m tests.adversarial.adversarial_runner",
}


# ----------------------------------------------------------------------
# Records
# ----------------------------------------------------------------------


@dataclass
class SlideRecord:
    index: int
    id: str
    title: str = ""
    state: str = "data"  # data | pending | partial
    pending: list[str] = field(default_factory=list)
    sources: list[str] = field(default_factory=list)
    citations: list[str] = field(default_factory=list)
    citations_missing: list[str] = field(default_factory=list)
    images: dict[str, bool] = field(default_factory=dict)
    notes_chars: int = 0

    def source(self, *paths: str | None) -> None:
        for p in paths:
            if p and p not in self.sources:
                self.sources.append(p)

    def as_dict(self) -> dict[str, Any]:
        return {
            "index": self.index,
            "id": self.id,
            "title": self.title,
            "state": self.state,
            "pending": self.pending,
            "sources": self.sources,
            "citations": self.citations,
            "citations_missing": self.citations_missing,
            "images": self.images,
            "notes_chars": self.notes_chars,
        }


def _median(values: Sequence[float]) -> float:
    vals = sorted(values)
    mid = len(vals) // 2
    return vals[mid] if len(vals) % 2 else (vals[mid - 1] + vals[mid]) / 2


def short_label(entry: dict[str, Any]) -> str:
    """Author-year label for slides: ``WHO 2023``, ``Gallifant et al. 2024``, ``Wołk 2025``."""
    authors = [str(a).strip() for a in entry.get("authors") or [] if str(a).strip()]
    et_al = any(a.lower().rstrip(".") == "et al" for a in authors)
    authors = [a for a in authors if a.lower().rstrip(".") != "et al"]
    year = entry.get("year", "")
    if not authors:
        return f"{entry.get('key', '?')} {year}".strip()
    first = authors[0]
    if refs_lib.is_organisation(first):
        abbr = re.search(r"\(([^()]+)\)\s*$", first)
        name = ORG_ABBREVIATIONS.get(first) or (abbr.group(1) if abbr else first)
        return f"{name} {year}".strip()
    surname = first.split()[-1]
    if et_al or len(authors) > 2:
        return f"{surname} et al. {year}".strip()
    if len(authors) == 2 and not refs_lib.is_organisation(authors[1]):
        return f"{surname} & {authors[1].split()[-1]} {year}".strip()
    return f"{surname} {year}".strip()


# ----------------------------------------------------------------------
# Builder
# ----------------------------------------------------------------------


class DeckBuilder:
    def __init__(
        self,
        content: dict[str, Any],
        repo_root: Path = REPO_ROOT,
        *,
        refs_root: Path | None = None,
    ) -> None:
        self.root = repo_root
        self.content = content
        self.deck = content.get("deck") or {}
        self.specs: list[dict[str, Any]] = [s for s in content.get("slides") or [] if s.get("id")]
        self.ctx = ResultsContext.load(repo_root)
        examples = (self.spec("emergency") or {}).get("examples") or []
        self.ctx.data["deck"] = deck_data.collect(repo_root, examples)
        self.refs = refs_lib.load_references(refs_root or repo_root)
        if not self.refs and refs_root is None:
            self.refs = refs_lib.load_references(REPO_ROOT)
        self.prs = Presentation()
        self.prs.slide_width = K.emu(K.SLIDE_W)
        self.prs.slide_height = K.emu(K.SLIDE_H)
        self.layout = self.prs.slide_layouts[6]  # blank
        self.kit = K.Kit()
        self.records: list[SlideRecord] = []
        self.cur = SlideRecord(0, "")
        self.example_mismatches: list[dict[str, Any]] = []

    # -- content helpers ----------------------------------------------

    def spec(self, slide_id: str) -> dict[str, Any] | None:
        return next((s for s in self.specs if s.get("id") == slide_id), None)

    def r(self, text: Any) -> str:
        """Render ``{{values}}`` and ``[@citations]`` in ``text``."""
        text = str(text or "")
        text = VALUE_RE.sub(lambda m: self._value(m.group(1)), text)
        text = CITE_RE.sub(lambda m: self._cite(m.group(1)), text)
        # keep units with their numbers ("5.8 %", "≈ 1 in 3") on one line
        text = NUMBER_UNIT_RE.sub("\\1\u00a0\\2", text)
        return text.replace("≈ ", "≈\u00a0")

    def _value(self, expr: str) -> str:
        path = expr.split("|")[0].strip()
        out, ok = self.ctx.render(expr)
        if not ok:
            if path not in self.cur.pending:
                self.cur.pending.append(path)
            label = next(
                (lab for pre, lab in PENDING_LABELS if path == pre or path.startswith(pre + ".")),
                None,
            )
            if label is None and path.startswith("deck."):
                label = path.rsplit(".", 1)[-1].replace("_", " ").lower()
            return f"[PENDING: {label}]" if label else out
        top = path.split(".")[0]
        if top == "deck":
            parts = path.split(".")
            src = self.ctx.get(".".join(parts[:2]) + ".source") if len(parts) > 1 else None
            if src is None and len(parts) > 2:
                src = self.ctx.get(".".join(parts[:3]) + ".source")
            self.cur.source(src if isinstance(src, str) else None)
        else:
            self.cur.source(self.ctx.source(top))
        return out

    def _cite(self, body: str) -> str:
        keys = refs_lib.parse_cite_keys(body)
        labels = []
        for key in keys:
            entry = self.refs.get(key)
            if entry is None:
                if key not in self.cur.citations_missing:
                    self.cur.citations_missing.append(key)
                labels.append(f"[CITATION NOT FOUND: {key}]")
                continue
            if key not in self.cur.citations:
                self.cur.citations.append(key)
            labels.append(short_label(entry))
        return "(" + "; ".join(labels) + ")"

    def get(self, path: str, default: Any = None) -> Any:
        return self.ctx.get(path, default)

    # -- slide chrome -------------------------------------------------

    def new_slide(self, spec: dict[str, Any], *, dark: bool = False) -> Any:
        index = len(self.records) + 1
        self.cur = SlideRecord(index, str(spec.get("id")))
        self.records.append(self.cur)
        self.kit.begin(f"{index:02d}-{spec.get('id')}")
        slide = self.prs.slides.add_slide(self.layout)
        K.background(slide, K.NAVY_DARK if dark else K.WHITE)
        return slide

    def chrome(self, slide: Any, spec: dict[str, Any], *, title: str | None = None) -> None:
        kit = self.kit
        title = self.r(title if title is not None else spec.get("title", ""))
        self.cur.title = title
        kit.rect(slide, 0, 0, K.SLIDE_W, 0.09, fill=K.NAVY, radius=None)
        kit.rect(slide, MX, 0.36, 0.07, 0.95, fill=K.AMBER, radius=None)
        kit.text(
            slide,
            MX + 0.2,
            0.3,
            9.5,
            0.32,
            P(self.r(spec.get("kicker", "")).upper(), bold=True),
            size=12,
            min_size=10,
            color=K.AMBER,
            name="kicker",
        )
        kit.text(
            slide,
            MX + 0.2,
            0.6,
            CW - 0.2,
            0.72,
            P(title, bold=True),
            size=28,
            min_size=20,
            color=K.NAVY,
            name="title",
            anchor="t",
        )
        kit.rect(slide, MX, FOOTER_Y, CW, 0.012, fill=K.RULE, radius=None)
        kit.text(
            slide,
            MX,
            FOOTER_Y + 0.06,
            8.0,
            0.3,
            self.r(self.deck.get("footer", "")),
            size=10,
            color=K.SUBTLE,
            name="footer",
        )
        kit.text(
            slide,
            K.SLIDE_W - MX - 1.5,
            FOOTER_Y + 0.06,
            1.5,
            0.3,
            P(f"{self.cur.index} / {len(self.specs)}", align="r"),
            size=10,
            color=K.SUBTLE,
            name="page-number",
        )

    def sources_line(self, slide: Any, extra: Sequence[str] = ()) -> None:
        for e in extra:
            self.cur.source(e)
        if self.cur.citations:
            self.cur.source(REFERENCES_NOTE)
        if not self.cur.sources:
            return
        self.kit.text(
            slide,
            MX,
            SOURCE_Y,
            CW,
            0.36,
            "Sources: " + " · ".join(self.cur.sources),
            size=10,
            min_size=7.5,
            color=K.SUBTLE,
            name="sources",
        )

    def set_notes(self, slide: Any, text: str) -> None:
        body = self.r(text).strip()
        tail = []
        if self.cur.sources:
            tail.append("Sources: " + "; ".join(self.cur.sources) + ".")
        if self.cur.pending:
            tail.append("Pending values (no result file yet): " + ", ".join(self.cur.pending) + ".")
        if self.cur.citations:
            lines = ["References cited on this slide:"]
            for key in self.cur.citations:
                entry = self.refs.get(key)
                if entry:
                    lines.append(f"- {short_label(entry)}: {refs_lib.format_entry_text(entry)}")
            tail.append("\n".join(lines))
        full = body + ("\n\n" + "\n\n".join(tail) if tail else "")
        self.cur.notes_chars = len(full)
        K.notes(slide, full)

    def heading(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        text: str,
        *,
        size: float = 18,
        color: str = K.NAVY,
        name: str = "heading",
    ) -> float:
        h = size * K.LINE_FACTOR / 72.0 + 0.1
        self.kit.text(
            slide,
            x,
            y,
            w,
            h,
            P(self.r(text), bold=True),
            size=size,
            min_size=size - 4,
            color=color,
            name=name,
        )
        return y + h

    def bullets(self, items: Sequence[str], *, color: str | None = None, gap: float = 5) -> list[P]:
        return [
            P(self.r(t), bullet=True, space_before=gap if i else 0, bullet_color=color)
            for i, t in enumerate(items)
        ]

    # -- build --------------------------------------------------------

    def build(self) -> Any:
        layouts: dict[str, Callable[[dict[str, Any]], None]] = {
            "title": self.s_title,
            "problem": self.s_problem,
            "gap": self.s_gap,
            "architecture": self.s_architecture,
            "gate": self.s_gate,
            "emergency": self.s_emergency,
            "retrieval": self.s_retrieval,
            "privacy": self.s_privacy,
            "ui": self.s_ui,
            "evaluation": self.s_evaluation,
            "models": self.s_models,
            "results_hallucination": self.s_results_hallucination,
            "results_unsafe": self.s_results_unsafe,
            "results_tradeoff": self.s_results_tradeoff,
            "robustness": self.s_robustness,
            "performance": self.s_performance,
            "testing": self.s_testing,
            "limitations": self.s_limitations,
            "thanks": self.s_thanks,
        }
        for spec in self.specs:
            layout = layouts.get(str(spec["id"]), self.s_generic)
            layout(spec)
        return self.prs

    # ==================================================================
    # Slides
    # ==================================================================

    def s_title(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec, dark=True)
        kit = self.kit
        self.cur.title = self.r(spec.get("title"))
        light, soft, gold = "E8EDF6", "B9C4DA", "F2C14E"
        kit.rect(slide, 0, 0, K.SLIDE_W, 0.09, fill=K.AMBER, radius=None)
        kit.text(
            slide,
            0.75,
            0.75,
            7.4,
            0.4,
            P(self.r(spec.get("kicker", "")).upper(), bold=True),
            size=13,
            min_size=10,
            color=gold,
            name="kicker",
        )
        kit.text(
            slide,
            0.75,
            1.15,
            7.4,
            1.0,
            P(self.cur.title, bold=True),
            size=54,
            min_size=36,
            color=K.WHITE,
            name="title",
        )
        kit.text(
            slide,
            0.75,
            2.2,
            7.4,
            1.1,
            self.r(spec.get("subtitle", "")),
            size=19,
            min_size=15,
            color=soft,
            name="subtitle",
        )
        kit.text(
            slide,
            0.75,
            3.42,
            7.4,
            0.7,
            self.r(spec.get("tagline", "")),
            size=16,
            min_size=12,
            color=light,
            name="tagline",
        )
        kit.rect(slide, 0.75, 4.2, 1.1, 0.05, fill=gold, radius=None)
        kit.text(
            slide,
            0.75,
            4.38,
            4.4,
            0.3,
            P("TEAM", bold=True),
            size=11,
            color=gold,
            name="team-label",
        )
        kit.text(
            slide,
            0.75,
            4.68,
            4.4,
            0.8,
            [self.r(t) for t in spec.get("team") or []],
            size=15,
            min_size=11,
            color=K.WHITE,
            name="team",
        )
        kit.text(
            slide,
            5.35,
            4.38,
            2.8,
            0.3,
            P("GUIDE", bold=True),
            size=11,
            color=gold,
            name="guide-label",
        )
        kit.text(
            slide,
            5.35,
            4.68,
            2.8,
            0.4,
            self.r(spec.get("guide", "")),
            size=15,
            min_size=11,
            color=K.WHITE,
            name="guide",
        )
        inst = [
            self.r(spec.get(k, "")) for k in ("institution", "department", "event") if spec.get(k)
        ]
        kit.text(
            slide, 0.75, 5.65, 7.4, 1.1, inst, size=14, min_size=10, color=light, name="institution"
        )
        # right: the five actions
        x0, w0 = 8.55, 4.1
        kit.text(
            slide,
            x0,
            1.2,
            w0,
            0.4,
            P("Every message ends in exactly one action", bold=True),
            size=14,
            min_size=11,
            color=soft,
            name="actions-label",
        )
        y = 1.72
        for action in spec.get("actions") or []:
            name = str(action.get("name", "")).upper()
            tone = ACTION_TONES.get(name, TONES["det"])
            kit.chip(slide, x0, y, 1.75, 0.56, name, tone, size=14, name=f"chip-{name}")
            kit.text(
                slide,
                x0 + 1.9,
                y,
                w0 - 1.9,
                0.56,
                self.r(action.get("text", "")),
                size=13,
                min_size=10,
                color=light,
                anchor="m",
                name=f"action-{name}",
            )
            y += 0.74
        kit.text(
            slide,
            x0,
            y + 0.1,
            w0,
            0.6,
            P("Two of the five never call the text generator.", italic=True),
            size=12,
            min_size=10,
            color=soft,
            name="actions-foot",
        )
        self.set_notes(slide, spec.get("notes", ""))

    def s_problem(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        stats = spec.get("stats") or []
        cw, ch, gap = 2.92, 2.42, 0.18
        for i, st in enumerate(stats[:4]):
            cx = MX + (i % 2) * (cw + gap)
            cy = BODY_TOP + (i // 2) * (ch + gap)
            tone = TONES.get(str(st.get("tone", "gate")), TONES["gate"])
            kit.stat(
                slide,
                cx,
                cy,
                cw,
                ch,
                self.r(st.get("value")),
                self.r(st.get("label")),
                tone,
                source=self.r(st.get("cite", "")),
                value_size=40,
                label_size=16,
                name=f"stat-{i}",
            )
        rx = MX + 2 * cw + gap + 0.45
        rw = MX + CW - rx
        y = self.heading(slide, rx, BODY_TOP, rw, spec.get("risks_title", ""), size=19)
        kit.text(
            slide,
            rx,
            y + 0.05,
            rw,
            BODY_BOTTOM - y - 0.05,
            self.bullets(spec.get("risks") or [], color=TONES["emergency"].line, gap=8),
            size=16,
            min_size=12,
            name="risks",
            claim=True,
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_gap(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        lw = 6.15
        y = self.heading(slide, MX, BODY_TOP, lw, spec.get("gap_title", ""), size=18)
        kit.text(
            slide,
            MX,
            y,
            lw,
            4.35 - y,
            self.bullets(spec.get("gap") or [], gap=6),
            size=15,
            min_size=11.5,
            name="gap",
            claim=True,
        )
        reasons = spec.get("reasons") or []
        rw = (lw - 0.2) / 2
        for i, rs in enumerate(reasons[:2]):
            kit.card(
                slide,
                MX + i * (rw + 0.2),
                4.5,
                rw,
                BODY_BOTTOM - 4.5,
                self.r(rs.get("title")),
                self.r(rs.get("text")),
                TONES["gate"],
                title_size=15,
                body_size=14,
                name=f"reason-{i}",
                title_h=0.62,
            )
        px = MX + lw + 0.35
        pw = MX + CW - px
        kit.rect(slide, px, BODY_TOP, pw, BODY_BOTTOM - BODY_TOP, fill=K.PANEL, line=K.RULE)
        self.kit.claim("objectives", px, BODY_TOP, pw, BODY_BOTTOM - BODY_TOP)
        y = self.heading(
            slide, px + 0.2, BODY_TOP + 0.12, pw - 0.4, spec.get("objectives_title", ""), size=18
        )
        paras = [
            P(f"**{i + 1}.**  {self.r(t)}", space_before=7 if i else 0)
            for i, t in enumerate(spec.get("objectives") or [])
        ]
        kit.text(
            slide,
            px + 0.2,
            y,
            pw - 0.4,
            BODY_BOTTOM - y - 0.1,
            paras,
            size=14.5,
            min_size=11,
            name="objectives-list",
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_architecture(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        fig = str(spec.get("figure", ""))
        fw = 3.75
        ok = kit.image(
            slide,
            self.root / fig,
            MX,
            BODY_TOP - 0.05,
            fw,
            BODY_BOTTOM - BODY_TOP + 0.4,
            placeholder_title=fig,
            placeholder_detail="python report/diagrams/build_all.py",
            name="architecture-figure",
        )
        self.cur.images[fig] = ok
        if ok:
            self.cur.source(fig)
        x = MX + fw + 0.3
        w = MX + CW - x
        stages = spec.get("stages") or []
        n = max(len(stages), 1)
        gap = 0.06
        rh = (BODY_BOTTOM + 0.35 - BODY_TOP - gap * (n - 1)) / n
        for i, st in enumerate(stages):
            tone = TONES.get(str(st.get("tone", "det")), TONES["det"])
            y = BODY_TOP + i * (rh + gap)
            kit.rect(slide, x, y, w, rh, fill=tone.fill, line=tone.line, line_w=0.9, radius=0.06)
            kit.chip(
                slide,
                x + 0.08,
                y + (rh - 0.36) / 2,
                0.42,
                0.36,
                str(st.get("n", "")),
                Tone(tone.line, tone.line, K.WHITE),
                size=13,
                name=f"stage-{i}-n",
            )
            kit.text(
                slide,
                x + 0.6,
                y,
                2.45,
                rh,
                P(self.r(st.get("name")), bold=True),
                size=14,
                min_size=11,
                color=tone.text,
                anchor="m",
                name=f"stage-{i}-name",
            )
            kit.text(
                slide,
                x + 3.1,
                y,
                w - 3.2,
                rh,
                self.r(st.get("text")),
                size=13,
                min_size=10.5,
                color=K.INK,
                anchor="m",
                name=f"stage-{i}-text",
            )
        self.kit.claim("stages", x, BODY_TOP, w, n * rh + (n - 1) * gap)
        self.cur.source(self.get("deck.config.source"))
        self.set_notes(slide, spec.get("notes", ""))

    def s_gate(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        header = [self.r(c) for c in spec.get("columns") or []]
        rows = [[self.r(c) for c in row] for row in spec.get("rows") or []]

        def style(r: int, c: int, value: str) -> K.CellStyle | None:
            if c == 0:
                tone = ACTION_TONES.get(value.upper())
                if tone:
                    return K.CellStyle(fill=tone.fill, color=tone.text, bold=True)
            if c == len(header) - 1 and value.lower().startswith("no"):
                return K.CellStyle(color=TONES["safe"].text, bold=True)
            return None

        fit = kit.table(
            slide,
            MX,
            BODY_TOP,
            CW,
            header,
            rows,
            widths=[1.35, 1.75, 3.75, 2.85, 1.25],
            size=15,
            min_size=11,
            max_h=4.3,
            align=["c", "l", "l", "l", "c"],
            style=style,
            name="actions-table",
        )
        strip = [self.r(s) for s in spec.get("strip") or []]
        y = BODY_TOP + fit.height_in + 0.2
        h = min(BODY_BOTTOM + 0.3 - y, 1.15)
        n = max(len(strip), 1)
        w = (CW - 0.18 * (n - 1)) / n
        for i, text in enumerate(strip):
            tone = TONES["gate"] if i < 2 else TONES["safe"]
            x = MX + i * (w + 0.18)
            kit.rect(slide, x, y, w, h, fill=tone.fill, line=tone.line)
            kit.text(
                slide,
                x + 0.1,
                y + 0.05,
                w - 0.2,
                h - 0.1,
                text,
                size=13.5,
                min_size=10.5,
                color=tone.text,
                anchor="m",
                name=f"strip-{i}",
            )
        self.kit.claim("strip", MX, y, CW, h)
        self.cur.source(
            "medbridge/ai/context_gate.py",
            "medbridge/ai/evidence_gate.py",
            self.get("deck.config.source"),
        )
        self.set_notes(slide, spec.get("notes", ""))

    def s_emergency(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        lw = 5.95
        y = self.heading(
            slide,
            MX,
            BODY_TOP,
            lw,
            spec.get("how_title", ""),
            size=17,
            color=TONES["emergency"].text,
        )
        kit.text(
            slide,
            MX,
            y,
            lw,
            4.3 - y,
            self.bullets(spec.get("how") or [], color=TONES["emergency"].line, gap=5),
            size=14.5,
            min_size=11,
            name="how",
            claim=True,
        )
        kit.card(
            slide,
            MX,
            4.42,
            lw,
            BODY_BOTTOM - 4.42,
            self.r(spec.get("nl_title", "")),
            self.bullets(spec.get("nl") or [], color=TONES["llm"].line, gap=3),
            TONES["llm"],
            title_size=15,
            body_size=13.5,
            name="nl-card",
        )
        # right column
        rx = MX + lw + 0.35
        rw = MX + CW - rx
        y = self.heading(
            slide, rx, BODY_TOP, rw, "Examples, classified when this deck was built", size=15
        )
        examples = self.get("deck.emergency_examples") or []
        rows = []
        for ex in examples:
            actual = ex.get("actual")
            if ex.get("match") is False:
                self.example_mismatches.append(ex)
            label = {"ESCALATE": "ESCALATE", "PASS-THROUGH": "to the gates"}.get(
                str(actual), "[PENDING: classifier]"
            )
            rows.append([f"“{ex.get('text')}”", label])

        def ex_style(r: int, c: int, value: str) -> K.CellStyle | None:
            if c == 1 and value == "ESCALATE":
                t = ACTION_TONES["ESCALATE"]
                return K.CellStyle(fill=t.fill, color=t.text, bold=True)
            if c == 1:
                return K.CellStyle(color=K.MUTED, bold=True)
            return None

        if rows:
            fit = kit.table(
                slide,
                rx,
                y,
                rw,
                ["Patient message", "Fast-path output"],
                rows,
                widths=[3.95, 1.45],
                size=13,
                min_size=10.5,
                max_h=2.45,
                align=["l", "c"],
                style=ex_style,
                name="examples",
            )
            y += fit.height_in + 0.18
            self.cur.source("medbridge/core/emergency_classifier.py")
        else:
            kit.placeholder(
                slide,
                rx,
                y,
                rw,
                1.6,
                "emergency examples",
                "medbridge.core.emergency_classifier could not be imported",
                name="examples",
            )
            y += 1.78
        y = self.heading(slide, rx, y, rw, "Labelled case files (in-sample)", size=15)
        cases = self.get("deck.emergency") or {}
        crow = []
        for key, label in (("english", "English"), ("hinglish", "Hinglish / Indian English")):
            c = cases.get(key)
            if not isinstance(c, dict):
                continue
            self.cur.source(c.get("source"))
            crow.append(
                [
                    label,
                    f"{c['cases']} ({c['positives']} / {c['negatives']})",
                    str(c["tp"]),
                    str(c["fp"]),
                    str(c["fn"]),
                    str(c["tn"]),
                ]
            )
        if crow:
            fit = kit.table(
                slide,
                rx,
                y,
                rw,
                ["Case file", "Messages (esc. / not)", "TP", "FP", "FN", "TN"],
                crow,
                widths=[1.75, 1.55, 0.52, 0.52, 0.52, 0.52],
                size=12.5,
                min_size=10,
                max_h=1.15,
                name="case-files",
            )
            y += fit.height_in + 0.08
        else:
            kit.placeholder(slide, rx, y, rw, 0.9, "emergency case files", name="case-files")
            y += 1.0
        indep = self.r(
            "Independent estimate, HalluBench test emergencies (C5 recall): "
            "{{study.per_condition.C5.emergency.recall_system|ci}}"
        )
        kit.text(
            slide,
            rx,
            y,
            rw,
            BODY_BOTTOM + 0.02 - y,
            [
                P(
                    "Written alongside the rules: coverage of intended phrasings, not real-world "
                    "accuracy.",
                    italic=True,
                    color=K.MUTED,
                    scale=0.95,
                ),
                P(indep, space_before=3),
            ],
            size=12.5,
            min_size=10,
            name="case-note",
            claim=True,
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_retrieval(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        kpis = spec.get("kpis") or []
        n = max(len(kpis), 1)
        kw = (CW - 0.2 * (n - 1)) / n
        tones = [TONES["storage"], TONES["storage"], TONES["llm"], TONES["gate"]]
        for i, kp in enumerate(kpis):
            kit.stat(
                slide,
                MX + i * (kw + 0.2),
                BODY_TOP,
                kw,
                1.22,
                self.r(kp.get("value")),
                self.r(kp.get("label")),
                tones[i % len(tones)],
                value_size=30,
                label_size=13,
                name=f"kpi-{i}",
            )
        modes = self.get("retrieval.modes") or {}
        ablation = self.get("reranker") or {}
        rows: list[list[str]] = []
        cells: list[list[dict[str, Any] | None]] = []
        labels = {
            "dense": "Dense only (bge-small-en-v1.5)",
            "sparse": "Sparse only (BM25)",
            "hybrid": "Hybrid: dense + BM25, RRF k = 60",
            "hybrid+rerank": "Hybrid + rerank: MiniLM-L-6 (production)",
        }
        for key in ("dense", "sparse", "hybrid", "hybrid+rerank"):
            m = modes.get(key)
            if isinstance(m, dict):
                rows.append(
                    [labels[key], *[tables.interval(m.get(k)) for k, _ in tables.RETRIEVAL_METRICS]]
                )
                cells.append([m.get(k) for k, _ in tables.RETRIEVAL_METRICS])
        for dirname, data in sorted(ablation.items()):
            m = ((data or {}).get("modes") or {}).get("hybrid+rerank")
            name = dirname.replace("_", "/", 1)
            if not isinstance(m, dict) or name == tables.DEFAULT_RERANKER:
                continue
            short = re.sub(r"^ms-marco-|-v\d+$|-Cross-Encoder$", "", name.split("/")[-1])
            rows.append(
                [
                    f"Hybrid + rerank: {short}",
                    *[tables.interval(m.get(k)) for k, _ in tables.RETRIEVAL_METRICS],
                ]
            )
            cells.append([m.get(k) for k, _ in tables.RETRIEVAL_METRICS])
        y = BODY_TOP + 1.42
        if rows:

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                if rows[r][0].endswith("(production)"):
                    return K.CellStyle(fill=TONES["gate"].fill, bold=True, color=TONES["gate"].text)
                return None

            fit = kit.table(
                slide,
                MX,
                y,
                CW,
                [
                    f"Configuration ({self.get('retrieval.modes.hybrid.n', '?')} queries)",
                    *[m for _, m in tables.RETRIEVAL_METRICS],
                ],
                rows,
                widths=[3.6, 2.1, 2.1, 2.1, 2.1],
                size=14,
                min_size=11,
                max_h=3.55,
                align=["l", "c", "c", "c", "c"],
                style=style,
                name="retrieval-table",
            )
            y += fit.height_in + 0.12
            self.cur.source(self.ctx.source("retrieval"), self.ctx.source("reranker"))
            overlap = []
            for j in range(len(tables.RETRIEVAL_METRICS)):
                ivs: list[dict[str, Any]] = [c[j] for c in cells if isinstance(c[j], dict)]  # type: ignore[misc]
                lows = [float(v["low"]) for v in ivs if isinstance(v.get("low"), (int, float))]
                highs = [float(v["high"]) for v in ivs if isinstance(v.get("high"), (int, float))]
                if len(lows) > 1 and len(highs) > 1:
                    overlap.append(max(lows) <= min(highs))
            if overlap and all(overlap):
                n_q = self.get("retrieval.modes.hybrid.n")
                query_set = f"this {n_q}-query set" if isinstance(n_q, int) else "this query set"
                reading = (
                    "**Reading:** the 95 % bootstrap CIs of all configurations overlap on every "
                    f"metric, so {query_set} shows no clear difference between retrievers or "
                    "rerankers (no paired test was run); the pipeline keeps hybrid fusion and the "
                    "fast MiniLM-L-6 cross-encoder."
                )
            elif overlap:
                reading = (
                    "**Reading:** some 95 % CIs do not overlap; compare the rows above before "
                    "drawing conclusions."
                )
            else:
                reading = ""
            if reading:
                kit.text(
                    slide,
                    MX,
                    y,
                    CW,
                    BODY_BOTTOM + 0.02 - y,
                    reading,
                    size=14,
                    min_size=11,
                    color=K.NAVY,
                    name="reading",
                    claim=True,
                )
        else:
            kit.placeholder(
                slide,
                MX,
                y,
                CW,
                BODY_BOTTOM - y,
                "retrieval metrics",
                "python -m evaluation.retrieval_eval",
                name="retrieval-table",
            )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_privacy(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        cards = spec.get("cards") or []
        lw = 6.45
        n = max(len(cards), 1)
        gap = 0.1
        ch = (BODY_BOTTOM + 0.02 - BODY_TOP - gap * (n - 1)) / n
        for i, cd in enumerate(cards):
            tone = TONES.get(str(cd.get("tone", "det")), TONES["det"])
            y = BODY_TOP + i * (ch + gap)
            kit.rect(slide, MX, y, lw, ch, fill=tone.fill, line=tone.line)
            kit.rect(slide, MX, y + 0.1, 0.07, ch - 0.2, fill=tone.line, radius=None)
            kit.text(
                slide,
                MX + 0.18,
                y,
                1.55,
                ch,
                P(self.r(cd.get("title")), bold=True),
                size=15,
                min_size=11,
                color=tone.text,
                anchor="m",
                name=f"card-{i}-title",
            )
            kit.text(
                slide,
                MX + 1.78,
                y,
                lw - 1.88,
                ch,
                self.r(cd.get("text")),
                size=13,
                min_size=10,
                anchor="m",
                name=f"card-{i}-text",
            )
        kit.claim("cards", MX, BODY_TOP, lw, n * ch + (n - 1) * gap)
        rx = MX + lw + 0.35
        rw = MX + CW - rx
        y = self.heading(slide, rx, BODY_TOP, rw, "Redaction quality (strict span match)", size=17)
        rows = []
        for label, path in (
            ("Development set (in-sample)", "pii.datasets.development.results.all"),
            ("Blind held-out set (frozen detector)", "pii.datasets.blind_v2.results.all"),
        ):
            res = self.get(path)
            if isinstance(res, dict):
                s = res.get("overall_strict") or {}
                rows.append(
                    [
                        label,
                        tables.integer(res.get("n_examples")),
                        tables.num(s.get("precision"), 3),
                        tables.num(s.get("recall"), 3),
                        tables.num(s.get("f1"), 3),
                    ]
                )
        if rows:
            fit = kit.table(
                slide,
                rx,
                y,
                rw,
                ["Test set", "Messages", "P", "R", "F1"],
                rows,
                widths=[2.3, 1.05, 0.72, 0.72, 0.72],
                size=14,
                min_size=10.5,
                max_h=1.7,
                name="pii-table",
            )
            y += fit.height_in + 0.18
            self.cur.source(self.ctx.source("pii"))
        else:
            kit.placeholder(
                slide,
                rx,
                y,
                rw,
                1.4,
                "PII redaction metrics",
                "python scripts/pii_eval.py",
                name="pii-table",
            )
            y += 1.58
        facts = [
            "The blind set was written by another author and evaluated once, with the detector frozen",
            "Weak spot, names without a cue: blind-set recall "
            "{{pii.datasets.blind_v2.results.all.per_category_strict.NAME.recall|f3}}",
            "Clinical values corrupted: {{pii.datasets.development.results.all.protected_values_corrupted|len}}"
            " (development), {{pii.datasets.blind_v2.results.all.protected_values_corrupted|len}} (blind)",
            "PII-free blind messages altered: {{pii.datasets.blind_v2.results.all.negatives_altered|len}} of "
            "{{pii.datasets.blind_v2.results.all.n_negative_examples|int}}",
            "Mean detection time per message: "
            "{{pii.datasets.blind_v2.results.all.mean_latency_ms|f3}} ms",
        ]
        kit.text(
            slide,
            rx,
            y,
            rw,
            BODY_BOTTOM + 0.02 - y,
            self.bullets(facts, color=TONES["safe"].line, gap=8),
            size=15,
            min_size=10.5,
            name="pii-facts",
            claim=True,
        )
        self.sources_line(slide, ["medbridge/privacy/"])
        self.set_notes(slide, spec.get("notes", ""))

    def s_ui(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        shots = self.get("deck.screenshots") or {}
        scenes = spec.get("scenes") or []
        cols = int(spec.get("columns") or (4 if len(scenes) > 6 else 3))
        rows_n = max(1, math.ceil(len(scenes) / cols))
        gx, gy = (0.2, 0.16) if cols > 3 else (0.25, 0.18)
        cw = (CW - gx * (cols - 1)) / cols
        cap_h = 0.5 if cols > 3 else 0.52
        ch = (BODY_BOTTOM + 0.02 - BODY_TOP - gy * (rows_n - 1)) / rows_n
        missing = 0
        for i, sc in enumerate(scenes):
            name = str(sc.get("name"))
            x = MX + (i % cols) * (cw + gx)
            y = BODY_TOP + (i // cols) * (ch + gy)
            rel = shots.get(name)
            path = self.root / rel if rel else None
            ok = kit.image(
                slide,
                path,
                x,
                y,
                cw,
                ch - cap_h,
                placeholder_title=f"screenshot {name}.png",
                placeholder_detail="capture against the real backend",
                name=f"shot-{name}",
            )
            self.cur.images[rel or f"{deck_data.SCREENSHOT_DIR}/{name}.png"] = ok
            if ok:
                kit.rect(slide, x, y, cw, ch - cap_h, fill=None, line=K.RULE, radius=None)
            else:
                missing += 1
            kit.text(
                slide,
                x,
                y + ch - cap_h,
                cw,
                cap_h,
                P(self.r(sc.get("caption")), align="c"),
                size=13 if cols <= 3 else 12,
                min_size=10 if cols > 3 else 10.5,
                color=K.INK,
                anchor="m",
                name=f"caption-{name}",
            )
        used_dirs = sorted(
            {
                str(Path(shots[str(s.get("name"))]).parent)
                for s in scenes
                if shots.get(str(s.get("name")))
            }
        )
        self.cur.source(*used_dirs)
        if missing:
            self.cur.state = "pending" if missing == len(scenes) else "partial"
            self.cur.source(f"screenshots: capture with {spec.get('capture', SCREENSHOT_CAPTURE)}")
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_evaluation(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        lw = 5.55
        bench = self.get("hallubench")
        head = (
            "MedBridge-HalluBench v1: {{hallubench.n_items|int}} items, {{hallubench.n_turns|int}} turns"
            if isinstance(bench, dict)
            else "MedBridge-HalluBench v1 (synthetic)"
        )
        y = self.heading(slide, MX, BODY_TOP, lw, head, size=16)
        if isinstance(bench, dict):
            rows = []
            for cat in tables.CATEGORY_ORDER:
                c = (bench.get("by_category") or {}).get(cat)
                if isinstance(c, dict):
                    rows.append(
                        [
                            tables.CATEGORY_NAMES[cat],
                            str(c.get("items", 0)),
                            str(c.get("dev", 0)),
                            str(c.get("test", 0)),
                        ]
                    )
            split = bench.get("by_split") or {}
            rows.append(
                [
                    "Total",
                    str(bench.get("n_items")),
                    str(split.get("dev", 0)),
                    str(split.get("test", 0)),
                ]
            )

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                return (
                    K.CellStyle(bold=True, fill=TONES["det"].fill) if r == len(rows) - 1 else None
                )

            fit = kit.table(
                slide,
                MX,
                y,
                lw,
                ["Category", "Items", "Dev", "Test"],
                rows,
                widths=[2.85, 0.9, 0.9, 0.9],
                size=13,
                min_size=10.5,
                max_h=2.75,
                style=style,
                name="hallubench",
            )
            y += fit.height_in + 0.18
            self.cur.source(self.ctx.source("hallubench"))
        else:
            kit.placeholder(slide, MX, y, lw, 2.4, "HalluBench dataset", name="hallubench")
            y += 2.58
        kit.card(
            slide,
            MX,
            y,
            lw,
            BODY_BOTTOM + 0.02 - y,
            "Metrics and statistics",
            self.bullets(spec.get("stats") or [], color=TONES["det"].line, gap=6),
            TONES["note"],
            title_size=15,
            body_size=14.5,
            name="stats-card",
            title_h=0.38,
        )
        # right
        rx = MX + lw + 0.35
        rw = MX + CW - rx
        y = self.heading(
            slide, rx, BODY_TOP, rw, "Conditions (same model, corpus and benchmark)", size=16
        )
        conds = spec.get("conditions") or []
        colw = (rw - 0.15) / 2
        per_col = math.ceil(len(conds) / 2)
        rh = 0.46 if per_col <= 4 else 0.4
        for i, cd in enumerate(conds):
            tone = TONES.get(str(cd.get("tone", "det")), TONES["det"])
            cx = rx + (i // per_col) * (colw + 0.15)
            cy = y + (i % per_col) * (rh + 0.08)
            kit.rect(
                slide,
                cx,
                cy,
                colw,
                rh,
                fill=tone.fill,
                line=tone.line,
                dashed=str(cd.get("tone")) == "note",
            )
            kit.text(
                slide,
                cx + 0.06,
                cy,
                0.55,
                rh,
                P(str(cd.get("name")), bold=True),
                size=14,
                min_size=11,
                color=tone.text,
                anchor="m",
                name=f"cond-{i}-name",
            )
            kit.text(
                slide,
                cx + 0.62,
                cy,
                colw - 0.68,
                rh,
                self.r(cd.get("text")),
                size=12.5,
                min_size=10,
                anchor="m",
                name=f"cond-{i}-text",
            )
        grid_h = per_col * (rh + 0.08) - 0.08
        kit.claim("conditions", rx, y, rw, grid_h)
        y += grid_h + 0.2
        kit.card(
            slide,
            rx,
            y,
            rw,
            BODY_BOTTOM + 0.02 - y,
            "Two complementary judges, plus people",
            self.bullets(spec.get("judges") or [], color=TONES["llm"].line, gap=6),
            TONES["llm"],
            title_size=15,
            body_size=14.5,
            name="judges-card",
            title_h=0.38,
        )
        self.cur.source("medbridge/core/pipeline_types.py", "report/figures/evaluation_design.png")
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_models(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        table = tables.generate("model_selection", self.ctx)
        chosen = str(self.get("deck.config.LLM_MODEL") or "")
        y = BODY_TOP
        if table is not None:
            renames = {
                "Parameters, quantisation (size)": "Params, quant (size)",
                "Turn p50 / p95 s": "Turn time p50 / p95 (s)",
            }
            header = [renames.get(h, h) for h in table.header]
            rows = [list(r) for r in table.rows]

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                if rows[r][0] == chosen:
                    return K.CellStyle(fill=TONES["gate"].fill, bold=True, color=TONES["gate"].text)
                return None

            fit = kit.table(
                slide,
                MX,
                y,
                CW,
                header,
                rows,
                widths=table.col_widths,
                size=14,
                min_size=11,
                max_h=1.9,
                style=style,
                name="model-table",
            )
            y += fit.height_in + 0.18
            self.cur.source(self.ctx.source("model_selection"))
        else:
            kit.placeholder(
                slide,
                MX,
                y,
                CW,
                2.0,
                "model-selection summary",
                "python -m evaluation.model_selection",
                name="model-table",
            )
            y += 2.2
        banner = self._model_banner(chosen)
        if banner:
            kit.rect(slide, MX, y, CW, 0.8, fill=TONES["gate"].fill, line=TONES["gate"].line)
            kit.text(
                slide,
                MX + 0.18,
                y,
                CW - 0.36,
                0.8,
                banner,
                size=15,
                min_size=11,
                color=TONES["gate"].text,
                anchor="m",
                name="decision",
            )
            kit.claim("decision", MX, y, CW, 0.8)
            y += 0.95
        fig = str(spec.get("figure", ""))
        h = BODY_BOTTOM + 0.02 - y
        fw = min(7.6, h * 2.92)
        ok = kit.image(
            slide,
            self.root / fig if fig else None,
            MX,
            y,
            fw,
            h,
            placeholder_title=fig,
            placeholder_detail="python -m evaluation.model_selection",
            name="model-figure",
        )
        self.cur.images[fig] = ok
        if ok:
            self.cur.source(fig)
        rx = MX + fw + 0.3
        rw = MX + CW - rx
        kit.text(
            slide,
            rx,
            y,
            rw,
            h,
            self.bullets(spec.get("rationale") or [], color=TONES["gate"].line, gap=8),
            size=14.5,
            min_size=10.5,
            name="rationale",
            claim=True,
            anchor="m",
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def _model_banner(self, chosen: str) -> str:
        results = self.get("model_selection.results") or {}
        mine = results.get(chosen) if isinstance(results, dict) else None
        if not isinstance(mine, dict):
            return ""
        others = {
            m: s
            for m, s in results.items()
            if m != chosen and isinstance(s, dict) and s.get("n_valid")
        }

        def strict(s: dict[str, Any]) -> float:
            v = ((s.get("routing") or {}).get("accuracy_strict") or {}).get("rate")
            return float(v) if isinstance(v, (int, float)) else -1.0

        text = (
            f"**Application model: {chosen}**, routing {tables.pct(strict(mine), ci=False)} % strict "
            f"on {mine.get('n_valid')} dev turns, turn p50 {tables.seconds(mine.get('latency_turn_p50_ms'))} s"
        )
        size = (mine.get("ollama") or {}).get("size_bytes")
        if isinstance(size, (int, float)):
            text += f", {size / 1e9:.1f} GB"
        if others:
            best_m, best = max(others.items(), key=lambda kv: strict(kv[1]))
            if strict(best) > strict(mine):
                text += (
                    f". Best router: {best_m} at {tables.pct(strict(best), ci=False)} %, turn p50 "
                    f"{tables.seconds(best.get('latency_turn_p50_ms'))} s"
                )
                decision = self.get("deck.decision") or {}
                if decision.get("no_significant_difference"):
                    text += "; paired test: no significant difference (DECISION.md)"
                    self.cur.source(decision.get("source"))
        return NUMBER_UNIT_RE.sub("\\1\u00a0\\2", text + ".")

    # -- results -------------------------------------------------------

    def _study(self) -> dict[str, Any] | None:
        per = self.get("study.per_condition")
        return per if isinstance(per, dict) and per else None

    def _mcnemar(
        self, control: str, outcome: str, judge: str = "llm", treatment: str = "C5"
    ) -> dict[str, Any] | None:
        for t in self.get("study.tests.mcnemar") or []:
            if (
                isinstance(t, dict)
                and t.get("treatment") == treatment
                and t.get("control") == control
                and t.get("outcome") == outcome
                and t.get("judge") == judge
            ):
                return t
        return None

    @staticmethod
    def _significant(t: dict[str, Any] | None) -> bool | None:
        if not t:
            return None
        p = t.get("pvalue_holm", t.get("pvalue"))
        rd = (t.get("risk_difference") or {}).get("estimate")
        if not isinstance(p, (int, float)) or not isinstance(rd, (int, float)):
            return None
        return p < 0.05 and rd < 0

    @staticmethod
    def _verdict_label(t: dict[str, Any] | None) -> str:
        """Three-state verdict of one paired test (Holm p < 0.05 and the sign of the RD)."""
        return deck_data.verdict_label(t)

    def _rd_text(self, t: dict[str, Any]) -> str:
        rd = t.get("risk_difference") or {}
        vals = [rd.get("estimate"), rd.get("low"), rd.get("high")]
        if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in vals):
            return "—"
        est, lo, hi = (100 * float(v) for v in vals)  # type: ignore[arg-type]
        return f"{est:+.1f} pp [{lo:+.1f}, {hi:+.1f}]"

    def _figure(self, spec: dict[str, Any]) -> str:
        """First existing file of ``figures`` (preference order); else the last one listed."""
        figs = [str(f) for f in spec.get("figures") or []]
        if not figs and spec.get("figure"):
            figs = [str(spec["figure"])]
        for fig in figs:
            if (self.root / fig).is_file():
                return fig
        return figs[-1] if figs else ""

    def _pending_results(self, spec: dict[str, Any], slide: Any, what: str, detail: str) -> None:
        """Pending variant of a results slide: placeholder, what will appear, no number."""
        kit = self.kit
        self.cur.state = "pending"
        self.chrome(slide, spec, title=spec.get("title_pending", spec.get("title")))
        kit.image(
            slide,
            None,
            MX,
            BODY_TOP,
            6.6,
            BODY_BOTTOM - BODY_TOP,
            placeholder_title=what,
            placeholder_detail=detail,
            name="figure",
        )
        rx = MX + 6.95
        rw = MX + CW - rx
        y = self.heading(slide, rx, BODY_TOP, rw, "What this slide will show", size=17)
        kit.text(
            slide,
            rx,
            y,
            rw,
            BODY_BOTTOM - y - 1.0,
            self.bullets(spec.get("pending_points") or [], gap=7),
            size=15,
            min_size=11,
            name="pending-points",
            claim=True,
        )
        kit.chip(
            slide,
            rx,
            BODY_BOTTOM - 0.75,
            rw,
            0.62,
            "Status: no result file yet, so no number is quoted",
            TONES["pending"],
            size=14,
            name="status",
        )
        self.cur.pending.append("study (evaluation/results/summary.json)")
        self.sources_line(slide, ["evaluation/results/summary.json (pending)"])
        self.set_notes(slide, spec.get("notes_pending", ""))

    def s_results_hallucination(self, spec: dict[str, Any]) -> None:
        per = self._study()
        slide = self.new_slide(spec)
        kit = self.kit
        fig = self._figure(spec)
        if per is None:
            self._pending_results(
                spec,
                slide,
                "hallucination study results",
                f"{fig} and evaluation/results/summary.json are written by "
                "python -m evaluation.analyze",
            )
            self.cur.images[fig] = False
            return
        self.chrome(slide, spec)
        fw, fh = 6.3, 2.85
        ok = kit.image(
            slide,
            self.root / fig,
            MX,
            BODY_TOP,
            fw,
            fh,
            placeholder_title=fig,
            placeholder_detail="python -m evaluation.analyze",
            name="figure",
            valign="t",
        )
        self.cur.images[fig] = ok
        if ok:
            self.cur.source(fig)
        else:
            self.cur.state = "partial"
        y = BODY_TOP + fh + 0.1
        kit.text(
            slide,
            MX,
            y,
            fw,
            BODY_BOTTOM + 0.02 - y,
            self.bullets(spec.get("points") or [], color=TONES["gate"].line, gap=5),
            size=13.5,
            min_size=10,
            name="points",
            claim=True,
        )
        default = str(spec.get("default_condition", "V2"))
        rows = []
        for cond, m in tables._study_conditions(self.ctx):
            llm = m.get("llm") or {}
            label = cond + (" (uncal.)" if tables._uncalibrated(self.ctx, cond) else "")
            rows.append(
                [
                    label,
                    tables.pct(llm.get("hallucination")),
                    tables.pct(llm.get("unsafe_answer")),
                    tables.pct(m.get("coverage_answerable"), ci=False),
                ]
            )

        def style(r: int, c: int, value: str) -> K.CellStyle | None:
            cond = rows[r][0].split()[0]
            if cond == "C5":
                return K.CellStyle(fill=TONES["gate"].fill, bold=True, color=TONES["gate"].text)
            if cond == default:
                return K.CellStyle(fill=TONES["safe"].fill, bold=True, color=TONES["safe"].text)
            return None

        rx = MX + fw + 0.3
        rw = MX + CW - rx
        fit = kit.table(
            slide,
            rx,
            BODY_TOP,
            rw,
            ["Cond.", "Hallucinated, % [95% CI]", "Unsafe answers, % [95% CI]", "Coverage %"],
            rows,
            widths=[0.72, 1.8, 1.8, 0.95],
            size=12.5,
            min_size=9.5,
            max_h=4.25,
            style=style,
            name="halluc-table",
        )
        y2 = BODY_TOP + fit.height_in + 0.06
        kit.text(
            slide,
            rx,
            y2,
            rw,
            BODY_BOTTOM + 0.02 - y2,
            self.r(spec.get("table_note", "")),
            size=11,
            min_size=8.5,
            color=K.MUTED,
            name="table-note",
            claim=True,
        )
        self.cur.source(self.ctx.source("study"))
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", "") + "\n\n" + self._study_notes())

    def _study_notes(self) -> str:
        lines = [
            "Values from evaluation/results/summary.json (test split, LLM judge, degraded turns "
            "excluded):"
        ]
        nli_ok = set(self.get("deck.stats.nli_reported") or [])
        for cond, m in tables._study_conditions(self.ctx):
            llm, nli = m.get("llm") or {}, m.get("nli") or {}
            line = (
                f"- {cond} {tables.CONDITION_NAMES.get(cond, cond)}: hallucination "
                f"{tables.pct(llm.get('hallucination'))} %, unsafe answers "
                f"{tables.pct(llm.get('unsafe_answer'))} %, coverage "
                f"{tables.pct(m.get('coverage_answerable'), ci=False)} %, claims per response "
                f"{tables.num(llm.get('claims_per_response'), 2)}, turns {m.get('n_valid')}"
            )
            if cond in nli_ok and isinstance(nli.get("hallucination"), dict):
                line += f"; NLI judge {tables.pct(nli.get('hallucination'))} %"
            lines.append(line)
        if nli_ok:
            lines.append(
                "NLI-judge rates are quoted only for the conditions it scored completely ("
                + ", ".join(sorted(nli_ok))
                + "); it was concluded early, and a partial subset of any other condition is "
                "not a result and is not quoted."
            )
        kappa = self.get("study.judge_agreement.kappa_response_hallucination")
        if isinstance(kappa, dict):
            lines.append(
                f"Judge agreement (Cohen's kappa, response level, "
                f"{self.get('study.judge_agreement.n_responses', '?')} paired responses): "
                f"{tables.interval(kappa)}"
            )
        return "\n".join(lines)

    def s_results_unsafe(self, spec: dict[str, Any]) -> None:
        per = self._study()
        slide = self.new_slide(spec)
        kit = self.kit
        if per is None:
            self.cur.state = "pending"
            self.chrome(slide, spec, title=spec.get("title_pending", spec.get("title")))
            fh = 2.95
            kit.image(
                slide,
                None,
                MX,
                BODY_TOP,
                CW,
                fh,
                placeholder_title="paired McNemar tests: C5 vs C1 and C2, ablation, variants",
                placeholder_detail="evaluation/results/summary.json is written by "
                "python -m evaluation.analyze",
                name="tests",
            )
            y = BODY_TOP + fh + 0.3
            kit.card(
                slide,
                MX,
                y,
                CW,
                BODY_BOTTOM + 0.02 - y,
                "Decision rule (fixed before the results)",
                self.r(spec.get("decision_rule", "")),
                TONES["gate"],
                title_size=16,
                body_size=15,
                name="decision-rule",
                title_h=0.42,
            )
            self.cur.pending.append("study (evaluation/results/summary.json)")
            self.sources_line(slide, ["evaluation/results/summary.json (pending)"])
            self.set_notes(slide, spec.get("notes_pending", ""))
            return
        self.chrome(slide, spec)
        outcome_names = {"hallucination": "Hallucination", "unsafe": "Unsafe answer"}
        rows = []
        for comp in spec.get("comparisons") or []:
            treat = str(comp.get("treatment", "C5"))
            control = str(comp.get("control", ""))
            outcome = str(comp.get("outcome", "hallucination"))
            t = self._mcnemar(control, outcome, treatment=treat)
            if not t:
                continue
            rows.append(
                [
                    f"{treat} vs {control}",
                    self.r(comp.get("role", "")),
                    outcome_names.get(outcome, outcome),
                    f"{t.get('b_treatment_only')} / {t.get('c_control_only')}",
                    self._rd_text(t),
                    tables.pvalue(t.get("pvalue_holm")),
                    self._verdict_label(t),
                ]
            )
        y = BODY_TOP
        if rows:

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                if c == 6:
                    tone = {
                        "significant reduction": TONES["safe"],
                        "significant increase": TONES["emergency"],
                    }.get(value, TONES["det"])
                    return K.CellStyle(fill=tone.fill, color=tone.text, bold=True)
                return None

            fit = kit.table(
                slide,
                MX,
                y,
                CW,
                [
                    "Comparison",
                    "What it tests",
                    "Outcome",
                    "b / c",
                    "Risk difference [95% CI]",
                    "p (Holm)",
                    "Verdict",
                ],
                rows,
                widths=[1.3, 2.75, 1.45, 0.85, 2.45, 1.25, 2.1],
                size=13,
                min_size=10,
                max_h=3.35,
                align=["l", "l", "l", "c", "c", "c", "c"],
                style=style,
                name="mcnemar",
            )
            y += fit.height_in + 0.15
            self.cur.source(self.ctx.source("study"))
        # The rule: a significant reduction against BOTH C1 and C2 in at least one outcome.
        rule = self.get("deck.stats.rule")
        if isinstance(rule, dict) and rule.get("primary_total"):
            phrases = {
                "hallucination": "the hallucination rate",
                "unsafe": "the unsafe-answer rate",
            }
            met = [phrases.get(o, o) for o in rule.get("outcomes_met") or []]
            k, n = rule.get("primary_significant"), rule.get("primary_total")
            if rule.get("met"):
                verdict = (
                    f"**Decision rule:** met. C5 significantly reduces {' and '.join(met)} "
                    f"against both C1 and C2 ({k} of {n} primary comparisons, Holm-corrected, "
                    "LLM judge)."
                )
            else:
                verdict = (
                    f"**Decision rule:** not met. {k} of {n} primary comparisons show a "
                    "significant reduction; the error analysis and the V1 / V2 variants are in "
                    "the report."
                )
            body = [P(verdict), P(self.r(spec.get("decision_followup", "")), space_before=6)]
        else:
            body = [P("**Decision rule:** [PENDING: hallucination study]")]
        lw = 6.3
        h = BODY_BOTTOM + 0.02 - y
        kit.card(
            slide,
            MX,
            y,
            lw,
            h,
            self.r(spec.get("decision_title", "Decision rule fixed in advance")),
            body,
            TONES["gate"],
            title_size=15,
            body_size=13.5,
            min_size=10,
            name="decision",
            title_h=0.4,
        )
        kit.card(
            slide,
            MX + lw + 0.2,
            y,
            CW - lw - 0.2,
            h,
            self.r(spec.get("ablation_title", "Reading the ablation")),
            self.bullets(spec.get("ablation") or [], color=TONES["llm"].line, gap=4),
            TONES["llm"],
            title_size=15,
            body_size=13.5,
            min_size=10,
            name="ablation",
            title_h=0.4,
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", "") + "\n\n" + self._study_notes())

    @staticmethod
    def _ratio_pct(d: Any) -> str:
        """``{estimate, low, high}`` fraction as ``26.3 [22.2, 30.8]``."""
        if not isinstance(d, dict):
            return tables.pct(d)
        return tables.pct({"rate": d.get("estimate"), "low": d.get("low"), "high": d.get("high")})

    def s_results_tradeoff(self, spec: dict[str, Any]) -> None:
        per = self._study()
        slide = self.new_slide(spec)
        kit = self.kit
        if per is None:
            self._pending_results(
                spec,
                slide,
                "coverage and claim profile of C1, C5 and V2",
                "evaluation/results/summary.json is written by python -m evaluation.analyze",
            )
            return
        self.chrome(slide, spec)
        conds = [str(c) for c in spec.get("conditions") or ["C1", "C5", "V2"]]
        conds = [c for c in conds if isinstance(per.get(c), dict)]
        labels = spec.get("condition_labels") or {}
        default = str(spec.get("default_condition", "V2"))

        def llm(m: dict[str, Any]) -> dict[str, Any]:
            return m.get("llm") or {}

        def claim_bearing(m: dict[str, Any]) -> str:
            n = (llm(m).get("hallucination_given_claims") or {}).get("n")
            total = m.get("n_valid")
            return f"{n} of {total}" if n is not None and total else "—"

        metrics: list[tuple[str, Callable[[dict[str, Any]], str]]] = [
            (
                "Hallucinated responses, all turns, %",
                lambda m: tables.pct(llm(m).get("hallucination")),
            ),
            ("Replies that make clinical claims", claim_bearing),
            (
                "… of those, hallucinated, %",
                lambda m: tables.pct(llm(m).get("hallucination_given_claims")),
            ),
            ("Claims per response", lambda m: tables.num(llm(m).get("claims_per_response"), 1)),
            (
                "Unsupported claims, % of all claims",
                lambda m: self._ratio_pct(llm(m).get("unsupported_claim_ratio")),
            ),
            (
                "Coverage of answerable turns, %",
                lambda m: tables.pct(m.get("coverage_answerable"), ci=False),
            ),
            (
                "Over-abstention on answerable turns, %",
                lambda m: tables.pct(m.get("over_abstention"), ci=False),
            ),
            (
                "Unsafe answers, unanswerable turns, %",
                lambda m: tables.pct(llm(m).get("unsafe_answer")),
            ),
        ]
        rows = [[label] + [fn(per[c]) for c in conds] for label, fn in metrics]
        header = [self.r(spec.get("table_title", "Metric (LLM judge, test split)"))] + [
            self.r(labels.get(c, c)) for c in conds
        ]

        def style(r: int, c: int, value: str) -> K.CellStyle | None:
            if c == 0:
                return None
            cond = conds[c - 1]
            if cond == default:
                return K.CellStyle(fill=TONES["safe"].fill, color=TONES["safe"].text, bold=True)
            if cond == "C5":
                return K.CellStyle(fill=TONES["gate"].fill, color=TONES["gate"].text)
            return None

        lw = 7.15
        fit = kit.table(
            slide,
            MX,
            BODY_TOP,
            lw,
            header,
            rows,
            widths=[2.55] + [(lw - 2.55) / max(len(conds), 1)] * len(conds),
            size=13,
            min_size=10,
            max_h=4.35,
            align=["l"] + ["c"] * len(conds),
            style=style,
            name="tradeoff-table",
        )
        y = BODY_TOP + fit.height_in + 0.08
        kit.text(
            slide,
            MX,
            y,
            lw,
            BODY_BOTTOM + 0.02 - y,
            self.r(spec.get("table_note", "")),
            size=11.5,
            min_size=9,
            color=K.MUTED,
            name="tradeoff-note",
            claim=True,
        )
        self.cur.source(self.ctx.source("study"))
        rx = MX + lw + 0.3
        rw = MX + CW - rx
        cards = spec.get("cards") or []
        gap = 0.18
        heights = [float(c.get("height", 0)) for c in cards]
        free = BODY_BOTTOM + 0.02 - BODY_TOP - gap * max(len(cards) - 1, 0)
        flex = [i for i, h in enumerate(heights) if h <= 0]
        rest = free - sum(h for h in heights if h > 0)
        for i in flex:
            heights[i] = rest / len(flex)
        cy = BODY_TOP
        for i, card in enumerate(cards):
            tone = TONES.get(str(card.get("tone", "note")), TONES["note"])
            kit.card(
                slide,
                rx,
                cy,
                rw,
                heights[i],
                self.r(card.get("title", "")),
                self.bullets(card.get("bullets") or [], color=tone.line, gap=4),
                tone,
                title_size=15,
                body_size=13,
                min_size=9.5,
                name=f"card-{i}",
                title_h=0.4,
            )
            cy += heights[i] + gap
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", "") + "\n\n" + self._study_notes())

    # -- robustness, performance, testing -------------------------------

    def _panel(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        title: str,
        tone: Tone,
        *,
        name: str,
    ) -> float:
        """Tinted panel with an accent bar and a one-line title; returns the content top."""
        kit = self.kit
        kit.claim(name, x, y, w, h)
        kit.rect(slide, x, y, w, h, fill=tone.fill, line=tone.line, line_w=1.0)
        kit.rect(slide, x, y + 0.12, 0.07, h - 0.24, fill=tone.line, radius=None)
        kit.text(
            slide,
            x + 0.18,
            y + 0.08,
            w - 0.3,
            0.4,
            P(title, bold=True),
            size=15.5,
            min_size=12,
            color=tone.text,
            name=f"{name}-title",
        )
        return y + 0.5

    def _robust_table(
        self, key: str
    ) -> tuple[list[str], list[list[str]], Callable[[int, int, str], K.CellStyle | None]] | None:
        """Header, rows and cell style of one robustness panel (None when no live run)."""
        good, bad = TONES["safe"].text, TONES["emergency"].text

        def pc(v: Any) -> str:
            return "—" if v is None else f"{tables.perc(v)} %"

        if key == "aq":
            m = self.get("aq.metrics")
            if not isinstance(m, dict):
                return None
            self.cur.source(self.ctx.source("aq"))
            targets = self.get("deck.robust.aq_targets") or {}
            acc = m.get("acceptance") or {}

            def target(name: str, text: str) -> str:
                ok = acc.get(name)
                return text if ok is None else f"{text}: {'met' if ok else 'missed'}"

            def tgt(name: str, prefix: str = "") -> str:
                v = targets.get(name)
                return "—" if v is None else f"{prefix}{v:g} %"

            rows = [
                [
                    "Routing accuracy",
                    pc(m.get("overall_routing_accuracy")),
                    target("routing_accuracy", tgt("routing_accuracy", "≥ ")),
                ],
                [
                    "Citation precision",
                    pc(m.get("overall_citation_precision")),
                    target("citation_precision", tgt("citation_precision", "≥ ")),
                ],
                [
                    "Emergency recall",
                    f"{pc(m.get('emergency_detection_recall'))} "
                    f"({tables.integer(m.get('emergency_detected_count'))} of "
                    f"{tables.integer(m.get('emergency_count'))})",
                    target("emergency_recall", tgt("emergency_recall")),
                ],
                [
                    "Out-of-domain answered",
                    f"{tables.integer(m.get('out_of_domain_answer_count'))} of "
                    f"{tables.integer(m.get('out_of_domain_count'))}",
                    target("ood_answer_rate", tgt("ood_answer_rate")),
                ],
            ]

            def aq_style(r: int, c: int, value: str) -> K.CellStyle | None:
                if c == 2 and value.endswith("missed"):
                    return K.CellStyle(color=bad, bold=True)
                if c == 2 and value.endswith("met"):
                    return K.CellStyle(color=good, bold=True)
                return None

            return ["Metric", "Live", "Target"], rows, aq_style
        if key == "rabbits":
            after = self.get("rabbits")
            if not isinstance(after, dict):
                return None
            before = self.get("deck.robust.rabbits_before") or {}
            self.cur.source(before.get("source"), self.ctx.source("rabbits"))

            def indian(d: dict[str, Any]) -> Any:
                return ((d.get("by_origin") or {}).get("indian") or {}).get(
                    "action_consistency_pct"
                )

            def jac(v: Any) -> str:
                return tables.num(v, 3) if v is not None else "—"

            rows = [
                [
                    "Same action, brand vs generic",
                    pc(before.get("action_consistency_pct")),
                    pc(after.get("action_consistency_pct")),
                ],
                ["… Indian brands only", pc(indian(before)), pc(indian(after))],
                [
                    "Same generic stored (personal)",
                    pc(before.get("extraction_consistency_pct")),
                    pc(after.get("extraction_consistency_pct")),
                ],
                [
                    "Brand stored as its generic",
                    pc(before.get("brand_normalised_pct")),
                    pc(after.get("brand_normalised_pct")),
                ],
                [
                    "Citation-set Jaccard",
                    jac(before.get("citation_jaccard_mean")),
                    jac(after.get("citation_jaccard_mean")),
                ],
            ]

            def rb_style(r: int, c: int, value: str) -> K.CellStyle | None:
                return K.CellStyle(color=good, bold=True) if c == 2 else None

            return ["Metric", "Before fix", "After fix"], rows, rb_style
        if key == "adversarial":
            after = self.get("adversarial.counts")
            if not isinstance(after, dict):
                return None
            before = (self.get("deck.robust.adversarial_before") or {}).get("counts") or {}
            self.cur.source(
                (self.get("deck.robust.adversarial_before") or {}).get("source"),
                self.ctx.source("adversarial"),
            )

            def of(d: dict[str, Any], k: str, n: str) -> str:
                if d.get(k) is None:
                    return "—"
                return f"{tables.integer(d.get(k))} / {tables.integer(d.get(n))}"

            def cnt(d: dict[str, Any], k: str) -> str:
                return "—" if d.get(k) is None else tables.integer(d.get(k))

            rows = [
                [
                    "Passed all checks",
                    of(before, "passed", "payloads"),
                    of(after, "passed", "payloads"),
                ],
                [
                    "System prompt leaked",
                    cnt(before, "leaked_system_prompt"),
                    cnt(after, "leaked_system_prompt"),
                ],
                [
                    "Fabricated prescription",
                    cnt(before, "fabricated_prescription"),
                    cnt(after, "fabricated_prescription"),
                ],
                ["Unsafe output", cnt(before, "unsafe_output"), cnt(after, "unsafe_output")],
                [
                    "Emergencies escalated",
                    of(before, "emergency_escalated", "emergency_payloads"),
                    of(after, "emergency_escalated", "emergency_payloads"),
                ],
            ]
            return ["Check", "Before fix", "After fix"], rows, lambda r, c, v: None
        return None

    def s_robustness(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        cols = spec.get("columns") or []
        n = max(len(cols), 1)
        gap = 0.2
        w = (CW - gap * (n - 1)) / n
        top_h = 4.05
        pending = 0
        for i, col in enumerate(cols):
            key = str(col.get("key"))
            tone = TONES.get(str(col.get("tone", "det")), TONES["det"])
            x = MX + i * (w + gap)
            y0 = self._panel(
                slide, x, BODY_TOP, w, top_h, self.r(col.get("title")), tone, name=f"robust-{key}"
            )
            ix, iw = x + 0.18, w - 0.3
            fit = kit.text(
                slide,
                ix,
                y0,
                iw,
                1.05,
                self.r(col.get("scope")),
                size=12,
                min_size=9.5,
                color=K.INK,
                name=f"robust-{key}-scope",
            )
            ty = y0 + min(fit.needed_in, 1.05) + 0.06
            table = self._robust_table(key)
            if table is None:
                pending += 1
                kit.text(
                    slide,
                    ix,
                    ty,
                    iw,
                    BODY_TOP + top_h - ty - 0.1,
                    P(f"[PENDING: live run, {LIVE_COMMANDS.get(key, key)}]"),
                    size=12.5,
                    min_size=9.5,
                    name=f"robust-{key}-pending",
                )
                continue
            header, rows, style = table
            kit.table(
                slide,
                ix,
                ty,
                iw,
                header,
                rows,
                widths=[1.75, 1.0, 1.0] if key != "aq" else [1.45, 1.2, 1.15],
                size=12,
                min_size=9,
                max_h=BODY_TOP + top_h - ty - 0.12,
                style=style,
                name=f"robust-{key}-table",
                claim=False,
            )
        y = BODY_TOP + top_h + 0.18
        h = BODY_BOTTOM + 0.02 - y
        kit.rect(slide, MX, y, CW, h, fill=TONES["safe"].fill, line=TONES["safe"].line)
        kit.text(
            slide,
            MX + 0.15,
            y,
            CW - 0.3,
            h,
            self.r(spec.get("hermetic", "")),
            size=13.5,
            min_size=10,
            color=TONES["safe"].text,
            anchor="m",
            name="hermetic",
        )
        kit.claim("hermetic", MX, y, CW, h)
        skipped = self.get("_skipped") or {}
        if pending:
            self.cur.state = "pending" if pending == len(cols) else "partial"
        if skipped:
            self.cur.source("mock-mode reports ignored: " + ", ".join(sorted(skipped)))
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    @staticmethod
    def _sec(v: Any) -> str:
        if not isinstance(v, (int, float)) or not math.isfinite(v):
            return "—"
        return f"{v:.2f}" if abs(v) < 1 else f"{v:.1f}"

    def _arrow(self, before: Any, after: Any) -> str:
        return f"{self._sec(before)} → {self._sec(after)}"

    @staticmethod
    def _change(row: dict[str, Any]) -> str:
        b, a = row.get("before_p50"), row.get("after_p50")
        if not isinstance(b, (int, float)) or not isinstance(a, (int, float)) or a <= 0 or b <= 0:
            return "—"
        d = a - b
        if abs(d) < 0.1:
            return "≈ 0 s"
        ratio = b / a if d < 0 else a / b
        text = f"{'−' if d < 0 else '+'}{abs(d):.1f} s"
        if ratio >= 1.15:
            text += f" ({ratio:.1f}× {'faster' if d < 0 else 'slower'})"
        return text

    def s_performance(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        perf = self.get("deck.perf")
        turns = (perf or {}).get("turns") or {}
        lw = 7.45
        y = self.heading(slide, MX, BODY_TOP, lw, spec.get("table_title", ""), size=15)
        if turns:
            has_after = any("after_p50" in r for r in turns.values())
            has_dep = any("deployed_p50" in r for r in turns.values())
            header = [
                "Scenario turn",
                "Action",
                "Cold: before → after",
                "Warm p50: before → after",
                "Warm change",
            ]
            if has_dep:
                header.append(f"Final on {perf.get('deployed_model')}, warm p50")
            names = spec.get("scenario_names") or {}
            rows = []
            for turn in sorted(turns):
                r = turns[turn]
                key = turn.split("-")[0]
                scen = str(names.get(key) or tables.SCENARIO_NAMES.get(key, ""))
                act = r.get("before_actions") or r.get("after_actions") or "—"
                if r.get("after_actions") and r.get("before_actions") and r["after_actions"] != act:
                    act += " → " + r["after_actions"]
                row = [
                    f"{turn} {scen}".strip(),
                    act,
                    self._arrow(r.get("before_cold"), r.get("after_cold")),
                    self._arrow(r.get("before_p50"), r.get("after_p50")),
                    self._change(r),
                ]
                if has_dep:
                    row.append(self._sec(r.get("deployed_p50")))
                rows.append(row)

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                if c == 4 and "faster" in value:
                    return K.CellStyle(color=TONES["safe"].text, bold=True)
                if c == 4 and value.startswith("+"):
                    return K.CellStyle(color=TONES["emergency"].text, bold=True)
                return None

            widths = [1.6, 1.05, 1.15, 1.2, 1.5] + ([0.95] if has_dep else [])
            fit = kit.table(
                slide,
                MX,
                y,
                lw,
                header,
                rows,
                widths=widths,
                size=12,
                min_size=9.5,
                max_h=2.95,
                style=style,
                name="perf-table",
            )
            y += fit.height_in + 0.1
            self.cur.source((perf or {}).get("source"))
            if not has_after:
                self.cur.state = "partial"
                self.cur.pending.append("deck.perf (after-optimisation probe)")
                paras = [
                    P(
                        "[PENDING: after-optimisation probe, python scripts/latency_probe.py "
                        "--label after_<model> --repeats 3]"
                    )
                ]
            else:
                paras = self.bullets(spec.get("honest") or [], color=TONES["det"].line, gap=4)
            kit.text(
                slide,
                MX,
                y,
                lw,
                BODY_BOTTOM + 0.02 - y,
                paras,
                size=12.5,
                min_size=9.5,
                name="perf-note",
                claim=True,
            )
        else:
            kit.placeholder(
                slide,
                MX,
                y,
                lw,
                BODY_BOTTOM - y,
                "latency probe files",
                "python scripts/latency_probe.py",
                name="perf-table",
            )
        rx = MX + lw + 0.3
        rw = MX + CW - rx
        top = BODY_TOP
        h1 = 2.55
        kit.rect(slide, rx, top, rw, h1, fill=K.PANEL, line=K.RULE)
        kit.claim("optimisations", rx, top, rw, h1)
        y2 = self.heading(
            slide, rx + 0.16, top + 0.08, rw - 0.32, spec.get("changes_title", ""), size=14.5
        )
        kit.text(
            slide,
            rx + 0.16,
            y2,
            rw - 0.32,
            top + h1 - y2 - 0.05,
            self.bullets(spec.get("optimisations") or [], color=TONES["llm"].line, gap=3),
            size=12,
            min_size=9.5,
            name="optimisations-list",
        )
        y3 = top + h1 + 0.15
        h2 = BODY_BOTTOM + 0.02 - y3
        tone = TONES["llm"]
        y4 = self._panel(
            slide,
            rx,
            y3,
            rw,
            h2,
            self.r(spec.get("load_title", "Burst load test")),
            tone,
            name="load",
        )
        levels = self.get("load.levels")
        if isinstance(levels, list) and levels:
            self.cur.source(self.ctx.source("load"))
            requests = [r for r in self.get("load.requests") or [] if isinstance(r, dict)]
            rows = []
            for lv in levels:
                if not isinstance(lv, dict):
                    continue
                llm = sorted(
                    float(r["latency_ms"])
                    for r in requests
                    if r.get("level") == lv.get("users")
                    and r.get("kind") not in FAST_KINDS
                    and isinstance(r.get("latency_ms"), (int, float))
                )
                emerg = ((lv.get("by_kind") or {}).get("emergency") or {}).get("p50_ms")
                rows.append(
                    [
                        tables.integer(lv.get("users")),
                        f"{tables.integer(lv.get('errors'))} / {tables.integer(lv.get('requests'))}",
                        tables.seconds(_median(llm)) if llm else "—",
                        tables.seconds(llm[-1]) if llm else "—",
                        "—" if emerg is None else f"{float(emerg):.0f} ms",
                    ]
                )
            fit = kit.table(
                slide,
                rx + 0.18,
                y4,
                rw - 0.3,
                ["Users", "Errors", "LLM p50 s", "LLM max s", "Emergency p50"],
                rows,
                widths=[0.7, 0.85, 0.8, 0.8, 1.05],
                size=11.5,
                min_size=9,
                max_h=y3 + h2 - y4 - 0.55,
                name="load-table",
                claim=False,
            )
            ny = y4 + fit.height_in + 0.05
            kit.text(
                slide,
                rx + 0.18,
                ny,
                rw - 0.3,
                y3 + h2 - ny - 0.06,
                self.r(spec.get("load_note", "")),
                size=11,
                min_size=8.5,
                color=K.MUTED,
                name="load-note",
            )
        else:
            kit.text(
                slide,
                rx + 0.18,
                y4,
                rw - 0.3,
                y3 + h2 - y4 - 0.1,
                P("[PENDING: load test, python scripts/load_test.py]"),
                size=12,
                min_size=9.5,
                name="load-pending",
            )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_testing(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        tiles = []
        totals = self.get("tests.totals") or self.get("tests.total")
        if isinstance(totals, dict) and totals.get("passed") is not None:
            tiles.append(
                (
                    f"{tables.integer(totals.get('passed'))} passed",
                    f"current run: {tables.integer(totals.get('failed', totals.get('failures', 0)))} failed, "
                    f"{tables.integer(totals.get('skipped', 0))} skipped",
                    TONES["safe"],
                    self.ctx.source("tests"),
                )
            )
        else:
            tiles.append(
                (
                    "Pending",
                    "current full test run: [PENDING: full test-run summary]",
                    TONES["pending"],
                    None,
                )
            )
            self.cur.pending.append("tests (reports/test_summary.json)")
        base = self.get("deck.baseline")
        if isinstance(base, dict):
            tiles.append(
                (
                    f"{base.get('passed', '?')} passed",
                    f"received prototype: {base.get('failed', 0)} failed in its own baseline run",
                    TONES["det"],
                    base.get("source"),
                )
            )
        smoke = self.get("smoke")
        if isinstance(smoke, dict) and isinstance(smoke.get("counts"), dict):
            c = smoke["counts"]
            n = sum(int(v) for v in c.values() if isinstance(v, int))
            tiles.append(
                (
                    f"{c.get('PASS', 0)} / {n} PASS",
                    f"end-to-end smoke test on a running server "
                    f"({str(smoke.get('generated_at', ''))[:10]})",
                    TONES["safe"],
                    self.ctx.source("smoke"),
                )
            )
        cr = self.get("deck.code_review")
        if isinstance(cr, dict):
            tiles.append(
                (
                    f"{cr['confirmed']} fixed",
                    f"review findings ({cr['high']} high, {cr['medium']} medium, "
                    f"{cr['low']} low) in commit {cr.get('commit')}; {cr['rejected']} rejected",
                    TONES["gate"],
                    cr.get("source"),
                )
            )
        n = max(len(tiles), 1)
        tw = (CW - 0.2 * (n - 1)) / n
        for i, (value, label, tone, src) in enumerate(tiles):
            kit.stat(
                slide,
                MX + i * (tw + 0.2),
                BODY_TOP,
                tw,
                1.42,
                value,
                label,
                tone,
                value_size=24,
                label_size=12.5,
                name=f"tile-{i}",
            )
            self.cur.source(src)
        y = BODY_TOP + 1.62
        lw = 6.7
        inv = self.get("deck.tests")
        if isinstance(inv, dict):
            rows = [[s["name"], str(s["files"]), str(s["tests_defined"])] for s in inv["suites"]]
            rows.append(
                [
                    "Total (Python + JavaScript)",
                    str(sum(s["files"] for s in inv["suites"])),
                    str(inv["python_tests_defined"] + inv["js_tests_defined"]),
                ]
            )

            def style(r: int, c: int, value: str) -> K.CellStyle | None:
                return (
                    K.CellStyle(bold=True, fill=TONES["det"].fill) if r == len(rows) - 1 else None
                )

            kit.table(
                slide,
                MX,
                y,
                lw,
                ["Test suite (static inventory)", "Files", "Test functions"],
                rows,
                widths=[4.3, 0.9, 1.5],
                size=12.5,
                min_size=10,
                max_h=BODY_BOTTOM - y - 0.05,
                style=style,
                name="inventory",
            )
            self.cur.source("tests/, medbridge/tests/, medbridge/web/tests/ (static count)")
        rx = MX + lw + 0.35
        rw = MX + CW - rx
        kit.rect(slide, rx, y, rw, BODY_BOTTOM + 0.02 - y, fill=K.PANEL, line=K.RULE)
        kit.claim("quality", rx, y, rw, BODY_BOTTOM + 0.02 - y)
        y2 = self.heading(slide, rx + 0.18, y + 0.1, rw - 0.36, "Quality practices", size=16)
        kit.text(
            slide,
            rx + 0.18,
            y2,
            rw - 0.36,
            BODY_BOTTOM - y2 - 0.05,
            self.bullets(spec.get("quality") or [], color=TONES["safe"].line, gap=8),
            size=13.5,
            min_size=10.5,
            name="quality-list",
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_limitations(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        kit = self.kit
        lw = 6.75
        y = self.heading(
            slide,
            MX,
            BODY_TOP,
            lw,
            "Limitations (stated honestly)",
            size=18,
            color=TONES["emergency"].text,
        )
        kit.text(
            slide,
            MX,
            y,
            lw,
            BODY_BOTTOM - y,
            self.bullets(spec.get("limitations") or [], color=TONES["emergency"].line, gap=10),
            size=16,
            min_size=10.5,
            name="limitations",
            claim=True,
        )
        rx = MX + lw + 0.35
        rw = MX + CW - rx
        kit.rect(
            slide,
            rx,
            BODY_TOP,
            rw,
            BODY_BOTTOM + 0.02 - BODY_TOP,
            fill=TONES["safe"].fill,
            line=TONES["safe"].line,
        )
        kit.claim("future", rx, BODY_TOP, rw, BODY_BOTTOM + 0.02 - BODY_TOP)
        y2 = self.heading(
            slide,
            rx + 0.18,
            BODY_TOP + 0.1,
            rw - 0.36,
            "Future work",
            size=18,
            color=TONES["safe"].text,
        )
        kit.text(
            slide,
            rx + 0.18,
            y2,
            rw - 0.36,
            BODY_BOTTOM - y2 - 0.05,
            self.bullets(spec.get("future") or [], color=TONES["safe"].line, gap=8),
            size=15,
            min_size=10.5,
            name="future-list",
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    def s_thanks(self, spec: dict[str, Any]) -> None:
        slide = self.new_slide(spec, dark=True)
        kit = self.kit
        self.cur.title = self.r(spec.get("title"))
        gold, light, soft = "F2C14E", "E8EDF6", "B9C4DA"
        kit.rect(slide, 0, 0, K.SLIDE_W, 0.09, fill=K.AMBER, radius=None)
        kit.text(
            slide,
            0.9,
            1.6,
            11.5,
            1.2,
            P(self.cur.title, bold=True),
            size=50,
            min_size=32,
            color=K.WHITE,
            name="title",
        )
        team = self.spec("title") or {}
        names = " · ".join(self.r(t) for t in team.get("team") or [])
        guide = self.r(team.get("guide", ""))
        kit.text(
            slide,
            0.9,
            2.95,
            11.5,
            0.9,
            [P(names), P(f"Guide: {guide}", color=soft, space_before=4)],
            size=18,
            min_size=13,
            color=light,
            name="team",
        )
        kit.rect(slide, 0.9, 4.05, 1.1, 0.05, fill=gold, radius=None)
        kit.text(
            slide,
            0.9,
            4.25,
            11.5,
            1.1,
            [self.r(t) for t in spec.get("lines") or []],
            size=15,
            min_size=11,
            color=light,
            name="lines",
        )
        kit.text(
            slide,
            0.9,
            6.2,
            11.5,
            0.6,
            P(self.r(spec.get("disclaimer", "")), italic=True),
            size=13,
            min_size=10,
            color=soft,
            name="disclaimer",
        )
        self.set_notes(slide, spec.get("notes", ""))

    def s_generic(self, spec: dict[str, Any]) -> None:
        """Unknown slide id: title plus ``bullets`` (lets the team add a simple slide)."""
        slide = self.new_slide(spec)
        self.chrome(slide, spec)
        self.kit.text(
            slide,
            MX,
            BODY_TOP,
            CW,
            BODY_BOTTOM - BODY_TOP,
            self.bullets(spec.get("bullets") or [], gap=8),
            size=18,
            min_size=12,
            name="bullets",
            claim=True,
        )
        self.sources_line(slide)
        self.set_notes(slide, spec.get("notes", ""))

    # -- log -------------------------------------------------------------

    def log(self) -> dict[str, Any]:
        return {
            "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "slides": [r.as_dict() for r in self.records],
            "slide_count": len(self.records),
            "layout_issues": [i.as_dict() for i in self.kit.issues],
            "citation_keys_missing": sorted({k for r in self.records for k in r.citations_missing}),
            "pending_slides": [r.id for r in self.records if r.state != "data"],
            "emergency_example_mismatches": self.example_mismatches,
            "results_sources": self.ctx.data.get("_sources", {}),
            "results_skipped": self.ctx.data.get("_skipped", {}),
            "deck_fact_errors": (self.ctx.data.get("deck") or {}).get("_errors", {}),
            "text_measurement": self.kit.measurer.source,
        }


# ----------------------------------------------------------------------
# Content, PDF and PNG
# ----------------------------------------------------------------------


def load_content(path: Path) -> dict[str, Any]:
    with path.open("rb") as fh:
        data = tomllib.load(fh)
    if not isinstance(data.get("slides"), list) or not data["slides"]:
        raise ValueError(f"{path}: no [[slides]] tables")
    return data


def convert_to_pdf(pptx: Path, pdf: Path, timeout: int = 300) -> tuple[Path | None, str]:
    """Headless LibreOffice conversion with a throw-away profile; never raises."""
    soffice = find_soffice()
    if soffice is None:
        return None, "LibreOffice (soffice) not found"
    with tempfile.TemporaryDirectory(prefix="medbridge_deck_") as tmp:
        work = Path(tmp) / "deck.pptx"
        shutil.copyfile(pptx, work)
        outdir = Path(tmp) / "out"
        outdir.mkdir()
        cmd = [
            soffice,
            f"-env:UserInstallation={(Path(tmp) / 'profile').resolve().as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(work),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, check=False)
        except (subprocess.TimeoutExpired, OSError) as exc:
            return None, str(exc)
        produced = outdir / "deck.pdf"
        if not produced.exists():
            return None, (proc.stderr or proc.stdout or "conversion failed").strip()[-400:]
        pdf.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), pdf)
    return pdf, "convert-to"


def render_pngs(pdf: Path, out_dir: Path, dpi: int = 80) -> list[Path]:
    try:
        import pymupdf
    except ImportError:  # older PyMuPDF releases only ship the "fitz" name
        import fitz as pymupdf  # type: ignore[no-redef]
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    with pymupdf.open(str(pdf)) as doc:
        for i, page in enumerate(doc, start=1):
            path = out_dir / f"slide-{i:02d}.png"
            page.get_pixmap(dpi=dpi).save(str(path))
            paths.append(path)
    return paths


def build(
    out: Path = DEFAULT_OUT,
    content_path: Path = DEFAULT_CONTENT,
    repo_root: Path = REPO_ROOT,
    *,
    refs_root: Path | None = None,
) -> tuple[Path, dict[str, Any]]:
    builder = DeckBuilder(load_content(content_path), repo_root, refs_root=refs_root)
    prs = builder.build()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    log = builder.log()
    log["output"] = _display_path(out, repo_root)
    return out, log


def _display_path(path: Path, root: Path) -> str:
    """Repo-relative path for the log (never a machine-specific absolute path)."""
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except ValueError:
        return path.name


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[1] if __doc__ else None)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output .pptx")
    ap.add_argument("--content", type=Path, default=DEFAULT_CONTENT, help="slide content (TOML)")
    ap.add_argument("--repo-root", type=Path, default=REPO_ROOT, help=argparse.SUPPRESS)
    ap.add_argument("--log", type=Path, default=DEFAULT_LOG, help="JSON build log")
    ap.add_argument("--no-log", action="store_true", help="do not write the build log")
    ap.add_argument(
        "--pdf", action="store_true", help="also write a PDF next to the PPTX (LibreOffice)"
    )
    ap.add_argument(
        "--png",
        type=Path,
        default=None,
        metavar="DIR",
        help="render every slide to DIR/slide-NN.png (LibreOffice + PyMuPDF)",
    )
    ap.add_argument(
        "--strict",
        action="store_true",
        help="exit 1 on layout issues, unknown citation keys or example mismatches",
    )
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args(argv)

    if not args.content.is_file():
        print(f"content file not found: {args.content}", file=sys.stderr)
        return 2
    try:
        out, log = build(args.out, args.content, args.repo_root)
    except (ValueError, tomllib.TOMLDecodeError) as exc:
        print(f"invalid content: {exc}", file=sys.stderr)
        return 2

    code = 0
    if args.pdf or args.png:
        pdf_path = out.with_suffix(".pdf") if args.pdf else Path(tempfile.mkdtemp()) / "deck.pdf"
        pdf, detail = convert_to_pdf(out, pdf_path)
        log["pdf"] = {
            "path": pdf.name if pdf else None,
            "detail": detail,
            "pages": pdf_page_count(pdf) if pdf else None,
        }
        if pdf is None:
            print(f"PDF conversion failed: {detail}", file=sys.stderr)
            code = 3
        elif args.png:
            pngs = render_pngs(pdf, args.png)
            log["png"] = {"dir": _display_path(args.png, args.repo_root), "count": len(pngs)}
            if not args.pdf:
                shutil.rmtree(pdf.parent, ignore_errors=True)
    if not args.no_log:
        args.log.parent.mkdir(parents=True, exist_ok=True)
        args.log.write_text(json.dumps(log, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    issues = log["layout_issues"]
    missing = log["citation_keys_missing"]
    mismatches = log["emergency_example_mismatches"]
    if not args.quiet:
        print(
            f"wrote {out} ({log['slide_count']} slides, text measured with {log['text_measurement']})"
        )
        if log["pending_slides"]:
            print("  pending / partial slides: " + ", ".join(log["pending_slides"]))
        pend = sorted({p for s in log["slides"] for p in s["pending"]})
        if pend:
            print("  pending values: " + ", ".join(pend))
        for issue in issues:
            print(
                f"  LAYOUT {issue['slide']} {issue['element']}: {issue['kind']} ({issue['detail']})"
            )
        if missing:
            print("  unknown citation keys: " + ", ".join(missing))
        for ex in mismatches:
            print(
                f"  EXAMPLE '{ex['text']}': expected {ex['expected']}, classifier says {ex['actual']}"
            )
        if log.get("results_skipped"):
            for key, why in log["results_skipped"].items():
                print(f"  ignored {key}: {why}")
        if "pdf" in log:
            print(f"  pdf: {log['pdf']}")
    if args.strict and (issues or missing or mismatches) and code == 0:
        code = 1
    return code


if __name__ == "__main__":
    sys.exit(main())
