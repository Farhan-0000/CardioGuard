#!/usr/bin/env python3
"""
Build the MedBridge IEEE conference paper as DOCX (two-column IEEE layout).

The paper is assembled from ``report/paper/content/*.md`` (same markup as the report,
see ``report/README.md`` and ``report/paper/README.md``). Every number, generated table
and figure comes from the real results files through the report library
(``report/lib``); citations resolve only against the verified reference list.

Layout (IEEE conference template): US Letter (or A4 with ``--paper a4``), Times New
Roman 10 pt body, single spacing, justified, first-line indents; the title block is set
in one column and the body in two columns of 3.5 in (section XML ``w:cols w:num="2"``).
Section headings are numbered I, II, ... in small capitals, subsections A, B, ... in
italics. Figure captions read "Fig. 1." below the figure, table captions "TABLE I" above
the table. References are in IEEE style, numbered by first citation. A figure or table
that does not fit one column is set across both columns (continuous section breaks).
The last page's columns are balanced by a closing continuous section break.

Usage::

    python report/paper/build_paper.py                 # -> report/paper/MedBridge_IEEE_Paper.docx
    python report/paper/build_paper.py --paper a4      # A4 instead of US Letter
    python report/paper/build_paper.py --anonymous     # double-blind review copy
    python report/paper/build_paper.py --pdf           # + PDF preview via LibreOffice
    python report/paper/build_paper.py --strict        # exit 1 on unknown citations / xrefs

Exit codes: 0 success, 1 strict check failed, 2 content directory missing,
3 ``--pdf`` requested but no PDF could be produced.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

PAPER_DIR = Path(__file__).resolve().parent
if str(PAPER_DIR) not in sys.path:
    sys.path.insert(0, str(PAPER_DIR))

import paper_common as pc  # noqa: E402  (also puts report/ on sys.path)
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from lib import docx_helpers as dh  # noqa: E402
from lib.content import (  # noqa: E402
    Algorithm,
    Block,
    CodeBlock,
    Figure,
    Heading,
    ListBlock,
    PageBreak,
    Paragraph,
    Pending,
    PendingTable,
    Table,
    tokenize,
)
from lib.references import format_entry  # noqa: E402

DEFAULT_OUT = PAPER_DIR / "MedBridge_IEEE_Paper.docx"
DEFAULT_LOG = PAPER_DIR / "build" / "docx_build_log.json"
FONT = "Times New Roman"
MONO = "Courier New"

#: IEEE conference template geometry (inches).
PAPER_SIZES: dict[str, dict[str, float]] = {
    "letter": {
        "width": 8.5,
        "height": 11.0,
        "top": 0.75,
        "bottom": 1.0,
        "side": 0.625,
        "gap": 0.25,
    },
    "a4": {
        "width": 8.27,
        "height": 11.69,
        "top": 0.75,
        "bottom": 1.69,
        "side": 0.551,
        "gap": 0.166,
    },
}
#: Unicode superscripts (p-values such as 5.4×10⁻¹⁶) -> plain characters in superscript runs.
_SUPERSCRIPT_MAP: dict[str, str] = {
    **{sup: str(d) for d, sup in enumerate("⁰¹²³⁴⁵⁶⁷⁸⁹")},
    "\u207b": "\u2212",
    "\u207a": "+",
}
_SUPERSCRIPT_TABLE = str.maketrans(_SUPERSCRIPT_MAP)
_SUPERSCRIPT_RUN = re.compile("([" + "".join(_SUPERSCRIPT_MAP) + "]+)")

_HIGHLIGHT_PREFIXES = (
    "[PENDING",
    "[TABLE PENDING",
    "[FIGURE PENDING",
    "[TO BE FILLED",
    "[TODO",
    "[TBD",
)


def _breakable(text: str) -> str:
    """Allow line breaks after "/" inside long slash-joined words (never in URLs or DOIs)."""

    def fix(m: re.Match[str]) -> str:
        word = m.group(0)
        if "://" in word or word.startswith(("www.", "10.")) or len(word) < 25:
            return word
        return word.replace("/", "/\u200b")

    return re.sub(r"\S+", fix, text)


def curly_quotes(text: str) -> str:
    """Typographic quotes for running text: "x" -> “x”, it's -> it’s."""
    text = re.sub(r'(^|[\s(\[{\u2014])"', "\\1\u201c", text)
    text = text.replace('"', "\u201d")
    text = re.sub(r"(^|[\s(\[{\u2014])'", "\\1\u2018", text)
    return text.replace("'", "\u2019")


# ----------------------------------------------------------------------
# Styles and sections
# ----------------------------------------------------------------------


def _style(doc: Any, name: str) -> Any:
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _fmt(
    style: Any,
    *,
    size: float,
    bold: bool = False,
    italic: bool = False,
    small_caps: bool = False,
    align: Any = WD_ALIGN_PARAGRAPH.JUSTIFY,
    before: float = 0,
    after: float = 0,
    first_indent: float | None = None,
    left_indent: float | None = None,
    keep_next: bool = False,
    font: str = FONT,
) -> None:
    style.font.name = font
    dh.set_rfonts(style.element.get_or_add_rPr(), font)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.small_caps = small_caps
    style.font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Inches(first_indent) if first_indent is not None else None
    pf.left_indent = Inches(left_indent) if left_indent is not None else None
    pf.keep_with_next = keep_next
    pf.page_break_before = False
    pf.widow_control = True


def setup_styles(doc: Any) -> None:
    """IEEE conference styles (sizes from the IEEE Word template)."""
    center, left = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    _fmt(doc.styles["Normal"], size=10)
    _fmt(_style(doc, "IEEE Body"), size=10, first_indent=0.14)
    _fmt(_style(doc, "IEEE Title"), size=24, align=center, after=12)
    _fmt(_style(doc, "IEEE Author"), size=11, align=center)
    _fmt(_style(doc, "IEEE Affiliation"), size=10, align=center)
    _fmt(_style(doc, "IEEE Abstract"), size=9, bold=True, after=6)
    _fmt(
        doc.styles["Heading 1"],
        size=10,
        small_caps=True,
        align=center,
        before=8,
        after=4,
        keep_next=True,
    )
    _fmt(
        doc.styles["Heading 2"], size=10, italic=True, align=left, before=6, after=3, keep_next=True
    )
    _fmt(
        doc.styles["Heading 3"],
        size=10,
        italic=True,
        align=left,
        before=3,
        after=0,
        first_indent=0.14,
        keep_next=True,
    )
    _fmt(_style(doc, "IEEE Caption"), size=8, align=center, before=3, after=6)
    _fmt(_style(doc, "IEEE Table Caption"), size=8, align=center, before=6, after=0, keep_next=True)
    _fmt(_style(doc, "IEEE Table"), size=8, align=left)
    _fmt(_style(doc, "IEEE Note"), size=7, italic=True, align=left, before=1, after=6)
    _fmt(_style(doc, "IEEE List"), size=10, left_indent=0.25, first_indent=-0.13)
    _fmt(_style(doc, "IEEE Code"), size=7, align=left, font=MONO)
    _fmt(_style(doc, "IEEE Reference"), size=8, left_indent=0.3, first_indent=-0.3)
    ref = doc.styles["IEEE Reference"].paragraph_format
    ref.tab_stops.add_tab_stop(Inches(0.3))
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        rpr = doc.styles[name].element.get_or_add_rPr()
        for tag in ("w:color",):
            for el in rpr.findall(qn(tag)):
                el.set(qn("w:val"), "000000")
                for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
                    if el.get(qn(attr)) is not None:
                        del el.attrib[qn(attr)]


def set_geometry(section: Any, geo: dict[str, float]) -> None:
    section.page_width = Inches(geo["width"])
    section.page_height = Inches(geo["height"])
    section.top_margin = Inches(geo["top"])
    section.bottom_margin = Inches(geo["bottom"])
    section.left_margin = Inches(geo["side"])
    section.right_margin = Inches(geo["side"])
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)


def set_columns(section: Any, num: int, gap_in: float) -> None:
    """``<w:cols w:num="2" w:space="..."/>`` in the section's ``sectPr``."""
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        # schema order: ... pgNumType, cols, formProt, vAlign, noEndnote, titlePg, ..., docGrid
        anchor = None
        for tag in (
            "w:formProt",
            "w:vAlign",
            "w:noEndnote",
            "w:titlePg",
            "w:textDirection",
            "w:bidi",
            "w:rtlGutter",
            "w:docGrid",
            "w:printerSettings",
            "w:sectPrChange",
        ):
            anchor = sect_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(cols)
        else:
            sect_pr.append(cols)
    for child in list(cols):
        cols.remove(child)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(round(gap_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def section_columns(section: Any) -> int:
    cols = section._sectPr.find(qn("w:cols"))
    return int(cols.get(qn("w:num"), "1")) if cols is not None else 1


# ----------------------------------------------------------------------
# Renderer
# ----------------------------------------------------------------------


class PaperRenderer:
    def __init__(self, paper: pc.Paper, *, paper_size: str = "letter") -> None:
        self.paper = paper
        self.geo = PAPER_SIZES[paper_size]
        self.text_width = self.geo["width"] - 2 * self.geo["side"]
        self.col_width = (self.text_width - self.geo["gap"]) / 2
        self.doc = Document()
        setup_styles(self.doc)
        set_geometry(self.doc.sections[0], self.geo)
        set_columns(self.doc.sections[0], 1, self.geo["gap"])
        self.wide = False
        self.letter = 0
        self.sub3 = 0
        self.pending: list[str] = []
        self.stats: dict[str, int] = {
            "figures": 0,
            "figures_pending": 0,
            "tables": 0,
            "tables_pending": 0,
            "algorithms": 0,
            "highlights": 0,
            "to_be_filled": 0,
        }

    # -- sections ---------------------------------------------------------

    def new_section(self, columns: int) -> None:
        section = self.doc.add_section(WD_SECTION.CONTINUOUS)
        set_geometry(section, self.geo)
        set_columns(section, columns, self.geo["gap"])

    def begin_wide(self) -> None:
        if not self.wide:
            self.new_section(1)
            self.wide = True

    def end_wide(self) -> None:
        if self.wide:
            self.new_section(2)
            self.wide = False

    @property
    def avail(self) -> float:
        return self.text_width if self.wide else self.col_width

    # -- inline -------------------------------------------------------------

    def _run(
        self,
        para: Any,
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        size: float | None = None,
        highlight: bool = False,
        small_caps: bool = False,
        font: str | None = None,
        superscript: bool = False,
    ) -> Any:
        if not superscript and _SUPERSCRIPT_RUN.search(text):
            # "5.4×10⁻¹⁶": Unicode superscripts become real superscript runs (fonts such
            # as Times New Roman lack some of those glyphs)
            run = None
            for part in _SUPERSCRIPT_RUN.split(text):
                if not part:
                    continue
                sup = part[0] in _SUPERSCRIPT_MAP
                run = self._run(
                    para,
                    part.translate(_SUPERSCRIPT_TABLE) if sup else part,
                    bold=bold,
                    italic=italic,
                    size=size,
                    highlight=highlight,
                    small_caps=small_caps,
                    font=font,
                    superscript=sup,
                )
            return run
        run = para.add_run(text)
        if superscript:
            run.font.superscript = True
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if small_caps:
            run.font.small_caps = True
        if size:
            run.font.size = Pt(size)
        if font:
            dh.set_run_font(run, font)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            self.stats["highlights"] += 1
        return run

    def inline(
        self,
        para: Any,
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        size: float | None = None,
        markup: bool = True,
    ) -> None:
        if not markup:
            hl = str(text).lstrip().upper().startswith(_HIGHLIGHT_PREFIXES)
            self._run(para, str(text), bold=bold, italic=italic, size=size, highlight=hl)
            return
        for tok in tokenize(text or ""):
            b, i = bold or tok.bold, italic or tok.italic
            if tok.kind == "text":
                self._run(para, curly_quotes(tok.text), bold=b, italic=i, size=size)
            elif tok.kind == "code":
                code = tok.text
                if len(code) > 20:  # long paths may wrap inside narrow columns
                    code = re.sub(r"([/._])(?=\S)", "\\1\u200b", code)
                self._run(para, code, bold=b, italic=i, size=(size or 10) - 1, font=MONO)
            elif tok.kind == "cite":
                self._cite(para, tok.keys, b, i, size)
            elif tok.kind == "value":
                value, ok = self.paper.ctx.render(tok.text)
                self._run(para, value, bold=b, italic=i, size=size, highlight=not ok)
                if not ok:
                    self.pending.append(f"value: {tok.text}")
            elif tok.kind == "fill":
                self._run(para, tok.text, bold=b, italic=i, size=size, highlight=True)
                if tok.text.upper().startswith("[TO BE FILLED"):
                    self.stats["to_be_filled"] += 1
                else:
                    self.pending.append(tok.text)
            elif tok.kind == "link":
                dh.add_hyperlink(para, tok.text, tok.url, bold=b, italic=i)
            elif tok.kind == "br":
                para.add_run().add_break()

    def _cite(
        self, para: Any, keys: list[str], bold: bool, italic: bool, size: float | None
    ) -> None:
        refs = [k for k in keys if not pc.is_xref(k)]
        parts: list[tuple[str, bool]] = []
        for key in (k for k in keys if pc.is_xref(k)):
            shown = self.paper.xref(key)
            if shown is None:
                parts.append((f"[REF NOT FOUND: {key}]", True))
                self.pending.append(f"cross-reference not found: {key}")
            else:
                parts.append((shown, False))
        if refs:
            label, unknown = self.paper.registry.label(refs)
            if label:
                parts.append((label, False))
            for key in unknown:
                parts.append((f"[CITATION NOT FOUND: {key}]", True))
                self.pending.append(f"citation not found: {key}")
        for n, (text, hl) in enumerate(parts):
            if n:
                self._run(para, ", ", bold=bold, italic=italic, size=size)
            self._run(para, text, bold=bold, italic=italic, size=size, highlight=hl)

    def pending_para(self, text: str) -> Any:
        para = self.doc.add_paragraph(style="Normal")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(para, text, bold=True, highlight=True)
        self.pending.append(text)
        return para

    # -- front matter ---------------------------------------------------------

    def front(self) -> None:
        paper = self.paper
        title = self.doc.add_paragraph(style="IEEE Title")
        self.inline(title, paper.title)
        authors = paper.authors or [pc.Author("[TO BE FILLED: authors]")]
        per_row = min(3, len(authors))
        rows = math.ceil(len(authors) / per_row)
        table = self.doc.add_table(rows=rows, cols=per_row)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        dh.set_column_widths(table, [self.text_width / per_row] * per_row)
        for idx, author in enumerate(authors):
            cell = table.rows[idx // per_row].cells[idx % per_row]
            name = cell.paragraphs[0]
            name.style = self.doc.styles["IEEE Author"]
            self.inline(name, author.name)
            for n, line in enumerate(author.affiliation):
                para = cell.add_paragraph(style="IEEE Affiliation")
                self.inline(para, line, italic=n == 0)
            if author.email:
                para = cell.add_paragraph(style="IEEE Affiliation")
                self.inline(para, author.email)
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(6)
        self.new_section(2)
        for n, text in enumerate(paper.abstract):
            para = self.doc.add_paragraph(style="IEEE Abstract")
            if n == 0:
                self._run(para, "Abstract", bold=True, italic=True)
                self._run(para, "\u2014", bold=True)
            else:
                para.paragraph_format.first_line_indent = Inches(0.14)
            self.inline(para, text, bold=True)
        if paper.index_terms:
            para = self.doc.add_paragraph(style="IEEE Abstract")
            self._run(para, "Index Terms", bold=True, italic=True)
            self._run(para, "\u2014", bold=True)
            self.inline(para, paper.index_terms, bold=True)

    # -- blocks -----------------------------------------------------------------

    def heading(self, h: Heading) -> None:
        el = self.paper.element(h)
        if h.level == 1:
            self.letter = 0
            para = self.doc.add_paragraph(style="Heading 1")
            prefix = f"{el.display}. " if el is not None else ""
            self.inline(para, prefix + h.text)
        elif h.level == 2:
            self.sub3 = 0
            para = self.doc.add_paragraph(style="Heading 2")
            if h.numbered and el is not None:
                self.letter = el.number
                prefix = f"{chr(64 + el.number)}. "
            else:
                prefix = ""
            self.inline(para, prefix + h.text)
        else:
            para = self.doc.add_paragraph(style="Heading 3")
            if h.level == 3 and h.numbered:
                self.sub3 += 1
                self.inline(para, f"{self.sub3}) {h.text}")
            else:
                self.inline(para, h.text)

    def paragraph(self, p: Paragraph) -> None:
        para = self.doc.add_paragraph(style="IEEE Body")
        if p.style == "quote":
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.right_indent = Inches(0.25)
            para.paragraph_format.first_line_indent = Inches(0)
        elif p.style == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.inline(para, p.text)

    def list_block(self, lb: ListBlock) -> None:
        counters: dict[int, int] = {}
        for item in lb.items:
            level = max(0, min(2, item.level))
            for deeper in [k for k in counters if k > level]:
                counters.pop(deeper)
            para = self.doc.add_paragraph(style="IEEE List")
            para.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
            if item.ordered:
                counters[level] = counters.get(level, 0) + 1
                marker = f"{counters[level]})"
            else:
                marker = "\u2022"
            para.paragraph_format.first_line_indent = Inches(-0.17 if item.ordered else -0.13)
            self._run(para, marker + "\u00a0")
            self.inline(para, item.text)

    def figure(self, f: Figure) -> None:
        el = self.paper.element(f)
        if el is not None and el.wide:
            self.begin_wide()
        path = self.paper.resolve_figure(f.path)
        width = min(f.width or self.avail, self.avail)
        if path is None:
            para = self.pending_para(f"[FIGURE PENDING: {f.path}]")
            dh.keep_with_next(para)
            self.stats["figures_pending"] += 1
        else:
            try:
                from PIL import Image

                with Image.open(path) as img:
                    w_px, h_px = img.size
                max_h = 0.62 * (self.geo["height"] - self.geo["top"] - self.geo["bottom"])
                if w_px and h_px and width * h_px / w_px > max_h:
                    width = max_h * w_px / h_px
            except Exception:  # size probing is best effort
                pass
            para = self.doc.add_paragraph(style="Normal")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(4)
            dh.keep_with_next(para)
            para.add_run().add_picture(str(path), width=Inches(width))
            self.stats["figures"] += 1
        cap = self.doc.add_paragraph(style="IEEE Caption")
        number = el.display if el is not None else "?"
        self._run(cap, f"Fig. {number}. ")
        self.inline(cap, f.caption or (path.name if path else f.path))
        if len(f.caption or "") > 110:
            cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if el is not None and el.wide:
            self.end_wide()

    def _table_caption(self, el: pc.Element | None, caption: str | None) -> None:
        number = el.display if el is not None else "?"
        p1 = self.doc.add_paragraph(style="IEEE Table Caption")
        self._run(p1, f"TABLE {number}")
        if caption:
            p2 = self.doc.add_paragraph(style="IEEE Table Caption")
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(3)
            for tok_text in [caption]:
                self.inline(p2, tok_text)
            for run in p2.runs:
                run.font.small_caps = True

    @staticmethod
    def _cell_border(cell: Any, edge: str, size: int) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = dh.get_or_add_ordered(tc_pr, "tcBorders", dh.TCPR_ORDER)
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")

    def table(self, t: Table) -> None:
        if not t.header and not t.rows:
            return
        el = self.paper.element(t)
        wide = el is not None and el.wide
        if wide:
            self.begin_wide()
        if el is not None:
            self._table_caption(el, t.caption)
        ncols = max(len(t.header), max((len(r) for r in t.rows), default=0))
        header = t.header + [""] * (ncols - len(t.header))
        rows = [r + [""] * (ncols - len(r)) for r in t.rows]
        size = t.font_size or (8.0 if ncols <= 5 else 7.5 if ncols <= 7 else 7.0)
        table = self.doc.add_table(rows=1 + len(rows), cols=ncols)
        compact = bool(t.font_size) and bool(t.col_widths) and len(t.col_widths) == ncols
        if compact:
            # a generated table with its own font size and column weights (paper_tables.py)
            # was measured by hand: use the weights as they are, with narrow cell margins
            widths = [w * self.avail / sum(t.col_widths) for w in t.col_widths]
        else:
            widths = dh.compute_column_widths(
                header, rows, total=self.avail, weights=t.col_widths, font_size=size, padding=0.1
            )
        dh.set_column_widths(table, widths)
        side = 30 if compact else 50
        dh.set_cell_margins(table, top=20, bottom=20, left=side, right=side)
        tbl_pr = table._tbl.tblPr
        borders = dh.get_or_add_ordered(tbl_pr, "tblBorders", dh.TBLPR_ORDER)
        for edge, val, sz in (
            ("top", "single", 8),
            ("left", "nil", 0),
            ("bottom", "single", 8),
            ("right", "nil", 0),
            ("insideH", "nil", 0),
            ("insideV", "nil", 0),
        ):
            b = OxmlElement(f"w:{edge}")
            b.set(qn("w:val"), val)
            if sz:
                b.set(qn("w:sz"), str(sz))
                b.set(qn("w:space"), "0")
                b.set(qn("w:color"), "000000")
            borders.append(b)
        align_map = {
            "l": WD_ALIGN_PARAGRAPH.LEFT,
            "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        aligns = (t.align + ["l"] * ncols)[:ncols]
        for c, text in enumerate(header):
            cell = table.rows[0].cells[c]
            para = cell.paragraphs[0]
            para.style = self.doc.styles["IEEE Table"]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            dh.keep_with_next(para)
            self.inline(para, text, bold=True, size=size, markup=t.markup)
            self._cell_border(cell, "bottom", 4)
        dh.mark_header_row(table.rows[0])
        dh.keep_row_together(table.rows[0])
        for r_idx, row in enumerate(rows, start=1):
            tr = table.rows[r_idx]
            dh.keep_row_together(tr)
            is_total = bool(row and str(row[0]).strip().lower() in {"total", "overall"})
            for c, text in enumerate(row):
                para = tr.cells[c].paragraphs[0]
                para.style = self.doc.styles["IEEE Table"]
                para.alignment = align_map.get(aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                # keep the whole (short) table in one column: every row but the last
                # keeps with the next one
                dh.keep_with_next(para, r_idx < len(rows) and len(rows) <= 25)
                self.inline(para, str(text), size=size, markup=t.markup, bold=is_total)
                if is_total:
                    self._cell_border(tr.cells[c], "top", 4)
        if t.note:
            note = self.doc.add_paragraph(style="IEEE Note")
            self.inline(note, t.note, markup=False)
        else:
            spacer = self.doc.add_paragraph(style="Normal")
            spacer.paragraph_format.space_after = Pt(4)
        self.stats["tables"] += 1
        if wide:
            self.end_wide()

    def pending_table(self, t: PendingTable) -> None:
        el = self.paper.element(t)
        self._table_caption(el, t.caption)
        self.pending_para(f"[TABLE PENDING: {t.key}]").paragraph_format.space_after = Pt(6)
        self.stats["tables_pending"] += 1

    def boxed(
        self, lines: list[str], *, caption: tuple[str, str] | None = None, fill: str | None = None
    ) -> None:
        if caption:
            cap = self.doc.add_paragraph(style="IEEE Table Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cap.paragraph_format.space_after = Pt(2)
            self._run(cap, caption[0], bold=True)
            self.inline(cap, caption[1])
        usable_pt = (self.avail - 0.16) * 72
        longest = max((len(line.rstrip()) for line in lines), default=1) or 1
        size = max(5.5, min(8.0, usable_pt / (0.6 * longest)))
        size = math.floor(size * 2) / 2
        table = self.doc.add_table(rows=1, cols=1)
        dh.set_column_widths(table, [self.avail])
        dh.set_table_borders(table, size=6, color="000000")
        dh.set_cell_margins(table, top=40, bottom=40, left=70, right=70)
        cell = table.rows[0].cells[0]
        if fill:
            dh.shade_cell(cell, fill)
        dh.keep_row_together(table.rows[0])
        for n, line in enumerate(lines or [""]):
            para = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
            para.style = self.doc.styles["IEEE Code"]
            dh.keep_with_next(para, n < len(lines) - 1)
            run = para.add_run(line if line else " ")
            dh.set_run_font(run, MONO)
            run.font.size = Pt(size)
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(4)

    def algorithm(self, a: Algorithm) -> None:
        el = self.paper.element(a)
        number = el.display if el is not None else "?"
        self.boxed(a.lines, caption=(f"Algorithm {number}: ", a.title))
        self.stats["algorithms"] += 1

    def render(self, blocks: list[Block]) -> None:
        for block in blocks:
            if isinstance(block, Heading):
                self.heading(block)
            elif isinstance(block, Paragraph):
                self.paragraph(block)
            elif isinstance(block, ListBlock):
                self.list_block(block)
            elif isinstance(block, Figure):
                self.figure(block)
            elif isinstance(block, Table):
                self.table(block)
            elif isinstance(block, PendingTable):
                self.pending_table(block)
            elif isinstance(block, Algorithm):
                self.algorithm(block)
            elif isinstance(block, CodeBlock):
                self.boxed(block.lines, fill=dh.CODE_FILL)
            elif isinstance(block, Pending):
                self.pending_para(block.text)
            elif isinstance(block, PageBreak):
                continue  # page breaks have no place in a conference paper

    def references(self) -> None:
        para = self.doc.add_paragraph(style="Heading 1")
        para.add_run("References")
        bib = self.paper.registry.bibliography()
        if not bib:
            self.pending_para("[PENDING: no citations in the content yet]")
            return
        for number, entry in bib:
            # left-aligned: long unbreakable tokens (URLs, "AHA/ACC/...") would otherwise
            # stretch justified lines in a 3.5 in column
            p = self.doc.add_paragraph(style="IEEE Reference")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(f"[{number}]\t")
            for seg in format_entry(entry):
                r = p.add_run(_breakable(seg.text))
                if seg.italic:
                    r.italic = True

    def build(self) -> Any:
        self.front()
        self.render(self.paper.blocks)
        self.end_wide()
        self.references()
        # A closing continuous break makes Word/LibreOffice balance the last page's columns.
        self.new_section(1)
        props = self.doc.core_properties
        props.title = re.sub(r"\[[^\]]*\]", "", self.paper.title).strip()
        props.author = "; ".join(a.name for a in self.paper.authors)
        props.subject = "IEEE conference paper draft (generated by report/paper/build_paper.py)"
        props.keywords = self.paper.index_terms
        return self.doc


# ----------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------


@dataclass
class BuildResult:
    docx: Path
    pdf: Path | None = None
    pages: int | None = None
    log: dict[str, Any] = field(default_factory=dict)


def build(
    content_dir: Path = pc.DEFAULT_CONTENT,
    out: Path = DEFAULT_OUT,
    *,
    repo_root: Path = pc.REPO_ROOT,
    paper_size: str = "letter",
    anonymous: bool = False,
    pdf: bool = False,
    log_path: Path | None = DEFAULT_LOG,
) -> BuildResult:
    paper = pc.load_paper(Path(content_dir), Path(repo_root), anonymous=anonymous)
    renderer = PaperRenderer(paper, paper_size=paper_size)
    doc = renderer.build()
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    result = BuildResult(docx=out)
    log = pc.base_log(paper)
    log.update(
        {
            "built_at": datetime.now(UTC).isoformat(timespec="seconds"),
            "output": pc._rel(out, Path(repo_root)),
            "paper_size": paper_size,
            "stats": renderer.stats,
            "pending_markers": sorted(set(renderer.pending)),
            "word_count_body": _word_count(paper),
        }
    )
    if pdf:
        from lib import libreoffice

        conv = libreoffice.convert_to_pdf(out, update_indexes=False)
        result.pdf = conv.pdf
        log["pdf"] = {
            "method": conv.method,
            "detail": conv.detail,
            "path": pc._rel(conv.pdf, Path(repo_root)) if conv.pdf else None,
        }
        if conv.pdf:
            result.pages = libreoffice.pdf_page_count(conv.pdf)
            log["pdf"]["pages"] = result.pages
    result.log = log
    if log_path is not None:
        log_path = Path(log_path)
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text(json.dumps(log, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return result


def _word_count(paper: pc.Paper) -> int:
    words = 0
    for block in paper.blocks:
        for text in pc._texts(block):
            words += len(re.findall(r"\w+", re.sub(r"\{\{[^}]*\}\}|\[@[^\]]*\]", "x", text)))
    return words


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the MedBridge IEEE paper (DOCX).")
    parser.add_argument("--content-dir", type=Path, default=pc.DEFAULT_CONTENT)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--paper", choices=sorted(PAPER_SIZES), default="letter")
    parser.add_argument("--anonymous", action="store_true", help="double-blind review copy")
    parser.add_argument("--pdf", action="store_true", help="also write a PDF preview (LibreOffice)")
    parser.add_argument("--log", type=Path, default=DEFAULT_LOG)
    parser.add_argument("--no-log", action="store_true")
    parser.add_argument(
        "--strict",
        action="store_true",
        help="exit 1 on unknown citation keys or unresolved cross-references",
    )
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args(argv)
    if not args.content_dir.is_dir():
        print(f"content directory not found: {args.content_dir}", file=sys.stderr)
        return 2
    result = build(
        args.content_dir,
        args.out,
        paper_size=args.paper,
        anonymous=args.anonymous,
        pdf=args.pdf,
        log_path=None if args.no_log else args.log,
    )
    log = result.log
    if not args.quiet:
        print(f"DOCX: {result.docx}")
        if result.pdf:
            print(f"PDF:  {result.pdf} ({result.pages} pages)")
        print(
            f"sections {log['counts']['section']}, figures {log['counts']['figure']}, "
            f"tables {log['counts']['table']}, algorithms {log['counts']['algorithm']}, "
            f"references {log['references']}, body words ~{log['word_count_body']}"
        )
        for key in ("citation_keys_missing", "unresolved_xrefs", "pending_markers", "warnings"):
            if log.get(key):
                print(f"{key}: {len(log[key])}")
                for item in log[key][:20]:
                    print(f"  - {item}")
    if args.pdf and result.pdf is None:
        return 3
    if args.strict and (log["citation_keys_missing"] or log["unresolved_xrefs"]):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
