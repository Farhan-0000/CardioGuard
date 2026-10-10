"""
Drawing primitives for the MedBridge review deck (python-pptx), with text-fit measurement.

python-pptx cannot lay text out, so every text box and table drawn through :class:`Kit`
is measured first: words are wrapped with the real Arial advance widths (Pillow reads the
TrueType font installed on macOS, Windows or Linux; a conservative estimate is used when
no font file is found). The kit picks the largest font size between ``size`` and
``min_size`` that fits the box, and records an *issue* when even the minimum size
overflows, when a shape leaves the slide, or when two layout regions overlap. The builder
prints those issues and ``--strict`` turns them into a non-zero exit code, so a deck that
regenerates with longer numbers never silently overflows.

Colours follow the report diagrams (``report/diagrams/_common.py``): amber for the
answerability gates, red for the emergency fast-path, blue for local LLM calls, grey for
deterministic code, green for pre-vetted templates and privacy controls, purple for
storage. Action badges use the colours of the web UI (``medbridge/web/css/base.css``).
"""

from __future__ import annotations

import re
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

FONT = "Arial"
SLIDE_W = 13.333  # inches (16:9)
SLIDE_H = 7.5
EMU_PER_IN = 914400
#: Line height as a multiple of the font size (Arial ascent + descent is about 1.15 em).
LINE_FACTOR = 1.2
#: Width safety factor: wrap a little earlier than the measured width.
WIDTH_SAFETY = 0.96


def emu(inches: float) -> Emu:
    return Emu(round(inches * EMU_PER_IN))


# ----------------------------------------------------------------------
# Palette
# ----------------------------------------------------------------------


@dataclass(frozen=True)
class Tone:
    """Fill, border and text colour (hex RGB without '#')."""

    fill: str
    line: str
    text: str


NAVY = "13234B"
NAVY_DARK = "0D1834"
INK = "1F2933"
MUTED = "5A6472"
SUBTLE = "8A94A6"
RULE = "D5DBE5"
PANEL = "F4F6FA"
WHITE = "FFFFFF"
AMBER = "B7740A"
HIGHLIGHT = "FFF176"
HIGHLIGHT_TEXT = "5C4300"

TONES: dict[str, Tone] = {
    "gate": Tone("FFE9B8", "B7740A", "4D3200"),
    "emergency": Tone("FAD4D0", "B83227", "5C1410"),
    "llm": Tone("D8E7FA", "2B62A8", "0F2F57"),
    "det": Tone("ECEFF3", "5A6472", "1F2933"),
    "safe": Tone("D7F0DE", "2E7D4F", "123D25"),
    "storage": Tone("E8DCF7", "6942B0", "2E1A57"),
    "note": Tone("F4F6FA", "A0AEC0", "2D3748"),
    "navy": Tone("13234B", "13234B", "FFFFFF"),
    "pending": Tone("FFFBEA", "C9A227", "5C4300"),
}

ACTION_TONES: dict[str, Tone] = {
    "ANSWER": Tone("E8F8EE", "16A34A", "166534"),
    "GENERALIZE": Tone("EAF1FE", "2563EB", "1E40AF"),
    "SOFT-ASK": Tone("FEF6E4", "D97706", "92400E"),
    "ABSTAIN": Tone("F1F1F3", "71717A", "3F3F46"),
    "ESCALATE": Tone("FDECEC", "DC2626", "991B1B"),
}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


# ----------------------------------------------------------------------
# Inline markup
# ----------------------------------------------------------------------

#: Team placeholders and pending markers get a yellow highlight.
PLACEHOLDER_RE = re.compile(
    r"\[(?:TO BE FILLED|PENDING|TABLE PENDING|FIGURE PENDING|CITATION NOT FOUND|TODO|TBD)"
    r"\b(?:[^\[\]]|\[[^\[\]]*\])*\]"
)


@dataclass(frozen=True)
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    highlight: bool = False
    color: str | None = None


def parse_runs(text: str, *, bold: bool = False, italic: bool = False) -> list[Run]:
    """``**bold**`` toggles bold; placeholder markers become highlighted runs."""
    runs: list[Run] = []
    pos = 0
    for match in PLACEHOLDER_RE.finditer(text):
        runs.extend(_bold_runs(text[pos : match.start()], bold, italic))
        runs.append(Run(match.group(0), bold=True, italic=italic, highlight=True))
        pos = match.end()
    runs.extend(_bold_runs(text[pos:], bold, italic))
    return [r for r in runs if r.text]


def _bold_runs(text: str, bold: bool, italic: bool) -> list[Run]:
    parts = text.split("**")
    return [
        Run(part, bold=(bold != (i % 2 == 1)), italic=italic)
        for i, part in enumerate(parts)
        if part
    ]


def has_placeholder(text: str) -> bool:
    return bool(PLACEHOLDER_RE.search(text))


@dataclass
class P:
    """One paragraph of a text box (``text`` uses the inline markup above)."""

    text: str
    bullet: bool = False
    level: int = 0
    scale: float = 1.0  # relative to the box font size
    bold: bool = False
    italic: bool = False
    color: str | None = None
    align: str = "l"  # l | c | r
    space_before: float = 0.0  # points
    bullet_color: str | None = None


def as_paras(content: str | P | Sequence[str | P]) -> list[P]:
    if isinstance(content, (str, P)):
        content = [content]
    out: list[P] = []
    for item in content:
        if isinstance(item, P):
            out.append(item)
        else:
            for line in str(item).split("\n"):
                out.append(P(line))
    return out


# ----------------------------------------------------------------------
# Measurement
# ----------------------------------------------------------------------

_REGULAR_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "~/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf",
)
_BOLD_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
    "~/Library/Fonts/Arial Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
    "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf",
)


class Measurer:
    """Advance widths of Arial text in points (Pillow; falls back to an estimate)."""

    def __init__(self) -> None:
        self._regular = self._load(_REGULAR_CANDIDATES)
        self._bold = self._load(_BOLD_CANDIDATES) or self._regular
        self.source = "truetype" if self._regular is not None else "estimate"
        self._cache: dict[tuple[str, bool], float] = {}

    @staticmethod
    def _load(candidates: Iterable[str]) -> Any | None:
        try:
            from PIL import ImageFont
        except ImportError:
            return None
        for cand in candidates:
            path = Path(cand).expanduser()
            if path.is_file():
                try:
                    return ImageFont.truetype(str(path), 100)
                except OSError:
                    continue
        try:  # DejaVu Sans ships with matplotlib; it is wider than Arial (conservative)
            import matplotlib

            dejavu = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf"
            if dejavu.is_file():
                return ImageFont.truetype(str(dejavu), 100)
        except (ImportError, OSError):
            pass
        return None

    def width(self, text: str, size: float, bold: bool = False) -> float:
        key = (text, bold)
        unit = self._cache.get(key)
        if unit is None:
            font = self._bold if bold else self._regular
            if font is None:
                unit = len(text) * (60.0 if bold else 56.0)
            else:
                unit = float(font.getlength(text))
            self._cache[key] = unit
        return unit * size / 100.0

    def count_lines(self, runs: Sequence[Run], size: float, avail_pt: float) -> int:
        """Greedy word wrap of ``runs`` at ``size`` into a line of ``avail_pt`` points."""
        avail = max(avail_pt * WIDTH_SAFETY, 1.0)
        chunks: list[tuple[str, bool]] = []
        for run in runs:
            # break only at ASCII spaces: a no-break space keeps "5.8 %" on one line
            for match in _CHUNK_RE.finditer(run.text):
                chunks.append((match.group(0), run.bold))
        if not chunks:
            return 1
        lines, cur = 1, 0.0
        for chunk, bold in chunks:
            word = chunk.rstrip()
            w_word = self.width(word, size, bold) if word else 0.0
            w_all = self.width(chunk, size, bold)
            if cur > 0 and cur + w_word > avail:
                lines += 1
                cur = 0.0
            if w_word > avail:  # a single token longer than the line breaks across lines
                extra = int(w_word // avail)
                lines += extra
                cur = w_word - extra * avail + (w_all - w_word)
            else:
                cur += w_all
        return lines

    def longest_token(self, runs: Sequence[Run], size: float) -> float:
        """Width in points of the widest unbreakable token."""
        widest = 0.0
        for run in runs:
            for token in _BREAK_RE.split(run.text):
                if token:
                    widest = max(widest, self.width(token, size, run.bold))
        return widest


_CHUNK_RE = re.compile(r"[^ \t\n]+[ \t\n]*|[ \t\n]+")
_BREAK_RE = re.compile(r"[ \t\n]+")


def balance_widths(widths: list[float], mins: list[float]) -> tuple[list[float], bool]:
    """Widen columns below their minimum by taking width from columns with slack."""
    w = list(widths)
    for _ in range(4):
        deficit = sum(max(0.0, m - x) for m, x in zip(mins, w, strict=True))
        if deficit <= 1e-6:
            return w, True
        slack = [max(0.0, x - m) for m, x in zip(mins, w, strict=True)]
        total = sum(slack)
        if total <= deficit:
            return w, False
        w = [
            m if x < m else x - deficit * s / total for x, m, s in zip(w, mins, slack, strict=True)
        ]
    return w, sum(max(0.0, m - x) for m, x in zip(mins, w, strict=True)) <= 1e-6


# ----------------------------------------------------------------------
# Results
# ----------------------------------------------------------------------


@dataclass
class TextFit:
    size: float
    needed_in: float
    available_in: float
    overflow: bool
    shape: Any = None


@dataclass
class TableFit:
    size: float
    height_in: float
    overflow: bool
    shape: Any = None


@dataclass
class Issue:
    slide: str
    element: str
    kind: str  # overflow | off-slide | overlap
    detail: str

    def as_dict(self) -> dict[str, str]:
        return {
            "slide": self.slide,
            "element": self.element,
            "kind": self.kind,
            "detail": self.detail,
        }


@dataclass
class CellStyle:
    fill: str | None = None
    color: str | None = None
    bold: bool | None = None


# ----------------------------------------------------------------------
# Low-level XML helpers
# ----------------------------------------------------------------------

_RPR_SUCCESSORS = tuple(
    qn(t)
    for t in (
        "a:uLnTx",
        "a:uLn",
        "a:uFillTx",
        "a:uFill",
        "a:latin",
        "a:ea",
        "a:cs",
        "a:sym",
        "a:hlinkClick",
        "a:hlinkMouseOver",
        "a:rtl",
        "a:extLst",
    )
)


def _highlight(run: Any, color: str = HIGHLIGHT) -> None:
    r_pr = run._r.get_or_add_rPr()
    hl = etree.Element(qn("a:highlight"))
    etree.SubElement(hl, qn("a:srgbClr"), val=color)
    for child in r_pr:
        if child.tag in _RPR_SUCCESSORS:
            child.addprevious(hl)
            break
    else:
        r_pr.append(hl)


def _set_bullet(paragraph: Any, level: int, indent_in: float, color: str | None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(int(emu(indent_in * (level + 1)))))
    p_pr.set("indent", str(-int(emu(indent_in))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buNone", "a:buChar", "a:buAutoNum"):
        for el in p_pr.findall(qn(tag)):
            p_pr.remove(el)
    if color:
        bu_clr = etree.SubElement(p_pr, qn("a:buClr"))
        etree.SubElement(bu_clr, qn("a:srgbClr"), val=color)
    etree.SubElement(p_pr, qn("a:buFont"), typeface=FONT)
    etree.SubElement(p_pr, qn("a:buChar"), char="•" if level == 0 else "–")


def _set_cell_borders(cell: Any, bottom: str | None, width_pt: float = 0.75) -> None:
    """Horizontal rule under the cell only (other sides explicitly empty)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tc_pr.findall(qn(tag)):
            tc_pr.remove(el)
    elements = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        color = bottom if tag == "a:lnB" else None
        ln = etree.Element(qn(tag), w=str(int(Pt(width_pt))) if color else "0")
        if color:
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), val=color)
        else:
            etree.SubElement(ln, qn("a:noFill"))
        elements.append(ln)
    for i, el in enumerate(elements):
        tc_pr.insert(i, el)


def _align(code: str) -> Any:
    return {"c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}.get(code, PP_ALIGN.LEFT)


# ----------------------------------------------------------------------
# Kit
# ----------------------------------------------------------------------


@dataclass
class _Region:
    name: str
    x: float
    y: float
    w: float
    h: float


@dataclass
class Kit:
    """Measured drawing operations; ``issues`` collects layout problems per slide."""

    measurer: Measurer = field(default_factory=Measurer)
    issues: list[Issue] = field(default_factory=list)
    label: str = ""
    _regions: dict[str, list[_Region]] = field(default_factory=dict)

    # -- bookkeeping ---------------------------------------------------

    def begin(self, label: str) -> None:
        self.label = label
        self._regions.setdefault(label, [])

    def issue(self, element: str, kind: str, detail: str) -> None:
        self.issues.append(Issue(self.label, element, kind, detail))

    def claim(self, name: str, x: float, y: float, w: float, h: float) -> None:
        """Register a top-level layout region; overlapping regions are reported."""
        if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
            self.issue(name, "off-slide", f"x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f} in")
        regions = self._regions.setdefault(self.label, [])
        for other in regions:
            dx = min(x + w, other.x + other.w) - max(x, other.x)
            dy = min(y + h, other.y + other.h) - max(y, other.y)
            if dx > 0.02 and dy > 0.02:
                self.issue(name, "overlap", f"overlaps '{other.name}' by {dx:.2f} x {dy:.2f} in")
        regions.append(_Region(name, x, y, w, h))

    # -- shapes --------------------------------------------------------

    def rect(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: str | None = PANEL,
        line: str | None = None,
        line_w: float = 1.0,
        radius: float | None = 0.08,
        dashed: bool = False,
        name: str = "",
    ) -> Any:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
        if radius:
            shape.adjustments[0] = max(0.0, min(0.5, radius / max(min(w, h), 0.01)))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(line_w)
            if dashed:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        # drop the theme style reference: LibreOffice otherwise draws its shadow effect
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        if name:
            shape.name = name
        return shape

    def line(
        self, slide: Any, x1: float, y1: float, x2: float, y2: float, color: str, width: float = 1.5
    ) -> Any:
        connector = slide.shapes.add_connector(1, emu(x1), emu(y1), emu(x2), emu(y2))
        connector.line.color.rgb = rgb(color)
        connector.line.width = Pt(width)
        return connector

    # -- text ----------------------------------------------------------

    def fit_paragraphs(
        self,
        paras: list[P],
        w: float,
        h: float,
        size: float,
        min_size: float,
        margins: tuple[float, float, float, float],
        indent_in: float,
    ) -> tuple[float, float, bool]:
        ml, mt, mr, mb = margins
        avail_h = h - mt - mb
        s = size
        while True:
            needed = 0.0
            for p in paras:
                ps = s * p.scale
                indent = indent_in * (p.level + 1) if p.bullet else 0.0
                avail_pt = (w - ml - mr - indent) * 72.0
                lines = self.measurer.count_lines(
                    parse_runs(p.text, bold=p.bold, italic=p.italic), ps, avail_pt
                )
                needed += (lines * ps * LINE_FACTOR + p.space_before) / 72.0
            if needed <= avail_h + 1e-6 or s <= min_size:
                return s, needed + mt + mb, needed > avail_h + 1e-6
            s = max(min_size, s - 0.5)

    def text(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        content: str | P | Sequence[str | P],
        *,
        size: float = 18,
        min_size: float | None = None,
        color: str = INK,
        anchor: str = "t",
        margins: tuple[float, float, float, float] = (0.06, 0.03, 0.06, 0.03),
        indent_in: float = 0.24,
        name: str = "text",
        claim: bool = False,
    ) -> TextFit:
        paras = as_paras(content)
        min_size = size if min_size is None else min_size
        s, needed, overflow = self.fit_paragraphs(paras, w, h, size, min_size, margins, indent_in)
        if overflow:
            self.issue(
                name,
                "overflow",
                f"needs {needed:.2f} in, box {h:.2f} in at the minimum size {s:.1f} pt",
            )
        if claim:
            self.claim(name, x, y, w, h)
        box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (
            emu(margins[0]),
            emu(margins[1]),
            emu(margins[2]),
            emu(margins[3]),
        )
        tf.vertical_anchor = {"m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}.get(
            anchor, MSO_ANCHOR.TOP
        )
        self.write(tf, paras, s, color, indent_in)
        return TextFit(s, needed, h, overflow, box)

    def write(
        self, tf: Any, paras: list[P], size: float, color: str, indent_in: float = 0.24
    ) -> None:
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _align(para.align)
            p.line_spacing = 1.0
            if para.space_before:
                p.space_before = Pt(para.space_before)
            if para.bullet:
                _set_bullet(p, para.level, indent_in, para.bullet_color or color)
            ps = size * para.scale
            for run in parse_runs(para.text, bold=para.bold, italic=para.italic):
                r = p.add_run()
                r.text = run.text
                f = r.font
                f.name = FONT
                f.size = Pt(ps)
                f.bold = run.bold
                f.italic = run.italic
                f.color.rgb = rgb(
                    HIGHLIGHT_TEXT if run.highlight else (run.color or para.color or color)
                )
                if run.highlight:
                    _highlight(r)
            if not p.runs:  # keep empty paragraphs at the right height
                p.font.size = Pt(ps)

    # -- tables --------------------------------------------------------

    def table(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        header: Sequence[str],
        rows: Sequence[Sequence[str]],
        *,
        widths: Sequence[float] | None = None,
        size: float = 14,
        min_size: float = 10,
        max_h: float | None = None,
        align: Sequence[str] | None = None,
        header_fill: str = NAVY,
        header_color: str = WHITE,
        zebra: str | None = PANEL,
        style: Callable[[int, int, str], CellStyle | None] | None = None,
        name: str = "table",
        claim: bool = True,
        cell_margin: tuple[float, float] = (0.07, 0.035),
    ) -> TableFit:
        n_cols = len(header)
        weights = list(widths or [1.0] * n_cols)
        total = sum(weights)
        base_w = [w * wt / total for wt in weights]
        align = list(align or (["l"] + ["c"] * (n_cols - 1)))
        mx, my = cell_margin
        grid = [list(header)] + [[str(c) for c in r] for r in rows]

        def is_bold(r_i: int, c_i: int, cell: str) -> bool:
            st = style(r_i - 1, c_i, cell) if (style and r_i > 0) else None
            return r_i == 0 or bool(st and st.bold)

        def min_widths(s: float) -> list[float]:
            out = []
            for c_i in range(n_cols):
                widest = 0.0
                for r_i, row in enumerate(grid):
                    for part in str(row[c_i]).split("\n"):
                        runs = parse_runs(part, bold=is_bold(r_i, c_i, row[c_i]))
                        widest = max(widest, self.measurer.longest_token(runs, s))
                out.append(widest / 72.0 / WIDTH_SAFETY + 2 * mx + 0.02)
            return out

        def heights(s: float, col_w: list[float]) -> list[float]:
            out = []
            for r_i, row in enumerate(grid):
                tallest = 0.0
                for c_i, cell in enumerate(row):
                    bold = is_bold(r_i, c_i, cell)
                    lines = 0
                    for part in str(cell).split("\n"):
                        lines += self.measurer.count_lines(
                            parse_runs(part, bold=bold), s, (col_w[c_i] - 2 * mx) * 72.0
                        )
                    tallest = max(tallest, lines * s * LINE_FACTOR / 72.0)
                out.append(tallest + 2 * my)
            return out

        s = size
        while True:
            col_w, words_fit = balance_widths(base_w, min_widths(s))
            row_h = heights(s, col_w)
            total_h = sum(row_h)
            fits = words_fit and (max_h is None or total_h <= max_h + 1e-6)
            if fits or s <= min_size:
                break
            s = max(min_size, s - 0.5)
        overflow = max_h is not None and total_h > max_h + 1e-6
        if overflow:
            self.issue(
                name,
                "overflow",
                f"table needs {total_h:.2f} in, region {max_h:.2f} in at {s:.1f} pt",
            )
        if not words_fit:
            self.issue(name, "overflow", f"a word does not fit its column at {s:.1f} pt")
        if claim:
            self.claim(name, x, y, w, total_h)
        frame = slide.shapes.add_table(len(grid), n_cols, emu(x), emu(y), emu(w), emu(total_h))
        frame.name = name
        tbl = frame.table
        tbl.first_row = True
        tbl.horz_banding = False
        tbl.vert_banding = False
        for c_i, cw in enumerate(col_w):
            tbl.columns[c_i].width = emu(cw)
        for r_i, row in enumerate(grid):
            tbl.rows[r_i].height = emu(row_h[r_i])
            for c_i, value in enumerate(row):
                cell = tbl.cell(r_i, c_i)
                st = style(r_i - 1, c_i, value) if (style and r_i > 0) else None
                if r_i == 0:
                    fill, color, bold = header_fill, header_color, True
                else:
                    fill = (
                        (zebra if (zebra and r_i % 2 == 0) else WHITE)
                        if not (st and st.fill)
                        else st.fill
                    )
                    color = st.color if st and st.color else INK
                    bold = bool(st.bold) if st and st.bold is not None else False
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill)
                cell.margin_left = cell.margin_right = emu(mx)
                cell.margin_top = cell.margin_bottom = emu(my)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                _set_cell_borders(cell, RULE if r_i > 0 else header_fill)
                tf = cell.text_frame
                tf.word_wrap = True
                paras = [
                    P(part, bold=bold, align=align[c_i] if c_i < len(align) else "l")
                    for part in str(value).split("\n")
                ]
                self.write(tf, paras, s, color)
        return TableFit(s, total_h, overflow, frame)

    # -- images --------------------------------------------------------

    def image(
        self,
        slide: Any,
        path: Path | None,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        placeholder_title: str = "Figure pending",
        placeholder_detail: str = "",
        name: str = "image",
        claim: bool = True,
        valign: str = "m",
    ) -> bool:
        """Contain-fit ``path`` in the box; draw a dashed placeholder when it is missing."""
        if claim:
            self.claim(name, x, y, w, h)
        if path is not None and path.is_file():
            try:
                from PIL import Image

                with Image.open(path) as im:
                    iw, ih = im.size
            except Exception:
                iw = ih = 0
            if iw and ih:
                scale = min(w / iw, h / ih)
                dw, dh = iw * scale, ih * scale
                top = {"t": y, "b": y + h - dh}.get(valign, y + (h - dh) / 2)
                pic = slide.shapes.add_picture(
                    str(path), emu(x + (w - dw) / 2), emu(top), emu(dw), emu(dh)
                )
                pic.name = name
                return True
        self.placeholder(slide, x, y, w, h, placeholder_title, placeholder_detail, name=name)
        return False

    def placeholder(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        title: str,
        detail: str = "",
        *,
        name: str = "placeholder",
    ) -> None:
        tone = TONES["pending"]
        self.rect(slide, x, y, w, h, fill=tone.fill, line=tone.line, dashed=True, line_w=1.25)
        paras = [P(f"[PENDING: {title}]", bold=True, align="c")]
        if detail:
            paras.append(P(detail, scale=0.8, align="c", color=MUTED, space_before=4))
        pad = min(0.25, w * 0.06)
        self.text(
            slide,
            x + pad,
            y + pad,
            w - 2 * pad,
            h - 2 * pad,
            paras,
            size=16 if h > 1.8 else 13,
            min_size=9,
            anchor="m",
            name=f"{name}-text",
        )

    # -- composites ----------------------------------------------------

    def chip(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        label: str,
        tone: Tone,
        *,
        size: float = 14,
        min_size: float = 9,
        bold: bool = True,
        name: str = "chip",
        dashed: bool = False,
    ) -> None:
        self.rect(
            slide,
            x,
            y,
            w,
            h,
            fill=tone.fill,
            line=tone.line,
            radius=min(h / 2, 0.14),
            dashed=dashed,
        )
        self.text(
            slide,
            x,
            y,
            w,
            h,
            P(label, bold=bold, align="c"),
            size=size,
            min_size=min_size,
            color=tone.text,
            anchor="m",
            margins=(0.05, 0.01, 0.05, 0.01),
            name=name,
        )

    def card(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        title: str,
        body: str | P | Sequence[str | P],
        tone: Tone,
        *,
        title_size: float = 16,
        body_size: float = 14,
        min_size: float = 10,
        name: str = "card",
        accent: bool = True,
        title_h: float | None = None,
    ) -> None:
        self.claim(name, x, y, w, h)
        self.rect(slide, x, y, w, h, fill=tone.fill, line=tone.line, line_w=1.0)
        if accent:
            self.rect(slide, x, y + 0.12, 0.07, h - 0.24, fill=tone.line, radius=None)
        pad = 0.14
        th = title_h if title_h is not None else (title_size * LINE_FACTOR / 72.0 + 0.1)
        self.text(
            slide,
            x + pad + 0.04,
            y + 0.08,
            w - 2 * pad - 0.04,
            th,
            P(title, bold=True),
            size=title_size,
            min_size=min(min_size + 2, title_size),
            color=tone.text,
            name=f"{name}-title",
        )
        self.text(
            slide,
            x + pad + 0.04,
            y + 0.08 + th,
            w - 2 * pad - 0.04,
            h - th - 0.16,
            body,
            size=body_size,
            min_size=min_size,
            color=INK,
            name=f"{name}-body",
        )

    def stat(
        self,
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        value: str,
        label: str,
        tone: Tone,
        *,
        source: str = "",
        value_size: float = 30,
        label_size: float = 13,
        name: str = "stat",
    ) -> None:
        self.claim(name, x, y, w, h)
        self.rect(slide, x, y, w, h, fill=WHITE, line=tone.line, line_w=1.25)
        self.rect(slide, x, y, w, 0.08, fill=tone.line, radius=None)
        vh = value_size * LINE_FACTOR / 72.0 + 0.06
        self.text(
            slide,
            x + 0.12,
            y + 0.14,
            w - 0.24,
            vh,
            P(value, bold=True),
            size=value_size,
            min_size=14,
            color=tone.text,
            name=f"{name}-value",
        )
        rest = h - vh - 0.2
        paras = [P(label)]
        if source:
            paras.append(P(source, scale=0.82, color=MUTED, space_before=2))
        self.text(
            slide,
            x + 0.12,
            y + 0.14 + vh,
            w - 0.24,
            rest,
            paras,
            size=label_size,
            min_size=9,
            name=f"{name}-label",
        )


def notes(slide: Any, text: str) -> None:
    """Speaker notes (plain text)."""
    slide.notes_slide.notes_text_frame.text = text.strip()


def background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def wrap_ok(kit: Kit) -> bool:
    return not any(i.kind in {"overflow", "off-slide", "overlap"} for i in kit.issues)
