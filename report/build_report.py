#!/usr/bin/env python3
"""
Build the MedBridge (CardioGuard) final report as DOCX (and optionally PDF).

The report is assembled from the content markup files in ``report/content/`` and
every number, table and figure is taken from real results files through the
results context (``report/lib/results_context.py``). See ``report/README.md`` for the
markup contract and the file-naming conventions.

Usage::

    python report/build_report.py                      # -> report/MedBridge_Final_Report.docx
    python report/build_report.py --pdf                # + PDF via LibreOffice (indexes refreshed)
    python report/build_report.py --out /tmp/r.docx --content-dir path/to/content
    python report/build_report.py --strict             # exit 1 on unknown citation keys

The layout follows the VIT Chennai BCSE497J template
(``report/templates/BCSE497J_report_template.docx``). Assembly order: cover page, title
page, declaration, certificate (no page numbers) | abstract, acknowledgement, contents,
list of tables, list of figures, list of symbols and abbreviations (i, ii, ...) |
chapters 1-5, references (numbered, APA style), appendices 1, 2, ... (1, 2, ...).

The template asks each student to submit a copy whose declaration, certificate and
acknowledgement carry only that student's name: ``--member REG`` builds one such copy and
``--all-members`` builds one per student in ``report/team.json``.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

REPORT_DIR = Path(__file__).resolve().parent
REPO_ROOT = REPORT_DIR.parent
if str(REPORT_DIR) not in sys.path:
    sys.path.insert(0, str(REPORT_DIR))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from lib import docx_helpers as dh  # noqa: E402
from lib import frontmatter as fm  # noqa: E402
from lib import libreoffice, vit_front  # noqa: E402
from lib.content import (  # noqa: E402
    Algorithm,
    Block,
    CodeBlock,
    Figure,
    GeneratedTable,
    Heading,
    ListBlock,
    PageBreak,
    Paragraph,
    ParseContext,
    Pending,
    PendingTable,
    Table,
    is_placeholder_file,
    parse_file,
    parse_markup,
    tokenize,
)
from lib.references import CitationRegistry, format_entry_apa, load_references  # noqa: E402
from lib.results_context import ResultsContext  # noqa: E402
from lib.tables import DEFAULT_CAPTIONS, generate  # noqa: E402

DEFAULT_OUT = REPORT_DIR / "MedBridge_Final_Report.docx"
DEFAULT_CONTENT = REPORT_DIR / "content"
DEFAULT_LOG = REPORT_DIR / "build" / "report_build_log.json"

#: Front-matter pages taken from content files. The cover, title, declaration, certificate
#: and acknowledgement pages come from the template (lib/vit_front.py, report/team.json).
FRONT_ORDER: tuple[str, ...] = ("abstract", "abbreviations")
#: Front-matter pages whose heading is listed in the table of contents.
FRONT_IN_TOC = frozenset({"abbreviations", "action_taken"})
#: Legacy content files whose pages the template now provides.
TEMPLATE_PAGES = frozenset({"title", "bonafide", "declaration", "acknowledgement"})
APPENDIX_HEADING = re.compile(r"^appendix\s+([A-Za-z]|\d+)\b\s*[:.\-–—]?\s*(.*)$", re.IGNORECASE)
CHAPTER_PREFIX = re.compile(r"^chapter\s+\d+\s*[:.\-–—]\s*", re.IGNORECASE)
#: Static pages placed after the list of abbreviations when no content file references them.
AUTO_FRONT_PAGES: tuple[tuple[str, str], ...] = (("action_taken", "action_taken.md"),)

#: Line spacing for preliminary pages that should fit on a single page.
COMPACT_FRONT: dict[str, float] = {"abstract": 1.0}  # single-spaced, as in the template sample
FRONT_TITLES: dict[str, str] = {
    "abstract": "Abstract",
    "abbreviations": "List of Symbols and Abbreviations",
    "action_taken": "Action Taken on Previous Review Comments",
}
FRONT_KEYWORDS: tuple[tuple[str, str], ...] = (
    ("title", r"title"),
    ("bonafide", r"bona\s*fide|certificate"),
    ("declaration", r"declaration"),
    ("acknowledgement", r"acknowledg"),
    ("abstract", r"abstract"),
    ("abbreviations", r"abbreviation|acronym|nomenclature|glossary"),
)
FRONT_EXACT = {
    "title",
    "title_page",
    "titlepage",
    "bonafide",
    "bonafide_certificate",
    "certificate",
    "declaration",
    "acknowledgement",
    "acknowledgements",
    "acknowledgment",
    "abstract",
    "abbreviations",
    "list_of_abbreviations",
}
EXPECTED_CHAPTERS = (1, 2, 3, 4, 5)
#: Body paragraphs open with this indent, as in the template's sample chapter page.
FIRST_LINE_INDENT_IN = 0.5
SEQ_KINDS = ("Figure", "Table", "Algorithm")


# ----------------------------------------------------------------------
# Content discovery
# ----------------------------------------------------------------------


@dataclass
class ContentPlan:
    front: dict[str, list[Block]] = field(default_factory=dict)
    front_source: dict[str, str] = field(default_factory=dict)
    chapters: list[tuple[str, list[Block]]] = field(default_factory=list)
    appendices: list[tuple[str, list[Block]]] = field(default_factory=list)
    include_only: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    auto_pages: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def classify(stem: str) -> tuple[str, Any]:
    """('front', slot) | ('chapter', (n, stem)) | ('appendix', stem) | ('skip', why) | ('include', None)."""
    s = stem.lower()
    m = re.match(r"^(?:front|pre|0\d)[_\-.\s]+(.+)$", s)
    if m or s in FRONT_EXACT:
        name = m.group(1) if m else s
        for slot, pattern in FRONT_KEYWORDS:
            if re.search(pattern, name):
                if slot == "abstract" and re.search(r"result|summary_", name):
                    break
                if slot in TEMPLATE_PAGES:
                    return "skip", "page generated from the VIT template and report/team.json"
                return "front", slot
    m = re.match(r"^(?:ch|chap|chapter)[_\-\s]?(\d+)", s)
    if m:
        return "chapter", (int(m.group(1)), s)
    if re.match(r"^(?:app|appendix|annex)(?:[_\-\s]?[a-z0-9]+)?", s) and not s.startswith("apple"):
        return "appendix", s
    if re.match(r"^(references|bibliography)\b", s):
        return "skip", "references are generated from report/references/*.json"
    return "include", None


def _referenced_names(content_dir: Path) -> set[str]:
    """File names pulled in with [[include:...]] and keys used as [[table:...]]."""
    names: set[str] = set()
    for path in content_dir.glob("*.md"):
        try:
            text = path.read_text(encoding="utf-8")
        except OSError:
            continue
        for kind, arg in re.findall(r"\[\[\s*(include|table)\s*:\s*([^\]|]+)", text):
            names.add(arg.strip() if kind == "include" else f"table:{arg.strip()}")
    return names


def discover(content_dir: Path) -> ContentPlan:
    plan = ContentPlan()
    pctx = ParseContext(content_dir=content_dir)
    files = sorted(
        p
        for p in content_dir.glob("*.md")
        if not p.name.startswith(("_", ".")) and p.name.lower() != "readme.md"
    )
    front_files: dict[str, list[Path]] = {}
    chapters: list[tuple[tuple[int, str], Path]] = []
    appendices: list[tuple[str, Path]] = []
    for path in files:
        kind, info = classify(path.stem)
        if kind == "front":
            front_files.setdefault(info, []).append(path)
        elif kind == "chapter":
            chapters.append((info, path))
        elif kind == "appendix":
            appendices.append((info, path))
        elif kind == "skip":
            plan.skipped.append(f"{path.name}: {info}")
        else:
            plan.include_only.append(path.name)

    def real_first(paths: list[Path]) -> list[Path]:
        real = [p for p in paths if not is_placeholder_file(p)]
        return real or paths[:1]

    for slot in FRONT_ORDER:
        paths = real_first(front_files.get(slot, []))
        if paths:
            blocks: list[Block] = []
            for path in paths:
                blocks.extend(parse_file(path, pctx))
            plan.front[slot] = blocks
            plan.front_source[slot] = ", ".join(p.name for p in paths)
        else:
            plan.front[slot] = parse_markup(fm.DEFAULTS[slot], pctx, source=f"default:{slot}")
            plan.front_source[slot] = "built-in default"

    by_number: dict[int, list[Path]] = {}
    for (number, _stem), path in sorted(chapters, key=lambda x: x[0]):
        by_number.setdefault(number, []).append(path)
    numbers = sorted(set(by_number) | set(EXPECTED_CHAPTERS))
    for number in numbers:
        paths = real_first(by_number.get(number, []))
        if not paths:
            plan.warnings.append(f"chapter {number}: no content file (placeholder chapter used)")
            text = fm.placeholder_chapter(number, f"report/content/ch{number}_*.md")
            plan.chapters.append(
                (
                    f"placeholder:ch{number}",
                    parse_markup(text, pctx, source=f"placeholder:ch{number}"),
                )
            )
            continue
        for path in paths:
            plan.chapters.append((path.name, parse_file(path, pctx)))
    groups: dict[str, list[Path]] = {}
    for stem, path in sorted(appendices):
        groups.setdefault(re.sub(r"\W", "", stem), []).append(path)
    for _key, paths in sorted(groups.items()):
        for path in real_first(paths):
            plan.appendices.append((path.name, parse_file(path, pctx)))
    referenced = _referenced_names(content_dir)
    for slot, name in AUTO_FRONT_PAGES:
        path = content_dir / name
        if path.is_file() and name not in referenced and f"table:{slot}" not in referenced:
            plan.front[slot] = parse_file(path, pctx)
            plan.front_source[slot] = f"{name} (auto: not referenced by any content file)"
            plan.auto_pages.append(slot)
    plan.warnings.extend(pctx.warnings)
    used = set(pctx.includes_used) | {
        name for slot, name in AUTO_FRONT_PAGES if slot in plan.auto_pages
    }
    plan.include_only = [n for n in plan.include_only if n not in used] + sorted(
        f"{n} (included)" for n in used if n in plan.include_only
    )
    return plan


# ----------------------------------------------------------------------
# Generated tables and citation numbering
# ----------------------------------------------------------------------


def expand_tables(blocks: list[Block], ctx: ResultsContext, errors: list[str]) -> list[Block]:
    out: list[Block] = []
    for block in blocks:
        if isinstance(block, GeneratedTable):
            table = generate(block.key, ctx, errors)
            caption = (
                block.caption
                or (table.caption if table else None)
                or DEFAULT_CAPTIONS.get(block.key)
            )
            if table is None:
                out.append(PendingTable(key=block.key, caption=caption))
            else:
                table.caption = caption
                out.append(table)
        else:
            out.append(block)
    return out


def texts_of(block: Block) -> list[str]:
    if isinstance(block, Heading | Paragraph):
        return [block.text]
    if isinstance(block, ListBlock):
        return [item.text for item in block.items]
    if isinstance(block, Figure):
        return [block.caption]
    if isinstance(block, Table):
        out = [block.caption or ""]
        if block.markup:
            out += block.header + [c for row in block.rows for c in row]
        return out
    if isinstance(block, PendingTable):
        return [block.caption or ""]
    if isinstance(block, Algorithm):
        return [block.title]
    return []


def register_citations(registry: CitationRegistry, blocks: list[Block]) -> None:
    for block in blocks:
        for text in texts_of(block):
            for tok in tokenize(text or ""):
                if tok.kind == "cite":
                    registry.register(tok.keys)


# ----------------------------------------------------------------------
# Renderer
# ----------------------------------------------------------------------


class Renderer:
    """Renders parsed blocks into a python-docx document with institutional formatting."""

    def __init__(
        self,
        doc: Any,
        ctx: ResultsContext,
        registry: CitationRegistry,
        content_dir: Path,
        repo_root: Path = REPO_ROOT,
    ) -> None:
        self.doc = doc
        self.ctx = ctx
        self.registry = registry
        self.content_dir = content_dir
        self.repo_root = repo_root
        self.part = "front"
        self.front_slot = ""
        self.chapter_no = 0
        self.appendix_no = 0
        self.label: str | None = None
        self.counters: Counter[str] = Counter()
        self.h2 = 0
        self.h3 = 0
        self.suppress_break = False
        self.stats: Counter[str] = Counter()
        self.pending: list[str] = []
        self.captions: dict[str, list[str]] = {k: [] for k in SEQ_KINDS}
        self.headings: list[str] = []

    # -- inline ------------------------------------------------------------

    def inline(
        self,
        para: Any,
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        size: float | None = None,
        markup: bool = True,
        upper: bool = False,
    ) -> None:
        def run(t: str, b: bool = False, i: bool = False, hl: bool = False) -> Any:
            r = para.add_run(t.upper() if upper and not hl else t)
            if bold or b:
                r.bold = True
            if italic or i:
                r.italic = True
            if size:
                r.font.size = Pt(size)
            if hl:
                dh.highlight(r)
                self.stats["highlights"] += 1
            return r

        if not markup:
            hl = (
                text.lstrip()
                .upper()
                .startswith(("[PENDING", "[TABLE PENDING", "[FIGURE PENDING", "[TO BE FILLED"))
            )
            run(text, hl=hl)
            return
        for tok in tokenize(text):
            if tok.kind == "text":
                run(tok.text, tok.bold, tok.italic)
            elif tok.kind == "code":
                code = tok.text
                if len(code) > 24:  # let long paths wrap instead of stretching justified lines
                    code = re.sub(r"([/._])(?=\S)", "\\1\u200b", code)
                r = run(code, tok.bold, tok.italic)
                dh.set_run_font(r, dh.MONO_FONT)
                r.font.size = Pt((size or 12) - 1.5)
            elif tok.kind == "cite":
                label, unknown = self.registry.label(tok.keys)
                if label:
                    run(label)
                for key in unknown:
                    run(f"[CITATION NOT FOUND: {key}]", hl=True)
                    self.pending.append(f"citation not found: {key}")
            elif tok.kind == "value":
                value, ok = self.ctx.render(tok.text)
                run(value, tok.bold, tok.italic, hl=not ok)
                if not ok:
                    self.pending.append(f"value: {tok.text}")
            elif tok.kind == "fill":
                run(tok.text, tok.bold, tok.italic, hl=True)
                if tok.text.upper().startswith("[TO BE FILLED"):
                    self.stats["to_be_filled"] += 1
                else:
                    self.pending.append(tok.text)
            elif tok.kind == "link":
                dh.add_hyperlink(
                    para, tok.text, tok.url, bold=tok.bold or bold, italic=tok.italic or italic
                )
            elif tok.kind == "br":
                para.add_run().add_break()

    def pending_para(self, text: str, *, align: Any = WD_ALIGN_PARAGRAPH.CENTER) -> Any:
        para = self.doc.add_paragraph()
        para.alignment = align
        r = para.add_run(text)
        r.bold = True
        dh.highlight(r)
        self.stats["highlights"] += 1
        self.pending.append(text)
        return para

    # -- headings ------------------------------------------------------------

    def _reset_chapter(self, label: str | None) -> None:
        self.label = label
        self.counters = Counter()
        self.h2 = self.h3 = 0

    def _heading_para(self, style: str) -> Any:
        para = self.doc.add_paragraph(style=style)
        if self.suppress_break:
            para.paragraph_format.page_break_before = False
            self.suppress_break = False
        return para

    def heading(self, h: Heading) -> None:
        text = h.text.strip()
        if self.part == "front":
            if h.level == 1:
                self.front_heading(
                    text,
                    in_toc=self.front_slot in FRONT_IN_TOC,
                    underline=self.front_slot == "abstract",
                )
            else:
                para = self.doc.add_paragraph()
                dh.keep_with_next(para)
                self.inline(para, text, bold=True)
            return
        if h.level == 1:
            explicit_app = APPENDIX_HEADING.match(text)
            if self.part == "main" and h.numbered:
                text = CHAPTER_PREFIX.sub("", text) or text  # "# Chapter 3: X" -> "X"
            if h.numbered and self.part == "main":
                self.chapter_no += 1
                self._reset_chapter(str(self.chapter_no))
                prefix = f"Chapter {self.chapter_no}"
            elif self.part == "appendix" and (h.numbered or explicit_app):
                self.appendix_no += 1
                number = str(self.appendix_no)
                if explicit_app:  # "#* Appendix 3: Test Cases" keeps the author's number
                    given = explicit_app.group(1)
                    number = given if given.isdigit() else str(ord(given.upper()) - ord("A") + 1)
                    text = explicit_app.group(2).strip() or text
                self._reset_chapter(f"A{number}")
                prefix = f"Appendix {number}"
            else:
                self._reset_chapter(None)
                prefix = ""
            # template: "Chapter n" in 12 pt bold above the title in 16 pt bold
            para = self._heading_para("Heading 1")
            if prefix:
                r = para.add_run(prefix)
                r.font.size = Pt(12)
                para.add_run().add_break()
            self.inline(para, text, upper=not prefix)
            self.headings.append(f"{prefix} {text}".strip() if prefix else text.upper())
            self.stats["chapters" if self.part == "main" else "appendices"] += 1
            return
        if h.level == 2:  # section: 14 pt, capitals
            number = ""
            if h.numbered and self.label:
                self.h2 += 1
                self.h3 = 0
                number = f"{self.label}.{self.h2}  "
            para = self._heading_para("Heading 2")
            self.inline(para, number + text, upper=True)
            self.headings.append(number + text.upper())
            return
        if h.level == 3:  # subsection: 12 pt, capitals
            number = ""
            if h.numbered and self.label and self.h2:
                self.h3 += 1
                number = f"{self.label}.{self.h2}.{self.h3}  "
            para = self._heading_para("Heading 3")
            self.inline(para, number + text, upper=True)
            self.headings.append(number + text.upper())
            return
        para = self._heading_para("Heading 4")
        self.inline(para, text)

    def front_heading(self, text: str, *, in_toc: bool, underline: bool = False) -> Any:
        """Heading of a preliminary page: centred (abstract, acknowledgement) or, for the
        contents and the lists, left-aligned as in the template, and then listed in the TOC."""
        para = self._heading_para("Heading 1" if in_toc else "Front Title")
        if in_toc:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self.headings.append(text.upper())
        r = para.add_run(text.upper())
        r.font.size = Pt(14)
        if underline:
            r.underline = True
        return para

    # -- captions ------------------------------------------------------------

    def caption(self, kind: str, text: str, *, keep_next: bool) -> Any:
        self.counters[kind] += 1
        n = self.counters[kind]
        para = self.doc.add_paragraph(style="Caption")
        prefix = f"{kind} {self.label}." if self.label else f"{kind} "
        r = para.add_run(prefix)
        r.bold = True
        # \s 1 = Word restarts at each Heading 1; \r 1 on the first caption of a chapter makes
        # LibreOffice (which ignores \s) restart too, so both agree with the computed number.
        instr = f"SEQ {kind} \\* ARABIC \\s 1" + (" \\r 1" if n == 1 else "")
        dh.add_field(para, instr, str(n), bold=True)
        r = para.add_run(": ")
        r.bold = True
        self.inline(para, text or "")
        dh.keep_with_next(para, keep_next)
        number = f"{self.label}.{n}" if self.label else str(n)
        self.captions[kind].append(f"{kind} {number}: {text}")
        return para

    # -- blocks ----------------------------------------------------------------

    def paragraph(self, p: Paragraph) -> None:
        if p.style == "quote":
            para = self.doc.add_paragraph(style="Quote")
        else:
            para = self.doc.add_paragraph()
        if self.part == "front" and self.front_slot in COMPACT_FRONT:
            # certificate-style pages are set tighter so each fits on one page
            para.paragraph_format.line_spacing = COMPACT_FRONT[self.front_slot]
            para.paragraph_format.space_after = Pt(6)
        if p.style == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p.style != "quote" and self.part != "front":
            para.paragraph_format.first_line_indent = Inches(FIRST_LINE_INDENT_IN)
        self.inline(para, p.text)

    def list_block(self, lb: ListBlock) -> None:
        num_ids: dict[int, int | None] = {}
        for item in lb.items:
            level = max(0, min(2, item.level))
            suffix = "" if level == 0 else f" {level + 1}"
            for deeper in [k for k in num_ids if k > level]:
                num_ids.pop(deeper)  # a nested list that starts again restarts at 1
            if item.ordered:
                style = f"List Number{suffix}"
                para = self.doc.add_paragraph(style=style)
                if level not in num_ids:
                    num_ids[level] = dh.restart_numbering(self.doc, style)
                num_id = num_ids.get(level)
                if num_id:
                    dh.set_paragraph_num(para, num_id, 0)
            else:
                para = self.doc.add_paragraph(style=f"List Bullet{suffix}")
            self.inline(para, item.text)

    def _resolve_figure(self, path_text: str) -> Path | None:
        raw = Path(path_text.strip())
        candidates = (
            [raw]
            if raw.is_absolute()
            else [self.repo_root / raw, self.content_dir / raw, self.repo_root / "report" / raw]
        )
        for cand in candidates:
            options = [cand]
            if cand.suffix.lower() in {".svg", ".pdf", ".eps"}:
                options = [cand.with_suffix(".png"), cand.with_suffix(".jpg")]
            for opt in options:
                if opt.is_file():
                    return opt
        return None

    def figure(self, f: Figure) -> None:
        path = self._resolve_figure(f.path)
        width = min(f.width or dh.TEXT_WIDTH_IN, dh.TEXT_WIDTH_IN)
        if path is None:
            para = self.pending_para(f"[FIGURE PENDING: {f.path}]")
            dh.keep_with_next(para)
            self.stats["figures_pending"] += 1
        else:
            try:
                from PIL import Image

                with Image.open(path) as img:
                    w_px, h_px = img.size
                max_h = 7.6
                if w_px and h_px and width * h_px / w_px > max_h:
                    width = max_h * w_px / h_px
            except Exception:  # size probing is best effort
                pass
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(2)
            dh.keep_with_next(para)
            try:
                para.add_run().add_picture(str(path), width=Inches(width))
                self.stats["figures"] += 1
            except Exception as exc:  # unsupported image format
                para.add_run(f"[FIGURE PENDING: {f.path} ({type(exc).__name__})]").bold = True
                para.runs[-1].font.highlight_color = WD_COLOR_INDEX.YELLOW
                self.pending.append(f"figure unreadable: {f.path}")
        self.caption(
            "Figure", f.caption or path.name if path else f.caption or f.path, keep_next=False
        )

    def table(self, t: Table) -> None:
        if not t.header and not t.rows:
            return
        show_caption = bool(t.caption) and self.part != "front"
        if show_caption:
            self.caption("Table", t.caption or "", keep_next=True)
        ncols = max(len(t.header), max((len(r) for r in t.rows), default=0))
        header = t.header + [""] * (ncols - len(t.header))
        rows = [r + [""] * (ncols - len(r)) for r in t.rows]
        table = self.doc.add_table(rows=1 + len(rows), cols=ncols)
        try:
            table.style = self.doc.styles["Table Grid"]
        except KeyError:
            pass
        size = t.font_size or (10 if ncols <= 5 else 9 if ncols <= 8 else 8)
        widths = dh.compute_column_widths(header, rows, weights=t.col_widths, font_size=size)
        dh.set_column_widths(table, widths)
        dh.set_table_borders(table)
        dh.set_cell_margins(table)
        align_map = {
            "l": WD_ALIGN_PARAGRAPH.LEFT,
            "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        aligns = (t.align + ["l"] * ncols)[:ncols]
        for c, text in enumerate(header):
            cell = table.rows[0].cells[c]
            dh.shade_cell(cell, dh.HEADER_FILL)
            para = cell.paragraphs[0]
            para.style = self.doc.styles["Table Text"]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            dh.keep_with_next(para)
            self.inline(para, text, bold=True, size=size, markup=t.markup)
        dh.mark_header_row(table.rows[0])
        dh.keep_row_together(table.rows[0])
        is_total = [
            bool(r and str(r[0]).strip().lower() in {"total", "overall"})
            or bool(len(r) > 1 and str(r[1]).strip().lower() == "total")
            for r in rows
        ]
        for r_idx, row in enumerate(rows, start=1):
            tr = table.rows[r_idx]
            dh.keep_row_together(tr)
            for c, text in enumerate(row):
                cell = tr.cells[c]
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Table Text"]
                para.alignment = align_map.get(aligns[c], WD_ALIGN_PARAGRAPH.LEFT)
                self.inline(para, str(text), size=size, markup=t.markup, bold=is_total[r_idx - 1])
        if t.note:
            note = self.doc.add_paragraph(style="Table Note")
            self.inline(note, t.note, markup=False)
        else:
            spacer = self.doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(6)
            spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        self.stats["tables"] += 1

    def pending_table(self, t: PendingTable) -> None:
        if t.caption and self.part != "front":
            self.caption("Table", t.caption, keep_next=True)
        self.pending_para(f"[TABLE PENDING: {t.key}]")
        self.stats["tables_pending"] += 1

    def boxed(self, lines: list[str], *, fill: str | None = None) -> None:
        # Courier is 0.6 em wide: shrink the font (9 -> 6.5 pt) until the longest line fits
        usable_pt = (dh.TEXT_WIDTH_IN - 0.2) * 72
        longest = max((len(line.rstrip()) for line in lines), default=1) or 1
        size = max(6.5, min(9.0, usable_pt / (0.6 * longest)))
        size = float(int(size * 2)) / 2  # half-point steps
        table = self.doc.add_table(rows=1, cols=1)
        dh.set_column_widths(table, [dh.TEXT_WIDTH_IN])
        dh.set_table_borders(table, size=8, color="404040")
        dh.set_cell_margins(table, top=80, bottom=80, left=140, right=140)
        cell = table.rows[0].cells[0]
        if fill:
            dh.shade_cell(cell, fill)
        dh.keep_row_together(table.rows[0])
        for i, line in enumerate(lines or [""]):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.style = self.doc.styles["Code Text"]
            dh.keep_with_next(para, i < len(lines) - 1)
            run = para.add_run(line if line else " ")
            dh.set_run_font(run, dh.MONO_FONT)
            run.font.size = Pt(size)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    def algorithm(self, a: Algorithm) -> None:
        self.caption("Algorithm", a.title, keep_next=True)
        self.boxed(a.lines)
        self.stats["algorithms"] += 1

    def render(self, blocks: list[Block]) -> None:
        for block in blocks:
            if isinstance(block, Heading):
                self.heading(block)
            elif isinstance(block, Paragraph):
                self.paragraph(block)
            elif isinstance(block, ListBlock):
                self.list_block(block)
            elif isinstance(block, Figure):
                self.figure(block)
            elif isinstance(block, Table):
                self.table(block)
            elif isinstance(block, PendingTable):
                self.pending_table(block)
            elif isinstance(block, Algorithm):
                self.algorithm(block)
            elif isinstance(block, CodeBlock):
                self.boxed(block.lines, fill=dh.CODE_FILL)
            elif isinstance(block, PageBreak):
                dh.add_page_break(self.doc.add_paragraph())
            elif isinstance(block, Pending):
                self.pending_para(block.text)
            elif isinstance(block, GeneratedTable):  # not expanded (should not happen)
                self.pending_table(PendingTable(key=block.key, caption=block.caption))

    # -- special pages -----------------------------------------------------------

    def front_page(self, slot: str, blocks: list[Block]) -> None:
        self.front_slot = slot
        if not any(isinstance(b, Heading) and b.level == 1 for b in blocks[:1]):
            title = FRONT_TITLES.get(slot, slot.title())
            self.heading(Heading(level=1, text=title, numbered=False))
        self.render(blocks)

    def index_page(self, title: str, instr: str, placeholder: str, *, in_toc: bool) -> None:
        self.front_heading(title, in_toc=in_toc)
        field = self.doc.add_paragraph()
        field.alignment = WD_ALIGN_PARAGRAPH.LEFT
        dh.add_index_field(field, instr, placeholder)

    def references(self) -> None:
        self._reset_chapter(None)
        para = self._heading_para("Heading 1")
        para.add_run("REFERENCES")
        self.headings.append("REFERENCES")
        bib = self.registry.bibliography()
        if not bib:
            self.pending_para("[PENDING: no citations in the content yet]")
            return
        for number, entry in bib:  # numbered list, each entry in APA style (template sample)
            p = self.doc.add_paragraph(style="Reference Entry")
            p.add_run(f"[{number}].\t")
            for seg in format_entry_apa(entry):
                r = p.add_run(seg.text)
                if seg.italic:
                    r.italic = True
        self.stats["references"] = len(bib)


# ----------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------


@dataclass
class BuildResult:
    docx: Path
    pdf: Path | None = None
    pdf_method: str = ""
    pdf_detail: str = ""
    pages: int | None = None
    log: dict[str, Any] = field(default_factory=dict)


def build(
    out: Path = DEFAULT_OUT,
    content_dir: Path = DEFAULT_CONTENT,
    *,
    repo_root: Path = REPO_ROOT,
    pdf: bool = False,
    log_path: Path | None = None,
    quiet: bool = False,
    member: str | None = None,
    team_file: Path | None = None,
) -> BuildResult:
    """Build one report. ``member`` (a register number from report/team.json) names the
    student on the declaration, certificate and acknowledgement; without it those pages
    keep a highlighted placeholder."""
    content_dir = content_dir.resolve()
    repo_root = repo_root.resolve()
    team = vit_front.load_team(repo_root, team_file)
    student = team.member(member) if member else None
    ctx = ResultsContext.load(repo_root)
    ctx.data["_content_dir"] = str(content_dir)  # static tables (action_taken) live here
    refs = load_references(repo_root)
    registry = CitationRegistry(refs)
    plan = discover(content_dir)

    table_errors: list[str] = []
    front = {slot: expand_tables(blocks, ctx, table_errors) for slot, blocks in plan.front.items()}
    chapters = [(name, expand_tables(blocks, ctx, table_errors)) for name, blocks in plan.chapters]
    appendices = [
        (name, expand_tables(blocks, ctx, table_errors)) for name, blocks in plan.appendices
    ]

    # citation numbers by first appearance, in final document order
    for slot in (*FRONT_ORDER, *plan.auto_pages):
        register_citations(registry, front[slot])
    for _, blocks in chapters:
        register_citations(registry, blocks)
    for _, blocks in appendices:
        register_citations(registry, blocks)

    doc = Document()
    dh.set_page_layout(doc.sections[0])
    dh.setup_styles(doc)
    r = Renderer(doc, ctx, registry, content_dir, repo_root)

    def count_fill(_text: str) -> None:
        r.stats["to_be_filled"] += 1
        r.stats["highlights"] += 1

    # --- template pages without page numbers -------------------------------
    r.part = "front"
    writer = vit_front.FrontWriter(doc, repo_root, count_fill)
    writer.cover_page(team)
    writer.title_page(team)
    writer.declaration(team, student)
    writer.certificate(team, student)

    # --- preliminary pages (lower-roman from the abstract) --------------------
    roman_section = dh.new_section(doc)
    r.suppress_break = True
    r.front_page("abstract", front["abstract"])
    r.front_slot = "acknowledgement"
    r.front_heading("Acknowledgement", in_toc=False)
    writer.acknowledgement(team, student)
    r.front_slot = "toc"
    r.index_page(
        "Contents",
        'TOC \\o "1-2" \\h \\z \\u',  # chapters and sections, as in the template sample
        "Right-click and choose Update Field (or press F9) to build the table of "
        "contents. Word refreshes it automatically when the document is opened.",
        in_toc=True,
    )
    r.index_page(
        "List of Tables",
        'TOC \\h \\z \\c "Table"',
        "Right-click and choose Update Field to build the list of tables.",
        in_toc=True,
    )
    r.index_page(
        "List of Figures",
        'TOC \\h \\z \\c "Figure"',
        "Right-click and choose Update Field to build the list of figures.",
        in_toc=True,
    )
    r.front_page("abbreviations", front["abbreviations"])
    for slot in plan.auto_pages:
        r.front_page(slot, front[slot])
    dh.set_page_numbering(roman_section, "lowerRoman", 1)

    # --- main matter (arabic, restarting at 1) ------------------------------
    main_section = dh.new_section(doc)
    dh.set_page_numbering(main_section, "decimal", 1)
    r.part = "main"
    r.suppress_break = True
    for _name, blocks in chapters:
        r.render(blocks)
    r.references()
    if appendices:
        r.part = "appendix"
        for _name, blocks in appendices:
            r.render(blocks)

    dh.clear_footer(doc.sections[0])  # cover, title, declaration and certificate
    dh.add_page_number_footer(doc.sections[1], italic=True)  # i, ii, ... as in the template
    dh.add_page_number_footer(doc.sections[2])
    dh.enable_update_fields(doc)
    dh.set_document_properties(
        doc,
        title=team.title,
        subject="MedBridge B.Tech project report (BCSE497J)",
        author=student.label if student else "; ".join(m.label for m in team.members),
        keywords="MedBridge; CardioGuard; RAG; hallucination; answerability; hypertension",
        comments="Generated by report/build_report.py from report/content and results files.",
    )
    out = out.resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    result = BuildResult(docx=out)
    if pdf:
        conv = libreoffice.convert_to_pdf(out)
        result.pdf, result.pdf_method, result.pdf_detail = conv.pdf, conv.method, conv.detail
        if conv.pdf:
            result.pages = libreoffice.pdf_page_count(conv.pdf)

    def rel(p: Path | None) -> str | None:
        if p is None:
            return None
        try:
            return p.resolve().relative_to(repo_root.resolve()).as_posix()
        except ValueError:
            return p.name

    log = {
        "generated_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "docx": rel(out),
        "pdf": rel(result.pdf),
        "pdf_method": result.pdf_method,
        "pdf_detail": result.pdf_detail,
        "pdf_pages": result.pages,
        "content_dir": rel(content_dir),
        "front_matter_sources": plan.front_source,
        "chapters": [name for name, _ in chapters],
        "appendices": [name for name, _ in appendices],
        "include_only_files": plan.include_only,
        "skipped_files": plan.skipped,
        "warnings": plan.warnings + table_errors,
        "stats": dict(r.stats),
        "citations_used": len(registry.order),
        "citation_keys_missing": registry.missing,
        "pending_markers": sorted(set(r.pending)),
        "pending_count": len(r.pending),
        "results_sources": ctx.data.get("_sources", {}),
        "results_skipped": ctx.data.get("_skipped", {}),
        "values_used": len(ctx.used),
        "values_missing": sorted(set(ctx.missing)),
        "captions": r.captions,
        "headings": r.headings,
    }
    result.log = log
    if log_path is not None:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text(json.dumps(log, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if not quiet:
        print_summary(result)
    return result


def print_summary(result: BuildResult) -> None:
    log = result.log
    stats = log["stats"]
    print(f"[report] wrote {log['docx']}")
    print(f"[report] chapters: {', '.join(log['chapters']) or '-'}")
    if log["appendices"]:
        print(f"[report] appendices: {', '.join(log['appendices'])}")
    print(f"[report] front matter: {log['front_matter_sources']}")
    print(
        f"[report] figures {stats.get('figures', 0)} (+{stats.get('figures_pending', 0)} pending)"
        f", tables {stats.get('tables', 0)} (+{stats.get('tables_pending', 0)} pending), "
        f"algorithms {stats.get('algorithms', 0)}, references {stats.get('references', 0)}"
    )
    print(
        f"[report] values resolved {log['values_used']}, pending {len(log['values_missing'])}; "
        f"[TO BE FILLED] placeholders {stats.get('to_be_filled', 0)}; "
        f"pending markers {log['pending_count']}"
    )
    if log["citation_keys_missing"]:
        print(f"[report] WARNING unknown citation keys: {', '.join(log['citation_keys_missing'])}")
    for warning in log["warnings"]:
        print(f"[report] warning: {warning}")
    for key, why in (log.get("results_skipped") or {}).items():
        print(f"[report] note: {key}: {why}")
    if log["include_only_files"]:
        print(f"[report] include-only files: {', '.join(log['include_only_files'])}")
    if result.pdf:
        print(
            f"[report] wrote {log['pdf']} ({log['pdf_pages']} pages, {log['pdf_method']}: "
            f"{log['pdf_detail']})"
        )
    elif log["pdf_method"]:
        print(
            f"[report] PDF not produced ({log['pdf_method']}: {log['pdf_detail']}); install "
            f"LibreOffice: {libreoffice.platform_hint()}"
        )


def member_out(out: Path, reg_no: str) -> Path:
    """``MedBridge_Final_Report.docx`` -> ``MedBridge_Final_Report_23BAI1160.docx``."""
    return out.with_name(f"{out.stem}_{reg_no}{out.suffix}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the MedBridge final report (DOCX/PDF)")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output .docx path")
    parser.add_argument("--content-dir", type=Path, default=DEFAULT_CONTENT)
    parser.add_argument("--pdf", action="store_true", help="also export a PDF with LibreOffice")
    parser.add_argument(
        "--log",
        type=Path,
        default=DEFAULT_LOG,
        help="JSON build log (pending markers, sources, warnings)",
    )
    parser.add_argument("--no-log", action="store_true")
    parser.add_argument(
        "--strict", action="store_true", help="exit 1 when unknown citation keys are used"
    )
    parser.add_argument("--quiet", action="store_true")
    parser.add_argument("--team-file", type=Path, default=None, help="default: report/team.json")
    who = parser.add_mutually_exclusive_group()
    who.add_argument(
        "--member",
        metavar="REG_NO",
        help="one student's copy (their name on the declaration, certificate and "
        "acknowledgement), written next to --out with the register number appended",
    )
    who.add_argument(
        "--all-members",
        action="store_true",
        help="the team copy plus one copy per student in report/team.json",
    )
    args = parser.parse_args(argv)
    if not args.content_dir.is_dir():
        print(f"[report] content directory not found: {args.content_dir}", file=sys.stderr)
        return 2
    try:
        team = vit_front.load_team(REPO_ROOT, args.team_file)
        if args.member:
            team.member(args.member)
    except (OSError, ValueError, KeyError) as exc:
        print(f"[report] team file problem: {exc}", file=sys.stderr)
        return 2
    jobs: list[tuple[Path, str | None]] = []
    if args.member:
        jobs.append((member_out(args.out, args.member), args.member))
    else:
        jobs.append((args.out, None))
        if args.all_members:
            jobs += [(member_out(args.out, m.reg_no), m.reg_no) for m in team.members]
    code = 0
    for i, (out, reg_no) in enumerate(jobs):
        result = build(
            out,
            args.content_dir,
            pdf=args.pdf,
            log_path=None if args.no_log or i > 0 else args.log,
            quiet=args.quiet,
            member=reg_no,
            team_file=args.team_file,
        )
        if args.strict and result.log["citation_keys_missing"]:
            code = 1
        if args.pdf and result.pdf is None:
            code = code or 3
    return code


if __name__ == "__main__":
    raise SystemExit(main())
