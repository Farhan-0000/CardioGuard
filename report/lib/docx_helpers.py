"""
Low-level python-docx helpers: page setup, institutional styles, Word fields
(TOC / SEQ / PAGE), section footers with roman/arabic page numbers, table
formatting, highlights, hyperlinks and restartable numbered lists.
"""

from __future__ import annotations

import copy
import math
import re
from typing import Any

from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_COLOR_INDEX,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph
from docx.text.run import Run

# A4 portrait, 1 inch margins with a 1.5 inch binding margin on the left.
PAGE_WIDTH_IN = 8.27
PAGE_HEIGHT_IN = 11.69
MARGIN_IN = 1.0
LEFT_MARGIN_IN = 1.5
TEXT_WIDTH_IN = PAGE_WIDTH_IN - LEFT_MARGIN_IN - MARGIN_IN  # 5.77
TEXT_HEIGHT_IN = PAGE_HEIGHT_IN - 2 * MARGIN_IN

BODY_FONT = "Times New Roman"
MONO_FONT = "Courier New"
HEADER_FILL = "D9E2F3"
CODE_FILL = "F2F2F2"


#: Schema order of the children we touch (Word rejects out-of-order elements).
TBLPR_ORDER = (
    "tblStyle",
    "tblpPr",
    "tblOverlap",
    "bidiVisual",
    "tblStyleRowBandSize",
    "tblStyleColBandSize",
    "tblW",
    "jc",
    "tblCellSpacing",
    "tblInd",
    "tblBorders",
    "shd",
    "tblLayout",
    "tblCellMar",
    "tblLook",
    "tblCaption",
    "tblDescription",
)
PPR_ORDER = (
    "pStyle",
    "keepNext",
    "keepLines",
    "pageBreakBefore",
    "framePr",
    "widowControl",
    "numPr",
    "suppressLineNumbers",
    "pBdr",
    "shd",
    "tabs",
    "suppressAutoHyphens",
    "kinsoku",
    "wordWrap",
    "overflowPunct",
    "topLinePunct",
    "autoSpaceDE",
    "autoSpaceDN",
    "bidi",
    "adjustRightInd",
    "snapToGrid",
    "spacing",
    "ind",
    "contextualSpacing",
    "mirrorIndents",
    "suppressOverlap",
    "jc",
    "textDirection",
    "textAlignment",
    "textboxTightWrap",
    "outlineLvl",
    "divId",
    "cnfStyle",
    "rPr",
    "sectPr",
    "pPrChange",
)
TCPR_ORDER = (
    "cnfStyle",
    "tcW",
    "gridSpan",
    "hMerge",
    "vMerge",
    "tcBorders",
    "shd",
    "noWrap",
    "tcMar",
    "textDirection",
    "tcFitText",
    "vAlign",
    "hideMark",
)
SETTINGS_AFTER_UPDATEFIELDS = (
    "hdrShapeDefaults",
    "footnotePr",
    "endnotePr",
    "compat",
    "docVars",
    "rsids",
    "mathPr",
    "attachedSchema",
    "themeFontLang",
    "clrSchemeMapping",
    "doNotIncludeSubdocsInStats",
    "doNotAutoCompressPictures",
    "forceUpgrade",
    "captions",
    "readModeInkLockDown",
    "smartTagType",
    "schemaLibrary",
    "shapeDefaults",
    "doNotEmbedSmartTags",
    "decimalSymbol",
    "listSeparator",
)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def get_or_add_ordered(parent: Any, name: str, order: tuple[str, ...]) -> Any:
    """Child ``w:name`` of ``parent``, created at its schema position if missing."""
    existing = parent.find(qn(f"w:{name}"))
    if existing is not None:
        return existing
    child = OxmlElement(f"w:{name}")
    later = set(order[order.index(name) + 1 :]) if name in order else set()
    for sibling in parent:
        if _local(sibling.tag) in later:
            sibling.addprevious(child)
            return child
    parent.append(child)
    return child


# ----------------------------------------------------------------------
# Fonts and styles
# ----------------------------------------------------------------------


def set_rfonts(rpr: Any, name: str) -> None:
    """Force a font on every script slot and drop theme-font overrides."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def set_run_font(run: Run, name: str) -> None:
    run.font.name = name
    set_rfonts(run._element.get_or_add_rPr(), name)


def _style(doc: DocumentObject, name: str, kind: WD_STYLE_TYPE = WD_STYLE_TYPE.PARAGRAPH) -> Any:
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, kind)


def _format_style(
    style: Any,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    align: Any = None,
    before: float | None = None,
    after: float | None = None,
    spacing: float | None = None,
    keep_next: bool | None = None,
    page_break: bool | None = None,
    font: str = BODY_FONT,
    color: RGBColor | None = RGBColor(0, 0, 0),
    base: str | None = None,
) -> None:
    if base is not None:
        style.base_style = style.part.styles[base] if hasattr(style, "part") else None
    style.font.name = font
    set_rfonts(style.element.get_or_add_rPr(), font)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if spacing is not None:
        pf.line_spacing = spacing
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if page_break is not None:
        pf.page_break_before = page_break


def setup_styles(doc: DocumentObject) -> None:
    """Institutional template: TNR 12 pt, 1.5 spacing, justified; headings 16/14/12 pt."""
    normal = doc.styles["Normal"]
    _format_style(normal, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5)
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    _format_style(
        h1,
        size=16,
        bold=True,
        italic=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=24,
        spacing=1.5,
        keep_next=True,
        page_break=True,
    )
    h2 = doc.styles["Heading 2"]
    _format_style(
        h2,
        size=14,
        bold=True,
        italic=False,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        before=18,
        after=6,
        spacing=1.5,
        keep_next=True,
        page_break=False,
    )
    h3 = doc.styles["Heading 3"]
    _format_style(
        h3,
        size=12,
        bold=True,
        italic=False,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        before=12,
        after=6,
        spacing=1.5,
        keep_next=True,
        page_break=False,
    )
    h4 = doc.styles["Heading 4"]
    _format_style(
        h4,
        size=12,
        bold=True,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        before=6,
        after=4,
        spacing=1.5,
        keep_next=True,
        page_break=False,
    )
    for h in (h1, h2, h3, h4):
        h.paragraph_format.left_indent = Pt(0)

    front = _style(doc, "Front Title")
    front.base_style = normal
    _format_style(
        front,
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=24,
        spacing=1.5,
        keep_next=True,
        page_break=True,
    )

    title = _style(doc, "Report Title")
    title.base_style = normal
    _format_style(title, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=6, spacing=1.15)

    caption = doc.styles["Caption"]
    _format_style(
        caption,
        size=11,
        bold=False,
        italic=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=6,
        after=12,
        spacing=1.0,
    )

    table_text = _style(doc, "Table Text")
    table_text.base_style = normal
    _format_style(
        table_text, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, spacing=1.0
    )

    note = _style(doc, "Table Note")
    note.base_style = normal
    _format_style(
        note, size=9, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=12, spacing=1.0
    )

    code = _style(doc, "Code Text")
    code.base_style = normal
    _format_style(
        code, size=9, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=0, spacing=1.0, font=MONO_FONT
    )

    quote = doc.styles["Quote"]
    _format_style(
        quote,
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        before=4,
        after=8,
        spacing=1.5,
    )
    quote.paragraph_format.left_indent = Inches(0.5)
    quote.paragraph_format.right_indent = Inches(0.5)

    ref = _style(doc, "Reference Entry")
    ref.base_style = normal
    _format_style(ref, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, spacing=1.15)
    ref.paragraph_format.left_indent = Inches(0.5)
    ref.paragraph_format.first_line_indent = Inches(-0.5)
    ref.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))

    for name in (
        "List Bullet",
        "List Bullet 2",
        "List Bullet 3",
        "List Number",
        "List Number 2",
        "List Number 3",
    ):
        st = doc.styles[name]
        _format_style(st, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3, spacing=1.5)

    # Index entry styles (Word "toc 1".."toc 3" and "table of figures").
    for level, name in enumerate(("toc 1", "toc 2", "toc 3"), start=1):
        st = _style(doc, name)
        st.base_style = normal
        _format_style(
            st,
            size=12,
            bold=level == 1,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            before=0,
            after=2 if level > 1 else 4,
            spacing=1.15,
        )
        st.paragraph_format.left_indent = Inches(0.3 * (level - 1))
        st.paragraph_format.tab_stops.add_tab_stop(
            Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
    tof = _style(doc, "table of figures")
    tof.base_style = normal
    _format_style(tof, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=2, spacing=1.15)
    tof.paragraph_format.tab_stops.add_tab_stop(
        Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    footer = doc.styles["Footer"]
    _format_style(footer, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, spacing=1.0)


# ----------------------------------------------------------------------
# Page layout, sections, footers
# ----------------------------------------------------------------------


def set_page_layout(section: Any) -> None:
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(LEFT_MARGIN_IN)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.gutter = Inches(0)


def set_page_numbering(section: Any, fmt: str, start: int | None = 1) -> None:
    """``fmt``: decimal | lowerRoman | upperRoman; restarts at ``start``."""
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        # schema order: ... pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols ...
        cols = sect_pr.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(pg)
        else:
            sect_pr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def new_section(doc: DocumentObject) -> Any:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_layout(section)
    return section


def add_page_number_footer(section: Any, *, blank_first_page: bool = False) -> None:
    """Centered PAGE field in the section's own (unlinked) footer."""
    section.different_first_page_header_footer = blank_first_page
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.style = (
        section.part.document.styles["Footer"] if hasattr(section.part, "document") else para.style
    )
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(para.runs):
        run._element.getparent().remove(run._element)
    add_field(para, "PAGE", "1")
    if blank_first_page:
        first = section.first_page_footer
        first.is_linked_to_previous = False
        for p in first.paragraphs:
            for run in list(p.runs):
                run._element.getparent().remove(run._element)


def enable_update_fields(doc: DocumentObject) -> None:
    """Ask Word to refresh TOC / LoF / LoT / SEQ / PAGE fields when the file is opened."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        for sibling in settings:
            if _local(sibling.tag) in SETTINGS_AFTER_UPDATEFIELDS:
                sibling.addprevious(el)
                break
        else:
            settings.append(el)
    el.set(qn("w:val"), "true")


def set_document_properties(doc: DocumentObject, **props: str) -> None:
    core = doc.core_properties
    for key, value in props.items():
        if hasattr(core, key):
            setattr(core, key, value)


# ----------------------------------------------------------------------
# Fields
# ----------------------------------------------------------------------


def _fld_char(kind: str, dirty: bool = False) -> Any:
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    if dirty:
        el.set(qn("w:dirty"), "true")
    return el


def add_field(
    paragraph: Paragraph,
    instr: str,
    result: str = "",
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    size: float | None = None,
    highlight: bool = False,
    dirty: bool = False,
) -> None:
    """Complex field ``{ instr }`` with a cached ``result`` (shown until fields update)."""

    def _run() -> Run:
        run = paragraph.add_run()
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if size is not None:
            run.font.size = Pt(size)
        return run

    _run()._element.append(_fld_char("begin", dirty))
    instr_run = _run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    instr_run._element.append(instr_el)
    _run()._element.append(_fld_char("separate"))
    res = _run()
    res.text = result
    if highlight:
        res.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _run()._element.append(_fld_char("end"))


def add_index_field(paragraph: Paragraph, instr: str, placeholder: str) -> None:
    """TOC-style field marked dirty so Word rebuilds it on open."""
    add_field(paragraph, instr, placeholder, dirty=True, italic=True)


# ----------------------------------------------------------------------
# Runs, highlights, hyperlinks
# ----------------------------------------------------------------------


def highlight(run: Run, color: WD_COLOR_INDEX = WD_COLOR_INDEX.YELLOW) -> Run:
    run.font.highlight_color = color
    return run


def add_hyperlink(
    paragraph: Paragraph, text: str, url: str, *, bold: bool = False, italic: bool = False
) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if bold:  # rPr order: b, i, ..., color, ..., u
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_page_break(paragraph: Paragraph) -> None:
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def keep_with_next(paragraph: Paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


# ----------------------------------------------------------------------
# Tables
# ----------------------------------------------------------------------


def shade_cell(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = get_or_add_ordered(tc_pr, "shd", TCPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_table_borders(table: Table, size: int = 4, color: str = "808080") -> None:
    tbl_pr = table._tbl.tblPr
    borders = get_or_add_ordered(tbl_pr, "tblBorders", TBLPR_ORDER)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(
    table: Table, top: int = 40, bottom: int = 40, left: int = 80, right: int = 80
) -> None:
    tbl_pr = table._tbl.tblPr
    mar = get_or_add_ordered(tbl_pr, "tblCellMar", TBLPR_ORDER)
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def mark_header_row(row: _Row) -> None:
    """Repeat the row at the top of every page the table spans."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def keep_row_together(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_column_widths(table: Table, widths_in: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for col, width in zip(cols, widths_in, strict=False):
        col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
    tbl_pr = table._tbl.tblPr
    tbl_w = get_or_add_ordered(tbl_pr, "tblW", TBLPR_ORDER)
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = get_or_add_ordered(tbl_pr, "tblLayout", TBLPR_ORDER)
    layout.set(qn("w:type"), "fixed")


def compute_column_widths(
    header: list[str],
    rows: list[list[str]],
    total: float = TEXT_WIDTH_IN,
    weights: list[float] | None = None,
    font_size: float = 10.0,
    padding: float = 0.14,
) -> list[float]:
    """Column widths: proportional to content (or ``weights``) but never narrower than the
    longest word of the column, so header words such as "Chunks" are not split."""
    n = len(header)
    if n == 0:
        return []
    char_in = font_size * 0.55 / 72.0

    def words(c: int) -> list[str]:
        texts = [header[c]] + [r[c] for r in rows if c < len(r)]
        # split on ordinary whitespace only: non-breaking spaces keep composites together
        return [w for t in texts for w in re.split(r"[ \t\r\n]+", str(t).replace("**", "")) if w]

    mins = []
    for c in range(n):
        # very long tokens (paths, URLs) wrap anyway: they must not starve other columns
        longest = min(max((len(w) for w in words(c)), default=1), 22)
        mins.append(min(longest * char_in + padding, total * 0.45))
    if weights and len(weights) == n:
        raw = [max(float(w), 0.05) for w in weights]
    else:
        raw = []
        for c in range(n):
            texts = [header[c]] + [r[c] for r in rows if c < len(r)]
            avg = sum(len(str(t)) for t in texts) / max(len(texts), 1)
            raw.append(math.sqrt(max(min(avg, 400.0), 3.0)))  # damped: long prose gets more
    widths = [w * total / sum(raw) for w in raw]
    for _ in range(12):
        deficit = sum(max(m - w, 0.0) for w, m in zip(widths, mins, strict=True))
        if deficit <= 1e-6:
            break
        widths = [max(w, m) for w, m in zip(widths, mins, strict=True)]
        slack_idx = [i for i in range(n) if widths[i] > mins[i] + 1e-9]
        slack = sum(widths[i] - mins[i] for i in slack_idx)
        if slack <= 1e-9:
            break
        take = min(deficit, slack)
        for i in slack_idx:
            widths[i] -= take * (widths[i] - mins[i]) / slack
    current = sum(widths)
    if current > total:
        widths = [w * total / current for w in widths]
    return widths


# ----------------------------------------------------------------------
# Numbered lists that restart at 1
# ----------------------------------------------------------------------


def restart_numbering(doc: DocumentObject, style_name: str = "List Number") -> int | None:
    """New ``w:num`` (same abstract definition as ``style_name``) starting again at 1."""
    try:
        style = doc.styles[style_name]
    except KeyError:
        return None
    num_pr = style.element.pPr.find(qn("w:numPr")) if style.element.pPr is not None else None
    if num_pr is None or num_pr.find(qn("w:numId")) is None:
        return None
    base_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == base_id:
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            break
    if abstract_id is None:
        return None
    existing = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = max(existing) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abs_el = OxmlElement("w:abstractNumId")
    abs_el.set(qn("w:val"), abstract_id)
    num.append(abs_el)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    nums = numbering.findall(qn("w:num"))
    if nums:  # w:num elements follow the abstractNums and precede numIdMacAtCleanup
        nums[-1].addnext(num)
    else:
        numbering.append(num)
    return new_id


def set_paragraph_num(paragraph: Paragraph, num_id: int, ilvl: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = get_or_add_ordered(p_pr, "numPr", PPR_ORDER)
    for child in list(num_pr):
        num_pr.remove(child)
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.append(lvl)
    num_pr.append(nid)


def clone_element(el: Any) -> Any:
    return copy.deepcopy(el)


def single_spacing(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
