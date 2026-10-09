"""
Human annotation subset (CLAUDE_AGENT_PROMPT §6.3.3).

``--export`` samples 50 items (seeded, stratified by category, proportional
allocation with largest remainders) among test items whose C5 or C1 response
contains at least one claim, and writes every claim of those responses with
its evidence to:

* ``evaluation/annotation/human_annotation_sheet.xlsx`` — an *Instructions*
  sheet (the rubric) and a *Claims* sheet with dropdown validation
  (SUPPORTED / UNSUPPORTED / CONTRADICTED) in the annotator columns A1 and A2;
* ``human_annotation_sheet.csv`` — the same rows;
* ``annotation_key.csv`` — the hidden mapping row → condition and judge labels
  (do not show it to annotators: responses are blinded by a response code).

``--score`` reads the filled sheet (xlsx or csv) and writes
``evaluation/results/human_agreement.json`` with human–human and human–judge
Cohen's κ (3-class and binary). An unfilled sheet yields a ``not_filled``
placeholder that the report shows as such.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import random
from collections import defaultdict
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from evaluation import metrics as mx
from evaluation.datasets.schema import CATEGORIES
from evaluation.io import iter_jsonl, load_all_judgments, load_run_records, write_json
from evaluation.judges.common import CLAIM_LABELS, EVIDENCE_SIDECAR, SUPPORTED
from evaluation.judges.nli_judge import load_thresholds
from evaluation.manifest import utc_now
from evaluation.outcomes import TurnOutcome, build_outcomes, filter_split
from evaluation.paths import (
    ANNOTATION_DIR,
    CALIBRATION_PATH,
    HUMAN_AGREEMENT_PATH,
    JUDGMENTS_DIR,
    RUNS_DIR,
)

SHEET_NAME = "human_annotation_sheet"
KEY_NAME = "annotation_key.csv"
COLUMNS: tuple[str, ...] = (
    "row_id",
    "response_code",
    "item_id",
    "turn",
    "category",
    "user_question",
    "response",
    "claim_index",
    "claim",
    "evidence",
    "A1_label",
    "A2_label",
    "A1_notes",
    "A2_notes",
)
ANNOTATORS: tuple[str, ...] = ("A1", "A2")
EVIDENCE_CHARS = 1500

INSTRUCTIONS: tuple[str, ...] = (
    "MedBridge-HalluBench — human claim annotation",
    "",
    "Each row of the 'Claims' sheet is ONE atomic clinical claim taken from an assistant response, "
    "with the guideline passages (evidence) the automatic judges used. Responses are anonymised "
    "(response_code); do not try to infer which system produced them.",
    "",
    "For every row, choose a label in your column (A1_label for annotator 1, A2_label for annotator 2):",
    "SUPPORTED — the evidence explicitly states the claim, or the claim is a direct paraphrase of it "
    "(same meaning, same numbers/units, same population and conditions).",
    "CONTRADICTED — the evidence explicitly states something incompatible with the claim "
    "(different threshold, dose, direction of effect, a contraindication).",
    "UNSUPPORTED — anything else: plausible or generally true but not stated in the evidence; adds "
    "details (numbers, doses, drug names, populations, time frames) the evidence does not contain; "
    "only partly supported.",
    "",
    "Rules: use ONLY the evidence column — not your own medical knowledge. When in doubt between "
    "SUPPORTED and UNSUPPORTED choose UNSUPPORTED. Label independently (do not discuss with the "
    "other annotator until both columns are complete). Optional notes go in A1_notes / A2_notes.",
    "",
    "Do not open annotation_key.csv: it reveals the system and the automatic labels.",
    "When finished, run: python -m evaluation.human_annotation --score",
)


def stratified_sample(item_categories: dict[str, str], n: int, seed: int = 42) -> list[str]:
    """Proportional stratified sample of item ids (largest-remainder allocation, seeded)."""
    by_cat: dict[str, list[str]] = defaultdict(list)
    for item_id, category in sorted(item_categories.items()):
        by_cat[category].append(item_id)
    total = len(item_categories)
    if n >= total:
        return sorted(item_categories)
    quotas = {c: n * len(ids) / total for c, ids in by_cat.items()}
    alloc = {c: int(q) for c, q in quotas.items()}
    remainder = n - sum(alloc.values())
    for c in sorted(quotas, key=lambda c: (-(quotas[c] - alloc[c]), c))[:remainder]:
        alloc[c] += 1
    rng = random.Random(seed)
    chosen: list[str] = []
    for category in sorted(by_cat, key=lambda c: CATEGORIES.index(c) if c in CATEGORIES else 99):
        ids = list(by_cat[category])
        rng.shuffle(ids)
        chosen.extend(ids[: alloc[category]])
    return sorted(chosen)


def _response_code(condition: str, item_id: str, turn: int, seed: int) -> str:
    digest = hashlib.sha256(f"{seed}|{condition}|{item_id}|{turn}".encode()).hexdigest()
    return "R" + digest[:8].upper()


def build_rows(
    outcomes: dict[str, list[TurnOutcome]],
    evidence_lookup: dict[str, dict[str, str]],
    sidecar: dict[str, str],
    n_items: int = 50,
    seed: int = 42,
    conditions: Sequence[str] = ("C5", "C1"),
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """(sheet rows, key rows) for the stratified sample."""
    candidates: dict[str, str] = {}
    for c in conditions:
        for o in outcomes.get(c, []):
            if o.valid and o.claims.get("llm"):
                candidates[o.item_id] = o.category
    sample = set(stratified_sample(candidates, n_items, seed))
    rows: list[dict[str, Any]] = []
    keys: list[dict[str, Any]] = []
    for c in conditions:
        for o in sorted(outcomes.get(c, []), key=lambda x: x.key):
            if o.item_id not in sample or not o.valid or not o.claims.get("llm"):
                continue
            code = _response_code(c, o.item_id, o.turn_index, seed)
            nli = {x.get("index"): x.get("label") for x in o.claims.get("nli") or []}
            texts = evidence_lookup.get(f"{c}|{o.item_id}|{o.turn_index}", {})
            for claim in o.claims["llm"] or []:
                evidence = []
                for i, chunk_id in enumerate(claim.get("evidence_ids") or [], start=1):
                    text = texts.get(chunk_id) or sidecar.get(chunk_id) or ""
                    evidence.append(f"[E{i} {chunk_id}] {text[:EVIDENCE_CHARS]}")
                row_id = f"{code}-{claim.get('index')}"
                rows.append(
                    {
                        "row_id": row_id,
                        "response_code": code,
                        "item_id": o.item_id,
                        "turn": o.turn_index,
                        "category": o.category,
                        "user_question": o.user_text,
                        "response": o.response_text,
                        "claim_index": claim.get("index"),
                        "claim": claim.get("claim"),
                        "evidence": "\n\n".join(evidence),
                        "A1_label": "",
                        "A2_label": "",
                        "A1_notes": "",
                        "A2_notes": "",
                    }
                )
                keys.append(
                    {
                        "row_id": row_id,
                        "condition": c,
                        "item_id": o.item_id,
                        "turn": o.turn_index,
                        "llm_label": claim.get("label"),
                        "nli_label": nli.get(claim.get("index"), ""),
                    }
                )
    rows.sort(key=lambda r: (r["response_code"], r["claim_index"]))
    return rows, keys


def write_sheet(
    rows: Sequence[dict[str, Any]], keys: Sequence[dict[str, Any]], out_dir: Path
) -> list[Path]:
    """Write the xlsx (Instructions + Claims with dropdowns), the csv and the hidden key."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = ins.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 120

    sheet = wb.create_sheet("Claims")
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F5")
    for row in rows:
        sheet.append([row.get(c, "") for c in COLUMNS])
    widths = {
        "user_question": 40,
        "response": 60,
        "claim": 50,
        "evidence": 90,
        "A1_notes": 25,
        "A2_notes": 25,
    }
    for idx, name in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=idx).column_letter
        sheet.column_dimensions[letter].width = widths.get(name, 14)
        if name in widths:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(CLAIM_LABELS) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Invalid label",
        error="Choose SUPPORTED, UNSUPPORTED or CONTRADICTED",
    )
    sheet.add_data_validation(validation)
    last = max(2, len(rows) + 1)
    for name in ("A1_label", "A2_label"):
        letter = sheet.cell(row=1, column=COLUMNS.index(name) + 1).column_letter
        validation.add(f"{letter}2:{letter}{last}")
    sheet.freeze_panes = "B2"
    xlsx = out_dir / f"{SHEET_NAME}.xlsx"
    wb.save(xlsx)

    csv_path = out_dir / f"{SHEET_NAME}.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(COLUMNS))
        writer.writeheader()
        writer.writerows(rows)
    key_path = out_dir / KEY_NAME
    with key_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(
            fh, fieldnames=["row_id", "condition", "item_id", "turn", "llm_label", "nli_label"]
        )
        writer.writeheader()
        writer.writerows(keys)
    return [xlsx, csv_path, key_path]


def read_filled(path: Path) -> list[dict[str, Any]]:
    """Rows of a filled sheet (xlsx 'Claims' sheet or csv)."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8") as fh:
            return list(csv.DictReader(fh))
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True)["Claims"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, r, strict=False)) for r in rows[1:]]


def _norm(label: Any) -> str | None:
    text = str(label or "").strip().upper()
    return text if text in CLAIM_LABELS else None


def _agreement(a: Sequence[str], b: Sequence[str]) -> dict[str, Any]:
    if not a:
        return {"n": 0}
    return {
        "n": len(a),
        "kappa_3class": mx.cohen_kappa(list(a), list(b), labels=CLAIM_LABELS).as_dict(),
        "kappa_binary": mx.cohen_kappa(
            [str(x == SUPPORTED) for x in a],
            [str(x == SUPPORTED) for x in b],
            labels=["True", "False"],
        ).as_dict(),
        "raw_agreement": sum(x == y for x, y in zip(a, b, strict=True)) / len(a),
    }


def score(sheet_path: Path, key_path: Path) -> dict[str, Any]:
    """Human–human and human–judge agreement (placeholder when the sheet is unfilled)."""
    rows = read_filled(sheet_path)
    keys = {}
    with key_path.open(encoding="utf-8") as fh:
        for k in csv.DictReader(fh):
            keys[k["row_id"]] = k
    labelled = {a: {str(r["row_id"]): _norm(r.get(f"{a}_label")) for r in rows} for a in ANNOTATORS}
    counts = {a: sum(1 for v in labelled[a].values() if v) for a in ANNOTATORS}
    result: dict[str, Any] = {
        "generated_at": utc_now(),
        "sheet": str(sheet_path),
        "n_rows": len(rows),
        "labelled": counts,
    }
    if not any(counts.values()):
        result["status"] = "not_filled"
        result["note"] = (
            "Human annotation sheet not filled yet; agreement will be computed with --score."
        )
        return result
    result["status"] = "complete" if all(c == len(rows) for c in counts.values()) else "partial"
    both = [rid for rid in labelled["A1"] if labelled["A1"][rid] and labelled["A2"].get(rid)]
    result["human_human"] = _agreement(
        [labelled["A1"][r] or "" for r in both], [labelled["A2"][r] or "" for r in both]
    )
    for judge in ("llm", "nli"):
        for a in ANNOTATORS:
            ids = [
                r
                for r, v in labelled[a].items()
                if v and _norm(keys.get(r, {}).get(f"{judge}_label"))
            ]
            result[f"{a}_vs_{judge}"] = _agreement(
                [labelled[a][r] or "" for r in ids],
                [_norm(keys[r][f"{judge}_label"]) or "" for r in ids],
            )
        consensus = [
            r
            for r in both
            if labelled["A1"][r] == labelled["A2"][r]
            and _norm(keys.get(r, {}).get(f"{judge}_label"))
        ]
        result[f"consensus_vs_{judge}"] = _agreement(
            [labelled["A1"][r] or "" for r in consensus],
            [_norm(keys[r][f"{judge}_label"]) or "" for r in consensus],
        )
    return result


def _evidence_lookup(runs: dict[str, list[dict[str, Any]]]) -> dict[str, dict[str, str]]:
    return {
        f"{c}|{r['item_id']}|{r['turn_index']}": {
            str(e["chunk_id"]): str(e.get("text", "")) for e in r.get("evidence_texts") or []
        }
        for c, recs in runs.items()
        for r in recs
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export / score the human annotation sheet")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--export", action="store_true")
    mode.add_argument("--score", action="store_true")
    parser.add_argument("--runs-dir", type=Path, default=RUNS_DIR)
    parser.add_argument("--judgments-dir", type=Path, default=JUDGMENTS_DIR)
    parser.add_argument("--calibration", type=Path, default=CALIBRATION_PATH)
    parser.add_argument("--out-dir", type=Path, default=ANNOTATION_DIR)
    parser.add_argument(
        "--sheet", type=Path, default=None, help="filled sheet (default: the exported xlsx)"
    )
    parser.add_argument("--output", type=Path, default=HUMAN_AGREEMENT_PATH)
    parser.add_argument("--n-items", type=int, default=50)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--split", choices=["dev", "test", "all"], default="test")
    args = parser.parse_args(argv)
    if args.export:
        runs = load_run_records(args.runs_dir, ["C5", "C1"])
        judgments = {
            j: load_all_judgments(args.judgments_dir, j, ["C5", "C1"]) for j in ("llm", "nli")
        }
        outcomes = filter_split(
            build_outcomes(runs, judgments, load_thresholds(args.calibration)), args.split
        )
        sidecar = {
            str(r["chunk_id"]): str(r.get("text", ""))
            for r in iter_jsonl(args.judgments_dir / EVIDENCE_SIDECAR)
        }
        rows, keys = build_rows(outcomes, _evidence_lookup(runs), sidecar, args.n_items, args.seed)
        for path in write_sheet(rows, keys, args.out_dir):
            print(f"[human_annotation] wrote {path}")
        print(
            f"[human_annotation] {len(rows)} claims from {len({r['item_id'] for r in rows})} items"
        )
        return 0
    sheet = args.sheet or args.out_dir / f"{SHEET_NAME}.xlsx"
    result = score(sheet, args.out_dir / KEY_NAME)
    write_json(args.output, result)
    print(f"[human_annotation] {result['status']} → {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
