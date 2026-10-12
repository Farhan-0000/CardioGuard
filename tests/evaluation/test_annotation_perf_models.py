"""Human annotation export/scoring, before/after latency table and model-selection summary."""

import csv
import json
import re
from collections import Counter

import pytest
from openpyxl import load_workbook

from evaluation import human_annotation as ha
from evaluation import model_selection as ms
from evaluation import perf
from evaluation.io import load_run_records
from tests.evaluation.synthetic import write_study


def test_stratified_sample_proportional_and_seeded():
    items = {f"a{i}": "answerable_complete" for i in range(30)} | {
        f"u{i}": "out_of_domain" for i in range(10)
    }
    sample = ha.stratified_sample(items, 8, seed=1)
    counts = Counter(items[i] for i in sample)
    assert counts == {"answerable_complete": 6, "out_of_domain": 2}
    assert sample == ha.stratified_sample(items, 8, seed=1)
    assert sample != ha.stratified_sample(items, 8, seed=2)
    assert ha.stratified_sample(items, 100) == sorted(items)


def test_export_and_score(tmp_path):
    paths = write_study(tmp_path, conditions=("C1", "C5"))
    out_dir = tmp_path / "annotation"
    assert (
        ha.main(
            [
                "--export",
                "--runs-dir",
                str(paths["runs"]),
                "--judgments-dir",
                str(paths["judgments"]),
                "--out-dir",
                str(out_dir),
                "--n-items",
                "10",
                "--calibration",
                str(tmp_path / "none"),
            ]
        )
        == 0
    )
    wb = load_workbook(out_dir / "human_annotation_sheet.xlsx")
    assert wb.sheetnames == ["Instructions", "Claims"]
    claims = wb["Claims"]
    header = [c.value for c in claims[1]]
    assert header == list(ha.COLUMNS)
    validations = claims.data_validations.dataValidation
    assert validations and "SUPPORTED,UNSUPPORTED,CONTRADICTED" in validations[0].formula1
    with (out_dir / "human_annotation_sheet.csv").open() as fh:
        rows = list(csv.DictReader(fh))
    assert rows and len({r["item_id"] for r in rows}) <= 10
    assert all(re.fullmatch(r"R[0-9A-F]{8}", r["response_code"]) for r in rows)  # blinded codes
    assert "Guideline statement" in rows[0]["evidence"] or "Judge-retrieved" in rows[0]["evidence"]
    with (out_dir / "annotation_key.csv").open() as fh:
        keys = {k["row_id"]: k for k in csv.DictReader(fh)}
    assert set(keys) == {r["row_id"] for r in rows} and {k["condition"] for k in keys.values()} <= {
        "C1",
        "C5",
    }

    # unfilled → placeholder
    assert (
        ha.main(["--score", "--out-dir", str(out_dir), "--output", str(tmp_path / "agree.json")])
        == 0
    )
    assert json.loads((tmp_path / "agree.json").read_text())["status"] == "not_filled"

    # fill: A1 copies the LLM judge, A2 disagrees on the first row
    for i, r in enumerate(rows):
        r["A1_label"] = keys[r["row_id"]]["llm_label"]
        r["A2_label"] = "CONTRADICTED" if i == 0 else keys[r["row_id"]]["llm_label"].lower()
    filled = out_dir / "filled.csv"
    with filled.open("w", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(ha.COLUMNS))
        writer.writeheader()
        writer.writerows(rows)
    result = ha.score(filled, out_dir / "annotation_key.csv")
    assert result["status"] == "complete"
    assert result["A1_vs_llm"]["raw_agreement"] == 1.0
    assert result["human_human"]["n"] == len(rows)
    assert result["human_human"]["raw_agreement"] == (len(rows) - 1) / len(rows)
    assert result["consensus_vs_llm"]["n"] == len(rows) - 1


def test_ai_annotator_import_and_score(tmp_path):
    """AI labels go to separate files; the human sheet is untouched and stays blind."""
    paths = write_study(tmp_path, conditions=("C1", "C5"))
    out_dir = tmp_path / "annotation"
    common = ["--runs-dir", str(paths["runs"]), "--judgments-dir", str(paths["judgments"])]
    common += ["--out-dir", str(out_dir), "--calibration", str(tmp_path / "none")]
    assert ha.main(["--export", "--n-items", "10", *common]) == 0
    with (out_dir / "human_annotation_sheet.csv").open() as fh:
        rows = list(csv.DictReader(fh))
    with (out_dir / "annotation_key.csv").open() as fh:
        keys = {k["row_id"]: k for k in csv.DictReader(fh)}
    human_before = (out_dir / "human_annotation_sheet.xlsx").read_bytes()

    # the AI agrees with the LLM judge except on the first row, which it calls CONTRADICTED
    labels = [
        {
            "row_id": r["row_id"],
            "label": "CONTRADICTED" if i == 0 else keys[r["row_id"]]["llm_label"],
            "note": "conflicts with E1" if i == 0 else "",
        }
        for i, r in enumerate(rows)
    ]
    half = len(labels) // 2
    batches = [tmp_path / "labels_00.json", tmp_path / "labels_01.json"]
    batches[0].write_text(json.dumps(labels[:half]))
    batches[1].write_text(json.dumps(labels[half:]))
    assert ha.main(["--import-ai", *map(str, batches), "--out-dir", str(out_dir)]) == 0
    assert (out_dir / "human_annotation_sheet.xlsx").read_bytes() == human_before
    wb = load_workbook(out_dir / "ai_annotation_sheet.xlsx")
    header = [c.value for c in wb["Claims"][1]]
    assert header[-2:] == ["AI_label", "AI_notes"] and header[: len(ha.COLUMNS)] == list(ha.COLUMNS)
    assert "not a human" in wb["Instructions"]["A1"].value

    out = tmp_path / "ai.json"
    assert ha.main(["--score-ai", "--out-dir", str(out_dir), "--output", str(out)]) == 0
    result = json.loads(out.read_text())
    assert "not a human" in result["annotator"]
    assert result["n_claims"] == len(rows) and result["label_counts"]["CONTRADICTED"] >= 1
    assert result["ai_vs_llm"]["n"] == len(rows)
    assert result["ai_vs_llm"]["raw_agreement"] == (len(rows) - 1) / len(rows)
    assert set(result["per_condition"]) <= {"C1", "C5"}
    assert result["paired_c5_vs_c1"]["n_pairs"] >= 0
    assert result["labels_file"] == "evaluation/annotation/ai_annotation_labels.csv"


def test_ai_annotator_rejects_incomplete_or_invalid_labels(tmp_path):
    paths = write_study(tmp_path, conditions=("C1", "C5"))
    out_dir = tmp_path / "annotation"
    common = ["--runs-dir", str(paths["runs"]), "--judgments-dir", str(paths["judgments"])]
    common += ["--out-dir", str(out_dir), "--calibration", str(tmp_path / "none")]
    assert ha.main(["--export", "--n-items", "10", *common]) == 0
    sheet = out_dir / "human_annotation_sheet.xlsx"
    with (out_dir / "human_annotation_sheet.csv").open() as fh:
        ids = [r["row_id"] for r in csv.DictReader(fh)]
    partial = tmp_path / "partial.json"
    partial.write_text(json.dumps([{"row_id": ids[0], "label": "SUPPORTED"}]))
    with pytest.raises(ValueError, match="missing"):
        ha.import_ai_labels([partial], sheet, out_dir)
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps([{"row_id": i, "label": "MAYBE"} for i in ids]))
    with pytest.raises(ValueError, match="invalid label"):
        ha.import_ai_labels([bad], sheet, out_dir)


def _probe(label, p50s, actions):
    warm_summary = {
        f"{t}:{a}": {"n": 3, "p50": v, "p95": v * 1.2, "mean": v}
        for (t, v), a in zip(p50s.items(), actions)
    }
    cold = [
        {"scenario": t.split("-")[0], "turn": int(t.split("-t")[1]), "latency_s": v * 2}
        for t, v in p50s.items()
    ]
    return {"label": label, "base_url": "x", "cold": cold, "warm": [], "warm_summary": warm_summary}


def test_perf_table_and_figure(tmp_path):
    (tmp_path / "before_a.json").write_text(
        json.dumps(_probe("before", {"S1-t1": 0.02, "S5-t1": 30.0}, ["ESCALATE", "ANSWER"]))
    )
    (tmp_path / "after_b.json").write_text(
        json.dumps(_probe("after", {"S1-t1": 0.01, "S5-t1": 6.0}, ["ESCALATE", "GENERALIZE"]))
    )
    (tmp_path / "notes.json").write_text("[]")
    assert perf.latest_with_label(tmp_path, "after").name == "after_b.json"
    rows = perf.build_table(
        json.loads((tmp_path / "before_a.json").read_text()),
        json.loads((tmp_path / "after_b.json").read_text()),
    )
    s5 = next(r for r in rows if r["Scenario turn"] == "S5-t1")
    assert (
        s5["Speed-up (p50)"] == "5.0×"
        and s5["Action after"] == "GENERALIZE"
        and s5["Cold before s"] == "60.00"
    )
    assert perf.main(["--perf-dir", str(tmp_path), "--figures-dir", str(tmp_path / "fig")]) == 0
    assert (tmp_path / "latency_table.md").exists() and (
        tmp_path / "fig" / "latency_before_after.png"
    ).exists()
    assert perf.main(["--perf-dir", str(tmp_path / "empty")]) == 2


def test_model_selection_summary(tmp_path):
    paths = write_study(tmp_path, conditions=("C5",))
    records = load_run_records(paths["runs"])["C5"]
    for rec in records[:5]:
        rec["llm_calls"][0]["cached"] = False  # synthetic extractor calls are memoised by default
    records[0]["llm_calls"][0]["attempts"] = 2
    summary = ms.summarize_model(
        "qwen3:8b", records, {"parameter_size": "8.2B", "size_bytes": 5_225_388_164}
    )
    roles = summary["roles"]
    extractor = roles["context_extractor"]
    assert extractor["n_cached"] == extractor["n_calls"] - 5
    assert extractor["valid_first_attempt"] == {**extractor["valid_first_attempt"], "k": 4, "n": 5}
    assert extractor["valid_any_attempt"]["k"] == 5
    assert roles["response_generator"]["tokens_per_s_median"] > 0
    assert summary["routing"]["accuracy_strict"]["n"] == len(records)
    assert summary["latency_turn_p50_ms"] > 0
    out = ms.write_outputs(
        {"split": "dev", "models": ["qwen3:8b"], "results": {"qwen3:8b": summary}},
        tmp_path / "ms",
        tmp_path / "fig",
    )
    assert {p.name for p in out} >= {
        "summary.json",
        "model_selection.md",
        "model_selection.csv",
        "model_selection.png",
        "model_selection.svg",
    }
    assert ms.safe_name("qwen3.5:4b") == "qwen3.5_4b"


def test_model_selection_never_trips_the_outage_breaker(tmp_path, monkeypatch):
    """LLM failures are the measured outcome of model selection, not an outage."""
    import subprocess
    import types

    seen = []
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda cmd, check=False: seen.append(cmd) or types.SimpleNamespace(returncode=0),
    )
    assert ms.run_model("m:1b", "dev", 4, tmp_path, tmp_path / "d.json") == 0
    cmd = seen[0]
    assert cmd[cmd.index("--max-consecutive-degraded") + 1] == "0"
